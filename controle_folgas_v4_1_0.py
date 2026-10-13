# -*- coding: utf-8 -*-
"""
Controle de Folgas e Banco de Horas - Windows 11 - Versão 4.1.0 - splash de inicialização, instância única e atualização automática via GitHub Releases
Responsável pelo projeto: Thiago Luiz Moro (desenvolvimento com assistência de IA)

Bibliotecas recomendadas:
    pip install customtkinter openpyxl pywin32 pillow

Para gerar EXE com splash de carregamento:
    1) python criar_splash_controle_folgas.py
    2) pyinstaller --clean --noconfirm Controle_de_Folgas_v4_1_0.spec

O arquivo .spec configura um splash nativo estático do PyInstaller durante a
extração do --onefile. Assim que o Python inicia, o programa troca automaticamente
para uma tela de carregamento própria, que mostra as etapas reais da inicialização.

Observação:
- O sistema salva os dados em um banco SQLite chamado controle_folgas.db na mesma pasta do programa.
- A preferência de tema é salva em controle_folgas_config.json na mesma pasta do programa.
- O saldo é calculado em minutos.
- Cada funcionário possui jornada individual de 6, 8 ou 12 horas.
- Backups automáticos são criados antes de alterações e ações administrativas.
- Exclusões usam Lixeira, senha administrativa, motivo obrigatório, confirmação, backup e auditoria.
- Exclusões definitivas da Lixeira exigem código de suporte assinado para cada operação.
- Alterações de lançamentos e funcionários são registradas na auditoria com os valores anteriores e novos.
- Cadastro e edição de funcionários e lançamentos usam janelas separadas para evitar confusão.
- Dashboard e Lançamentos usam tabelas profissionais; Funcionários usa uma lista administrativa responsiva em CustomTkinter.
- A recuperação de suporte usa desafio/resposta Ed25519, sem senha mestra dentro do aplicativo.
"""

import csv
import json
import os
import re
import sqlite3
import sys
import webbrowser
import ctypes
import hashlib
import hmac
import secrets
import tempfile
import socket
import atexit
import queue
import subprocess
import threading
import urllib.error
import urllib.request
from datetime import datetime, timedelta
from pathlib import Path
from calendar import monthrange
from dataclasses import dataclass
from typing import Callable

from database_backup import create_sqlite_snapshot
from support_security import (
    SupportAuthorizationError, issue_request, verify_and_consume,
)


# ---------------------------------------------------------------------------
# Inicialização rápida: splash nativo do PyInstaller + instância única
# ---------------------------------------------------------------------------
# O splash do PyInstaller aparece ainda durante a extração do executável
# --onefile, antes da interface principal ficar pronta. O módulo pyi_splash
# só existe no EXE criado com a opção --splash; durante o desenvolvimento pelo
# Python estas funções simplesmente não fazem nada.
_STARTUP_SPLASH = None
_APP_STARTUP_WINDOW = None
_APP_STARTUP_STATUS_LABEL = None
_APP_STARTUP_IMAGE = None
_SINGLE_INSTANCE_MUTEX_HANDLE = None
_EARLY_APP_NAME = "Controle de Folgas"
_SINGLE_INSTANCE_MUTEX_NAME = r"Local\ControleDeFolgas_SingleInstance"


def _startup_splash_module():
    global _STARTUP_SPLASH
    if _STARTUP_SPLASH is False:
        return None
    if _STARTUP_SPLASH is not None:
        return _STARTUP_SPLASH
    try:
        import pyi_splash  # type: ignore[import-not-found]

        if pyi_splash.is_alive():
            _STARTUP_SPLASH = pyi_splash
            return pyi_splash
    except Exception:
        pass
    _STARTUP_SPLASH = False
    return None


def startup_splash_update(message: str) -> None:
    """Atualiza o splash nativo e a tela própria de carregamento, quando existirem."""
    text = str(message)

    splash = _startup_splash_module()
    if splash is not None:
        try:
            splash.update_text(text)
        except Exception:
            pass

    label = globals().get("_APP_STARTUP_STATUS_LABEL")
    window = globals().get("_APP_STARTUP_WINDOW")
    if label is not None:
        try:
            label.configure(text=text)
        except Exception:
            pass
    if window is not None:
        try:
            window.update_idletasks()
            window.update()
        except Exception:
            pass


def _close_native_startup_splash() -> None:
    """Fecha somente o splash do bootloader do PyInstaller."""
    global _STARTUP_SPLASH
    splash = _startup_splash_module()
    if splash is not None:
        try:
            splash.close()
        except Exception:
            pass
    _STARTUP_SPLASH = False


def startup_splash_close() -> None:
    """Fecha todas as telas de carregamento da inicialização."""
    global _APP_STARTUP_WINDOW, _APP_STARTUP_STATUS_LABEL, _APP_STARTUP_IMAGE
    _close_native_startup_splash()
    window = _APP_STARTUP_WINDOW
    _APP_STARTUP_WINDOW = None
    _APP_STARTUP_STATUS_LABEL = None
    _APP_STARTUP_IMAGE = None
    if window is not None:
        try:
            window.destroy()
        except Exception:
            pass


def _release_single_instance_mutex() -> None:
    global _SINGLE_INSTANCE_MUTEX_HANDLE
    handle = _SINGLE_INSTANCE_MUTEX_HANDLE
    _SINGLE_INSTANCE_MUTEX_HANDLE = None
    if not handle or not sys.platform.startswith("win"):
        return
    try:
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel32.CloseHandle.argtypes = [ctypes.c_void_p]
        kernel32.CloseHandle.restype = ctypes.c_int
        kernel32.CloseHandle(handle)
    except Exception:
        pass


def acquire_single_instance() -> bool:
    """Garante apenas uma instância do Controle de Folgas por sessão do Windows.

    O mutex é criado antes dos imports pesados da interface. Se o usuário clicar
    várias vezes no atalho enquanto o primeiro processo ainda está carregando,
    as execuções extras são encerradas sem abrir outro banco ou outra janela.
    """
    global _SINGLE_INSTANCE_MUTEX_HANDLE

    if not sys.platform.startswith("win"):
        return True

    try:
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel32.CreateMutexW.argtypes = [ctypes.c_void_p, ctypes.c_int, ctypes.c_wchar_p]
        kernel32.CreateMutexW.restype = ctypes.c_void_p
        kernel32.CloseHandle.argtypes = [ctypes.c_void_p]
        kernel32.CloseHandle.restype = ctypes.c_int

        ctypes.set_last_error(0)
        handle = kernel32.CreateMutexW(None, 0, _SINGLE_INSTANCE_MUTEX_NAME)
        if not handle:
            # Falha ao criar o mutex não deve impedir o uso normal do sistema.
            return True

        error = ctypes.get_last_error()
        ERROR_ALREADY_EXISTS = 183
        if error == ERROR_ALREADY_EXISTS:
            try:
                kernel32.CloseHandle(handle)
            except Exception:
                pass
            return False

        _SINGLE_INSTANCE_MUTEX_HANDLE = handle
        atexit.register(_release_single_instance_mutex)
        return True
    except Exception:
        # Proteção defensiva: se a API do Windows falhar, o programa continua.
        return True


def show_already_running_message() -> None:
    """Informa de forma leve que o primeiro clique já está sendo processado."""
    if not sys.platform.startswith("win"):
        return
    try:
        MB_OK = 0x00000000
        MB_ICONINFORMATION = 0x00000040
        ctypes.windll.user32.MessageBoxW(
            None,
            "O Controle de Folgas já está aberto ou está sendo iniciado.\n\n"
            "Aguarde a primeira abertura ser concluída.",
            _EARLY_APP_NAME,
            MB_OK | MB_ICONINFORMATION,
        )
    except Exception:
        pass


startup_splash_update("Iniciando Controle de Folgas...")

if not acquire_single_instance():
    startup_splash_close()
    show_already_running_message()
    raise SystemExit(0)

startup_splash_update("Carregando interface...")

import tkinter as tk
import tkinter.font as tkfont
from tkinter import filedialog, ttk

startup_splash_update("Carregando componentes visuais...")

try:
    import customtkinter as ctk
except ImportError:
    raise SystemExit("Instale a biblioteca customtkinter: pip install customtkinter")


startup_splash_update("Carregando recursos gráficos...")

try:
    from PIL import Image
except ImportError:
    raise SystemExit("Instale a biblioteca Pillow: pip install pillow")


startup_splash_update("Carregando relatórios e planilhas...")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins
except Exception:
    Workbook = None

startup_splash_update("Preparando integração com o Windows...")

try:
    import win32com.client as win32
except ImportError:
    win32 = None

APP_NAME = "Controle de Folgas"
APP_VERSION = "4.1.0"

# Atualização automática via GitHub Releases.
# O repositório deve ser público e usado somente para publicar os executáveis.
GITHUB_OWNER = "thiagolmoro"
GITHUB_REPOSITORY = "controle-de-folgas-releases"
GITHUB_API_VERSION = "2026-03-10"
GITHUB_LATEST_RELEASE_URL = (
    f"https://api.github.com/repos/{GITHUB_OWNER}/{GITHUB_REPOSITORY}/releases/latest"
)
UPDATE_REQUEST_TIMEOUT_SECONDS = 12
UPDATE_PREFERRED_ASSET_NAMES = (
    "Controle de Folgas.exe",
    "Controle_de_Folgas.exe",
    "Controle.Folgas.exe",
    "Controle.de.Folgas.exe",
)

DB_NAME = "controle_folgas.db"
CONFIG_NAME = "controle_folgas_config.json"
LOGO_DIR_NAME = "Logo"
MAIN_LOGO_NAME = "logo 3D.png"
APP_ICON_ICO_NAME = "logo.ico"
APP_ICON_PNG_NAME = "logo.png"
STATUS_DIR_NAME = "Status_Programa"
STATUS_HEARTBEAT_MS = 60_000
DEFAULT_HOURS_PER_DAY = 12
ADMIN_MAX_ATTEMPTS = 5
ADMIN_LOCK_MINUTES = 5
PBKDF2_ITERATIONS = 310_000

TIPOS_PADRAO = sorted([
    "Banco de horas",
    "Capacitação",
    "Chegou mais cedo",
    "Chegou mais tarde",
    "Curso",
    "Entrou mais cedo",
    "Entrou mais tarde",
    "Evento",
    "Exame médico",
    "Faltou",
    "Feriado",
    "Folga",
    "Reunião",
    "Reunião geral",
    "Saiu mais cedo",
    "Saiu mais tarde",
    "Trabalhou fora do plantão",
    "Troca de plantão",
], key=str.casefold) + ["Outro"]

# Categorias que podem representar tanto crédito quanto débito.
# Nessas opções, o usuário escolhe explicitamente como o lançamento afeta o saldo.
CATEGORIAS_MOVIMENTO_MANUAL = {
    "banco de horas",
    "exame médico",
    "troca de plantão",
    "outro",
}

CATEGORIA_PLACEHOLDER = "Selecione uma categoria"

MESES_FILTRO = [
    "Todos",
    "01 - Janeiro",
    "02 - Fevereiro",
    "03 - Março",
    "04 - Abril",
    "05 - Maio",
    "06 - Junho",
    "07 - Julho",
    "08 - Agosto",
    "09 - Setembro",
    "10 - Outubro",
    "11 - Novembro",
    "12 - Dezembro",
]

MESES_NOMES = {
    "01": "Janeiro",
    "02": "Fevereiro",
    "03": "Março",
    "04": "Abril",
    "05": "Maio",
    "06": "Junho",
    "07": "Julho",
    "08": "Agosto",
    "09": "Setembro",
    "10": "Outubro",
    "11": "Novembro",
    "12": "Dezembro",
}

CATEGORIAS_CREDITO = {
    "feriado",
    "capacitação",
    "curso",
    "reunião",
    "reunião geral",
    "banco de horas",
    "saiu mais tarde",
    "chegou mais cedo",
    "entrou mais cedo",
    "trabalhou fora do plantão",
    "troca de plantão",
    "evento",
}

CATEGORIAS_DEBITO = {
    "folga",
    "faltou",
    "saiu mais cedo",
    "chegou mais tarde",
    "entrou mais tarde",
    "ausentou-se",
}


MOVIMENTO_CREDITO = "Crédito"
MOVIMENTO_DEBITO = "Débito"
MOVIMENTO_CREDITO_LABEL = "Crédito (+) — Aumenta o saldo"
MOVIMENTO_DEBITO_LABEL = "Débito (−) — Diminui o saldo"

MOVIMENTO_CREDITO_AJUDA = (
    "Aumenta o saldo do funcionário. Ex.: trabalhou a mais, chegou mais cedo ou saiu mais tarde."
)
MOVIMENTO_DEBITO_AJUDA = (
    "Diminui o saldo do funcionário. Ex.: folga, saiu mais cedo, chegou mais tarde ou faltou."
)


# Links oficiais dos tutoriais em vídeo.
# Mantidos em um único ponto para facilitar futuras substituições.
TUTORIAIS_YOUTUBE = [
    {
        "numero": "Tutorial 1 de 6",
        "titulo": "Primeiro Acesso e Apresentação do Programa",
        "descricao": "Primeiro acesso, cadastro do nome do Serviço e apresentação das principais telas do sistema.",
        "videos": [
            ("Assistir tutorial", "https://youtu.be/EtnG59NAn8U"),
        ],
    },
    {
        "numero": "Tutorial 2 de 6",
        "titulo": "Cadastro e Edição de Funcionários",
        "descricao": "Cadastro, pesquisa e edição de funcionários, regimes de trabalho, jornadas e situação Ativo/Inativo.",
        "videos": [
            ("Assistir tutorial", "https://youtu.be/NJWjp6g710k"),
        ],
    },
    {
        "numero": "Tutorial 3 de 6",
        "titulo": "Cadastro, Edição e Exclusão de Lançamentos",
        "descricao": "Aprenda a cadastrar movimentações e, na segunda parte, editar registros e enviá-los para a Lixeira.",
        "videos": [
            ("Parte 1 — Cadastro", "https://youtu.be/YL0vK2svUBs"),
            ("Parte 2 — Edição e Exclusão", "https://youtu.be/lj7M64gU-lc"),
        ],
    },
    {
        "numero": "Tutorial 4 de 6",
        "titulo": "Exportando para o Excel",
        "descricao": "Uso dos filtros, geração do arquivo Excel e consulta das planilhas de lançamentos, resumo e assinatura.",
        "videos": [
            ("Assistir tutorial", "https://youtu.be/85uR8AHW3m8"),
        ],
    },
    {
        "numero": "Tutorial 5 de 6",
        "titulo": "Impressão Direta",
        "descricao": "Impressão das folhas de assinatura diretamente pelo programa, com seleção de período e quantidade de cópias.",
        "videos": [
            ("Assistir tutorial", "https://youtu.be/iniFKCzDZqU"),
        ],
    },
    {
        "numero": "Tutorial 6 de 6",
        "titulo": "Opções Administrativas",
        "descricao": "Lixeira, auditoria, nome do Serviço, senha administrativa e recuperação de acesso.",
        "videos": [
            ("Parte 1 — Administração", "https://youtu.be/dtKZr5V-Oew"),
            ("Parte 2 — Senha e Recuperação", "https://youtu.be/H61QiMEQK9E"),
        ],
    },
]

# Se desejar exibir também um botão para a playlist completa, informe o link abaixo.
# Enquanto estiver vazio, o botão da playlist não aparece na interface.
PLAYLIST_YOUTUBE_URL = ""


def movimento_para_exibicao(value: str) -> str:
    """Converte o valor interno Crédito/Débito para o texto explicativo da interface."""
    value = normalize_spaces(value)
    if value.startswith(MOVIMENTO_DEBITO):
        return MOVIMENTO_DEBITO_LABEL
    return MOVIMENTO_CREDITO_LABEL


def normalizar_movimento(value: str) -> str:
    """Mantém no banco apenas Crédito/Débito, mesmo com rótulos explicativos na interface."""
    value = normalize_spaces(value)
    if value.startswith(MOVIMENTO_CREDITO):
        return MOVIMENTO_CREDITO
    if value.startswith(MOVIMENTO_DEBITO):
        return MOVIMENTO_DEBITO
    raise ValueError("Selecione um tipo de movimento válido: Crédito ou Débito.")


def app_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def resource_path(*parts) -> Path:
    """Localiza recursos tanto no código-fonte quanto dentro do EXE do PyInstaller."""
    relative = Path(*parts)

    # Em --onefile, os arquivos adicionados com --add-data são extraídos
    # temporariamente para sys._MEIPASS.
    bundle_dir = getattr(sys, "_MEIPASS", None)
    if bundle_dir:
        bundled = Path(bundle_dir) / relative
        if bundled.exists():
            return bundled

    # Durante o desenvolvimento e também como fallback ao lado do executável.
    return app_dir() / relative


def create_app_startup_window(master) -> None:
    """Exibe uma tela de carregamento própria assim que o Python já iniciou.

    O splash nativo do PyInstaller cobre a fase de extração do --onefile. Esta
    janela assume em seguida e permanece durante banco, interface e leitura dos
    dados. Usa somente Tk, portanto não depende do tema ou do CustomTkinter para
    ser desenhada.
    """
    global _APP_STARTUP_WINDOW, _APP_STARTUP_STATUS_LABEL, _APP_STARTUP_IMAGE

    if _APP_STARTUP_WINDOW is not None:
        return

    try:
        window = tk.Toplevel(master)
        window.withdraw()
        window.overrideredirect(True)
        window.configure(bg="#0c182c", highlightthickness=1, highlightbackground="#24476f")
        try:
            window.attributes("-topmost", True)
        except Exception:
            pass

        width, height = 620, 392
        screen_w = window.winfo_screenwidth()
        screen_h = window.winfo_screenheight()
        x = max(0, (screen_w - width) // 2)
        y = max(0, (screen_h - height) // 2)
        window.geometry(f"{width}x{height}+{x}+{y}")

        image_path = resource_path(LOGO_DIR_NAME, "splash.png")
        image_loaded = False
        if image_path.exists():
            try:
                photo = tk.PhotoImage(file=str(image_path))
                image_label = tk.Label(window, image=photo, bd=0, highlightthickness=0, bg="#0c182c")
                image_label.place(x=0, y=0, width=620, height=360)
                _APP_STARTUP_IMAGE = photo
                image_loaded = True
            except Exception:
                _APP_STARTUP_IMAGE = None

        if not image_loaded:
            tk.Label(
                window,
                text="Controle de Folgas e Banco de Horas",
                font=("Segoe UI", 20, "bold"),
                bg="#0c182c",
                fg="#f2f7ff",
            ).pack(pady=(92, 8))
            tk.Label(
                window,
                text=f"Versão {APP_VERSION}",
                font=("Segoe UI", 11),
                bg="#0c182c",
                fg="#76beff",
            ).pack()

        status = tk.Label(
            window,
            text="Inicializando o sistema...",
            font=("Segoe UI", 10),
            bg="#101f36",
            fg="#d7e3f4",
            anchor="w",
            padx=18,
        )
        status.place(x=0, y=360, width=620, height=32)

        _APP_STARTUP_WINDOW = window
        _APP_STARTUP_STATUS_LABEL = status
        window.deiconify()
        window.lift()
        window.update_idletasks()
        window.update()

        # A tela própria já está desenhada; não precisamos manter a janela
        # experimental do bootloader aberta em paralelo.
        _close_native_startup_splash()
    except Exception:
        _APP_STARTUP_WINDOW = None
        _APP_STARTUP_STATUS_LABEL = None
        _APP_STARTUP_IMAGE = None


def main_logo_path() -> Path:
    return resource_path(LOGO_DIR_NAME, MAIN_LOGO_NAME)


def icon_resource_path() -> Path | None:
    """Prioriza o ICO e aceita PNG como fallback durante o desenvolvimento."""
    for filename in (APP_ICON_ICO_NAME, APP_ICON_PNG_NAME, MAIN_LOGO_NAME):
        path = resource_path(LOGO_DIR_NAME, filename)
        if path.exists():
            return path
    return None


def apply_window_icon(window) -> None:
    """Aplica o ícone do sistema ao root, CTkToplevel e Toplevel do Tk.

    O CustomTkinter pode aplicar o próprio ícone aproximadamente 200 ms após
    criar uma CTk/CTkToplevel no Windows. Por isso, além da aplicação imediata,
    repetimos o iconbitmap depois desse intervalo para garantir que o ícone
    personalizado seja o último a ser definido na barra de título.
    """

    ico_path = resource_path(LOGO_DIR_NAME, APP_ICON_ICO_NAME)

    def _apply():
        try:
            if not window.winfo_exists():
                return
        except Exception:
            return

        icon_applied = False

        # No Windows, o formato ICO é o caminho mais confiável para o ícone da
        # barra de título. Usamos iconbitmap sem default= para aplicar à janela
        # concreta que acabou de ser criada.
        if ico_path.exists():
            try:
                window.iconbitmap(str(ico_path))
                icon_applied = True
            except Exception:
                try:
                    window.wm_iconbitmap(str(ico_path))
                    icon_applied = True
                except Exception:
                    pass

        # Marcar explicitamente a chamada evita que versões do CustomTkinter
        # que consultam este atributo restaurem o ícone padrão posteriormente.
        try:
            if hasattr(window, "_iconbitmap_method_called"):
                window._iconbitmap_method_called = True
        except Exception:
            pass

        # Fallback para plataformas/ambientes em que iconbitmap não aceite ICO.
        # No Windows, se o ICO funcionou, não chamamos iconphoto em seguida para
        # não substituir o ícone de título que já foi aplicado corretamente.
        if not icon_applied:
            for png_name in (APP_ICON_PNG_NAME, MAIN_LOGO_NAME):
                png_path = resource_path(LOGO_DIR_NAME, png_name)
                if not png_path.exists():
                    continue
                try:
                    photo = tk.PhotoImage(file=str(png_path))
                    window.iconphoto(False, photo)
                    # Tk exige que a referência permaneça viva.
                    window._app_icon_photo = photo
                    break
                except Exception:
                    continue

    # Aplicação imediata para a janela principal e Toplevels Tk puros.
    _apply()

    # CTk/CTkToplevel pode recolocar o próprio ícone por volta de 200 ms.
    # As reaplicações abaixo acontecem depois desse ponto e também cobrem casos
    # em que o handle da janela é recriado por DPI, maximização ou RDP.
    for delay in (260, 420, 700):
        try:
            window.after(delay, _apply)
        except Exception:
            pass



class UpdateError(RuntimeError):
    """Erro controlado do mecanismo de atualização automática."""


@dataclass(frozen=True)
class ReleaseInfo:
    version: str
    tag_name: str
    title: str
    notes: str
    html_url: str
    asset_name: str
    download_url: str
    asset_size: int
    sha256: str


@dataclass(frozen=True)
class DownloadedUpdate:
    release: ReleaseInfo
    file_path: Path


UpdateProgressCallback = Callable[[int, int], None]


def _github_api_request(url: str) -> urllib.request.Request:
    return urllib.request.Request(
        url,
        headers={
            "Accept": "application/vnd.github+json",
            "User-Agent": "Controle-de-Folgas-Updater",
            "X-GitHub-Api-Version": GITHUB_API_VERSION,
        },
    )


def _github_download_request(url: str) -> urllib.request.Request:
    return urllib.request.Request(
        url,
        headers={
            "Accept": "application/octet-stream",
            "User-Agent": "Controle-de-Folgas-Updater",
        },
    )


def normalize_version_text(value: str) -> str:
    value = str(value or "").strip()
    if value.lower().startswith("v"):
        value = value[1:]
    return value.strip()


def version_key(value: str) -> tuple[int, int, int, tuple]:
    """Converte versões semânticas MAJOR.MINOR.PATCH em uma chave comparável."""
    value = normalize_version_text(value)
    match = re.fullmatch(
        r"(\d+)(?:\.(\d+))?(?:\.(\d+))?(?:-([0-9A-Za-z.-]+))?(?:\+[0-9A-Za-z.-]+)?",
        value,
    )
    if not match:
        raise UpdateError(f"Versão inválida recebida: {value!r}")

    major = int(match.group(1) or 0)
    minor = int(match.group(2) or 0)
    patch = int(match.group(3) or 0)
    prerelease = match.group(4)

    if prerelease is None:
        prerelease_key = (1,)
    else:
        parts = []
        for part in prerelease.split("."):
            if part.isdigit():
                parts.append((0, int(part)))
            else:
                parts.append((1, part.casefold()))
        prerelease_key = (0, *parts)

    return major, minor, patch, prerelease_key


def is_newer_version(latest: str, current: str) -> bool:
    return version_key(latest) > version_key(current)


def _select_update_asset(assets: list[dict]) -> dict:
    uploaded = [
        asset
        for asset in assets
        if str(asset.get("state") or "").lower() == "uploaded"
    ]

    for preferred in UPDATE_PREFERRED_ASSET_NAMES:
        for asset in uploaded:
            if str(asset.get("name") or "").casefold() == preferred.casefold():
                return asset

    executables = [
        asset
        for asset in uploaded
        if str(asset.get("name") or "").lower().endswith(".exe")
    ]
    if len(executables) == 1:
        return executables[0]
    if not executables:
        raise UpdateError(
            "A Release mais recente não possui um executável (.exe) para atualização."
        )
    raise UpdateError(
        "A Release mais recente possui mais de um executável e não foi possível "
        "identificar automaticamente o arquivo do Controle de Folgas."
    )


def _extract_release_sha256(asset: dict) -> str:
    digest = str(asset.get("digest") or "").strip()
    if not digest:
        raise UpdateError(
            "O GitHub não informou o SHA-256 do executável desta Release. "
            "A atualização automática foi bloqueada por segurança."
        )

    algorithm, separator, value = digest.partition(":")
    if not separator or algorithm.casefold() != "sha256":
        raise UpdateError(
            "O hash publicado para esta Release não está no formato SHA-256 esperado."
        )

    value = value.strip().lower()
    if not re.fullmatch(r"[0-9a-f]{64}", value):
        raise UpdateError("O SHA-256 publicado para esta Release é inválido.")
    return value


def get_latest_release() -> ReleaseInfo:
    try:
        with urllib.request.urlopen(
            _github_api_request(GITHUB_LATEST_RELEASE_URL),
            timeout=UPDATE_REQUEST_TIMEOUT_SECONDS,
        ) as response:
            payload = response.read().decode("utf-8", errors="replace")
    except urllib.error.HTTPError as exc:
        if exc.code == 404:
            raise UpdateError(
                "Nenhuma Release publicada foi encontrada no repositório de atualizações."
            ) from exc
        raise UpdateError(
            f"O GitHub respondeu com erro HTTP {exc.code} ao verificar atualizações."
        ) from exc
    except urllib.error.URLError as exc:
        raise UpdateError(
            "Não foi possível acessar o servidor de atualizações. Verifique a conexão com a internet."
        ) from exc
    except TimeoutError as exc:
        raise UpdateError("A verificação de atualizações excedeu o tempo limite.") from exc

    try:
        data = json.loads(payload)
    except json.JSONDecodeError as exc:
        raise UpdateError("O GitHub retornou uma resposta inválida.") from exc

    tag_name = str(data.get("tag_name") or "").strip()
    if not tag_name:
        raise UpdateError("A Release mais recente não possui uma versão identificável.")

    version = normalize_version_text(tag_name)
    version_key(version)

    assets = data.get("assets") or []
    if not isinstance(assets, list):
        assets = []
    asset = _select_update_asset(assets)

    download_url = str(asset.get("browser_download_url") or "").strip()
    if not download_url.startswith("https://"):
        raise UpdateError("A Release mais recente não possui uma URL de download válida.")

    try:
        asset_size = int(asset.get("size") or 0)
    except (TypeError, ValueError):
        asset_size = 0

    return ReleaseInfo(
        version=version,
        tag_name=tag_name,
        title=str(data.get("name") or tag_name).strip(),
        notes=str(data.get("body") or "").strip(),
        html_url=str(data.get("html_url") or "").strip(),
        asset_name=str(asset.get("name") or "Controle de Folgas.exe").strip(),
        download_url=download_url,
        asset_size=max(0, asset_size),
        sha256=_extract_release_sha256(asset),
    )


def current_version_for_update_check(real_version: str) -> str:
    """Permite simular uma versão antiga somente ao executar pelo Python.

    Exemplo no PowerShell, somente para testar a DETECÇÃO de uma atualização:
        $env:CONTROLE_FOLGAS_UPDATE_TEST_VERSION="3.10.0"
        python controle_folgas_v4_1_0.py

    Em um EXE PyInstaller a variável é ignorada.
    """
    if not getattr(sys, "frozen", False):
        override = str(os.environ.get("CONTROLE_FOLGAS_UPDATE_TEST_VERSION") or "").strip()
        if override:
            return normalize_version_text(override)
    return normalize_version_text(real_version)


def download_update(
    release: ReleaseInfo,
    progress_callback: UpdateProgressCallback | None = None,
) -> DownloadedUpdate:
    update_dir = Path(tempfile.gettempdir()) / "Controle_Folgas_Update"
    update_dir.mkdir(parents=True, exist_ok=True)

    destination = update_dir / f"Controle_de_Folgas_{release.version}.exe"
    partial = destination.with_suffix(".download")

    for candidate in (partial, destination):
        try:
            candidate.unlink(missing_ok=True)
        except Exception:
            pass

    downloaded = 0
    digest = hashlib.sha256()

    try:
        with urllib.request.urlopen(
            _github_download_request(release.download_url),
            timeout=60,
        ) as response, partial.open("wb") as output:
            header_size = response.headers.get("Content-Length")
            try:
                total = int(header_size or release.asset_size or 0)
            except (TypeError, ValueError):
                total = release.asset_size

            while True:
                chunk = response.read(1024 * 256)
                if not chunk:
                    break
                output.write(chunk)
                digest.update(chunk)
                downloaded += len(chunk)
                if progress_callback is not None:
                    progress_callback(downloaded, total)
    except urllib.error.URLError as exc:
        try:
            partial.unlink(missing_ok=True)
        except Exception:
            pass
        raise UpdateError(
            "Não foi possível baixar a nova versão. Verifique a conexão com a internet."
        ) from exc
    except Exception as exc:
        try:
            partial.unlink(missing_ok=True)
        except Exception:
            pass
        if isinstance(exc, UpdateError):
            raise
        raise UpdateError(f"Falha ao baixar a atualização: {exc}") from exc

    calculated = digest.hexdigest().lower()
    if calculated != release.sha256.lower():
        try:
            partial.unlink(missing_ok=True)
        except Exception:
            pass
        raise UpdateError(
            "O arquivo baixado não passou na verificação SHA-256. "
            "A atualização foi cancelada por segurança."
        )

    try:
        partial.replace(destination)
    except Exception as exc:
        raise UpdateError(f"Não foi possível preparar o arquivo de atualização: {exc}") from exc

    return DownloadedUpdate(release=release, file_path=destination)


def can_auto_install_update() -> bool:
    return bool(getattr(sys, "frozen", False)) and sys.platform.startswith("win")


def current_executable_path() -> Path:
    if not can_auto_install_update():
        raise UpdateError(
            "A instalação automática só é executada pelo arquivo .exe no Windows."
        )
    return Path(sys.executable).resolve()


def _powershell_quote(value: str | Path) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def _verify_update_directory_writable(executable: Path) -> None:
    probe = executable.parent / ".controle_folgas_update_write_test.tmp"
    try:
        probe.write_text("ok", encoding="utf-8")
        probe.unlink(missing_ok=True)
    except Exception as exc:
        raise UpdateError(
            "A pasta onde o programa está instalado não permite substituir o executável. "
            "Mova o programa para uma pasta gravável ou execute-o com permissão adequada."
        ) from exc


def launch_update_installer(downloaded: DownloadedUpdate) -> Path:
    """Abre um PowerShell oculto que troca somente o EXE após o programa encerrar."""
    executable = current_executable_path()
    _verify_update_directory_writable(executable)

    update_dir = downloaded.file_path.parent
    script_path = update_dir / "controle_folgas_apply_update.ps1"
    backup_path = Path(str(executable) + ".bak")
    log_path = update_dir / "controle_folgas_update.log"
    pid = os.getpid()

    script = f"""$ErrorActionPreference = 'Stop'
$CurrentExe = {_powershell_quote(executable)}
$NewExe = {_powershell_quote(downloaded.file_path)}
$BackupExe = {_powershell_quote(backup_path)}
$LogFile = {_powershell_quote(log_path)}
$ProcessId = {pid}

function Write-UpdateLog([string]$Message) {{
    try {{
        $stamp = Get-Date -Format 'yyyy-MM-dd HH:mm:ss'
        Add-Content -LiteralPath $LogFile -Value \"[$stamp] $Message\" -Encoding UTF8
    }} catch {{}}
}}

Write-UpdateLog 'Aguardando o Controle de Folgas encerrar.'
for ($i = 0; $i -lt 120; $i++) {{
    $running = Get-Process -Id $ProcessId -ErrorAction SilentlyContinue
    if (-not $running) {{ break }}
    Start-Sleep -Milliseconds 500
}}

try {{
    if (Test-Path -LiteralPath $BackupExe) {{
        Remove-Item -LiteralPath $BackupExe -Force -ErrorAction SilentlyContinue
    }}

    Copy-Item -LiteralPath $CurrentExe -Destination $BackupExe -Force
    Write-UpdateLog 'Backup do executável anterior criado.'

    $updated = $false
    for ($attempt = 1; $attempt -le 30; $attempt++) {{
        try {{
            Copy-Item -LiteralPath $NewExe -Destination $CurrentExe -Force
            $updated = $true
            break
        }} catch {{
            Start-Sleep -Milliseconds 700
        }}
    }}

    if (-not $updated) {{
        throw 'Não foi possível substituir o executável após várias tentativas.'
    }}

    Write-UpdateLog 'Executável atualizado com sucesso.'

    Remove-Item Env:_MEIPASS2 -ErrorAction SilentlyContinue
    Get-ChildItem Env: |
        Where-Object {{ $_.Name -like '_PYI_*' }} |
        ForEach-Object {{
            Remove-Item (\"Env:\" + $_.Name) -ErrorAction SilentlyContinue
        }}
    $env:PYINSTALLER_RESET_ENVIRONMENT = '1'

    Start-Process -FilePath $CurrentExe -WorkingDirectory (Split-Path -Parent $CurrentExe)
    Start-Sleep -Seconds 2
    Remove-Item -LiteralPath $NewExe -Force -ErrorAction SilentlyContinue
    Write-UpdateLog 'Nova versão iniciada.'
}} catch {{
    Write-UpdateLog (\"Falha na atualização: \" + $_.Exception.Message)
    try {{
        if (Test-Path -LiteralPath $BackupExe) {{
            Copy-Item -LiteralPath $BackupExe -Destination $CurrentExe -Force
            Write-UpdateLog 'Versão anterior restaurada.'
        }}
    }} catch {{
        Write-UpdateLog (\"Falha ao restaurar backup: \" + $_.Exception.Message)
    }}

    try {{
        Remove-Item Env:_MEIPASS2 -ErrorAction SilentlyContinue
        Get-ChildItem Env: |
            Where-Object {{ $_.Name -like '_PYI_*' }} |
            ForEach-Object {{
                Remove-Item (\"Env:\" + $_.Name) -ErrorAction SilentlyContinue
            }}
        $env:PYINSTALLER_RESET_ENVIRONMENT = '1'
        Start-Process -FilePath $CurrentExe -WorkingDirectory (Split-Path -Parent $CurrentExe)
    }} catch {{}}
}}
"""

    try:
        script_path.write_text(script, encoding="utf-8-sig")
    except Exception as exc:
        raise UpdateError(f"Não foi possível criar o instalador temporário: {exc}") from exc

    creation_flags = 0
    if sys.platform.startswith("win"):
        creation_flags |= int(getattr(subprocess, "CREATE_NO_WINDOW", 0))
        creation_flags |= int(getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0))

    clean_env = os.environ.copy()
    for env_name in list(clean_env):
        upper_name = env_name.upper()
        if upper_name == "_MEIPASS2" or upper_name.startswith("_PYI_"):
            clean_env.pop(env_name, None)
    clean_env["PYINSTALLER_RESET_ENVIRONMENT"] = "1"

    try:
        subprocess.Popen(
            [
                "powershell.exe",
                "-NoLogo",
                "-NoProfile",
                "-NonInteractive",
                "-ExecutionPolicy",
                "Bypass",
                "-WindowStyle",
                "Hidden",
                "-File",
                str(script_path),
            ],
            cwd=str(executable.parent),
            creationflags=creation_flags,
            close_fds=True,
            env=clean_env,
        )
    except Exception as exc:
        raise UpdateError(f"Não foi possível iniciar o atualizador: {exc}") from exc

    return script_path


def cleanup_previous_executable_backup() -> None:
    """Remove o .bak deixado depois que a nova versão já conseguiu iniciar."""
    if not can_auto_install_update():
        return
    try:
        executable = Path(sys.executable).resolve()
        backup = Path(str(executable) + ".bak")
        backup.unlink(missing_ok=True)
    except Exception:
        pass

def db_path() -> Path:
    return app_dir() / DB_NAME


def config_path() -> Path:
    return app_dir() / CONFIG_NAME


def load_app_config() -> dict:
    """Carrega preferências visuais locais. Se o arquivo não existir ou for inválido, usa o tema escuro."""
    default = {"tema": "dark"}
    path = config_path()
    if not path.exists():
        return default.copy()
    try:
        with path.open("r", encoding="utf-8") as file:
            data = json.load(file)
        if not isinstance(data, dict):
            return default.copy()
        tema = str(data.get("tema", "dark")).strip().lower()
        if tema not in ("dark", "light"):
            tema = "dark"
        return {"tema": tema}
    except (OSError, json.JSONDecodeError, TypeError, ValueError):
        return default.copy()


def save_app_config(config: dict) -> bool:
    """Salva o tema no JSON ao lado do programa sem interferir no banco de dados."""
    path = config_path()
    tema = str(config.get("tema", "dark")).strip().lower()
    if tema not in ("dark", "light"):
        tema = "dark"
    data = {"tema": tema}
    temp_path = path.with_suffix(path.suffix + ".tmp")
    try:
        with temp_path.open("w", encoding="utf-8") as file:
            json.dump(data, file, ensure_ascii=False, indent=4)
            file.write("\n")
        temp_path.replace(path)
        return True
    except OSError:
        try:
            if temp_path.exists():
                temp_path.unlink()
        except OSError:
            pass
        return False


def backup_dir() -> Path:
    folder = app_dir() / "Backups"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def create_database_backup(reason: str = "alteracao") -> Path | None:
    source = db_path()
    if not source.exists():
        return None
    safe_reason = re.sub(r"[^A-Za-z0-9_-]+", "_", normalize_spaces(reason)) or "backup"
    destination = backup_dir() / f"controle_folgas_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{safe_reason}.db"
    create_sqlite_snapshot(source, destination)
    backups = sorted(backup_dir().glob("controle_folgas_*.db"), key=lambda item: item.stat().st_mtime, reverse=True)
    for old in backups[50:]:
        try:
            old.unlink()
        except OSError:
            pass
    return destination


def parse_date(value: str):
    if not value:
        return None
    value = str(value).strip()
    match = re.search(r"(\d{1,2}/\d{1,2}/\d{4})", value)
    if not match:
        return None
    txt = match.group(1)
    try:
        return datetime.strptime(txt, "%d/%m/%Y").date().isoformat()
    except ValueError:
        return None


def br_date(iso_date: str) -> str:
    if not iso_date:
        return ""
    try:
        return datetime.strptime(iso_date, "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return iso_date


def only_text_after_date(value: str) -> str:
    value = str(value or "").replace("\n", " ").strip()
    value = re.sub(r"\s+", " ", value)
    value = re.sub(r"^\d{1,2}/\d{1,2}/\d{4}", "", value).strip(" -|\t")
    return value.strip()


def normalize_spaces(txt: str) -> str:
    return re.sub(r"\s+", " ", str(txt or "").replace("\n", " ")).strip()


def format_program_open_time(started_at: datetime, ended_at: datetime | None = None) -> str:
    """Formata o tempo de execução para leitura rápida no arquivo de status."""
    ended_at = ended_at or datetime.now()
    total_seconds = max(0, int((ended_at - started_at).total_seconds()))
    total_minutes = total_seconds // 60
    days, remainder = divmod(total_minutes, 24 * 60)
    hours, minutes = divmod(remainder, 60)

    parts = []
    if days:
        parts.append(f"{days}d")
    if hours:
        parts.append(f"{hours}h")
    if minutes or not parts:
        parts.append(f"{minutes}min")
    return " ".join(parts)


class ProgramStatusTracker:
    """Mantém um arquivo de status por computador na pasta compartilhada do programa.

    O recurso é intencionalmente silencioso: falhas de gravação não interrompem o
    funcionamento principal do sistema. O heartbeat permite identificar arquivos
    que ficaram como ABERTO após queda de energia ou encerramento forçado.
    """

    def __init__(self, base_dir: Path | None = None, computer_name: str | None = None, now_func=None):
        self.base_dir = Path(base_dir) if base_dir is not None else app_dir()
        self.now_func = now_func or datetime.now
        detected_name = (
            normalize_spaces(computer_name)
            if computer_name
            else normalize_spaces(os.environ.get("COMPUTERNAME") or socket.gethostname() or "COMPUTADOR")
        )
        self.computer_name = detected_name or "COMPUTADOR"
        safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1F]+', "_", self.computer_name).strip(" .") or "COMPUTADOR"
        self.folder = self.base_dir / STATUS_DIR_NAME
        self.path = self.folder / f"{safe_name}.json"
        self.opened_at = None
        self.closed = False
        self._atexit_registered = False

    @staticmethod
    def _format_datetime(value: datetime) -> str:
        return value.strftime("%d/%m/%Y %H:%M:%S")

    def _payload(self, status: str, moment: datetime) -> dict:
        opened_at = self.opened_at or moment
        return {
            "status": status,
            "computador": self.computer_name,
            "versao": APP_VERSION,
            "aberto_em": self._format_datetime(opened_at),
            "ultima_confirmacao": self._format_datetime(moment),
            "fechado_em": self._format_datetime(moment) if status == "FECHADO" else None,
            "tempo_aberto": format_program_open_time(opened_at, moment),
        }

    def _write(self, payload: dict) -> bool:
        """Grava de forma atômica sempre que possível, sem exibir erros ao usuário."""
        temp_path = None
        try:
            self.folder.mkdir(parents=True, exist_ok=True)
            temp_path = self.path.with_name(f".{self.path.name}.{os.getpid()}.tmp")
            with temp_path.open("w", encoding="utf-8") as file:
                json.dump(payload, file, ensure_ascii=False, indent=4)
                file.write("\n")
            temp_path.replace(self.path)
            return True
        except Exception:
            # Em pastas sincronizadas um arquivo pode ficar brevemente bloqueado.
            # Uma tentativa direta evita perder a atualização por esse motivo.
            try:
                with self.path.open("w", encoding="utf-8") as file:
                    json.dump(payload, file, ensure_ascii=False, indent=4)
                    file.write("\n")
                return True
            except Exception:
                return False
        finally:
            if temp_path is not None:
                try:
                    if temp_path.exists():
                        temp_path.unlink()
                except Exception:
                    pass

    def start(self) -> bool:
        self.opened_at = self.now_func()
        self.closed = False
        written = self._write(self._payload("ABERTO", self.opened_at))
        if not self._atexit_registered:
            try:
                atexit.register(self.close)
                self._atexit_registered = True
            except Exception:
                pass
        return written

    def heartbeat(self) -> bool:
        if self.closed or self.opened_at is None:
            return False
        moment = self.now_func()
        return self._write(self._payload("ABERTO", moment))

    def close(self) -> bool:
        if self.closed or self.opened_at is None:
            return False
        moment = self.now_func()
        written = self._write(self._payload("FECHADO", moment))
        self.closed = True
        return written


def normalize_recovery_code(value: str) -> str:
    """Remove separadores e padroniza o código de recuperação."""
    return re.sub(r"[^A-Za-z0-9]", "", str(value or "")).upper()


def generate_recovery_code() -> str:
    """Gera um código legível, evitando caracteres visualmente ambíguos."""
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
    raw = "".join(secrets.choice(alphabet) for _ in range(20))
    return "-".join(raw[i:i + 4] for i in range(0, len(raw), 4))


def hash_secret(secret: str) -> str:
    """Cria hash PBKDF2-SHA256 com salt aleatório."""
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac(
        "sha256",
        str(secret).encode("utf-8"),
        salt,
        PBKDF2_ITERATIONS,
    )
    return f"pbkdf2_sha256${PBKDF2_ITERATIONS}${salt.hex()}${digest.hex()}"


def verify_secret(secret: str, stored_hash: str) -> bool:
    """Compara uma senha/código com o hash armazenado sem expor o valor original."""
    try:
        algorithm, iterations, salt_hex, digest_hex = str(stored_hash).split("$", 3)
        if algorithm != "pbkdf2_sha256":
            return False
        digest = hashlib.pbkdf2_hmac(
            "sha256",
            str(secret).encode("utf-8"),
            bytes.fromhex(salt_hex),
            int(iterations),
        )
        return hmac.compare_digest(digest.hex(), digest_hex)
    except Exception:
        return False


def month_value(label: str) -> str:
    """Retorna o número do mês no formato MM ou 'Todos'."""
    label = str(label or "Todos").strip()
    if not label or label == "Todos":
        return "Todos"
    match = re.match(r"^(\d{1,2})", label)
    if match:
        return match.group(1).zfill(2)
    return "Todos"


def month_name_from_value(value: str) -> str:
    value = month_value(value)
    if value == "Todos":
        return "Todos"
    return MESES_NOMES.get(value, value)


def safe_filename(texto: str) -> str:
    texto = normalize_spaces(texto)
    texto = re.sub(r'[<>:"/\\|?*]', "", texto)
    texto = re.sub(r"\s+", "_", texto)
    return texto.strip("_") or "Controle_Folgas"


def safe_sheet_name(texto: str) -> str:
    texto = normalize_spaces(texto)
    texto = re.sub(r"[\[\]\*\?/\\:]", "", texto)
    return (texto or "Planilha")[:31]


def period_label(year: str = "Todos", month: str = "Todos") -> str:
    mes = month_value(month)
    ano = str(year or "Todos")
    if mes != "Todos" and ano != "Todos":
        return f"{MESES_NOMES.get(mes, mes)} de {ano}"
    if ano != "Todos":
        return f"Ano {ano}"
    if mes != "Todos":
        return MESES_NOMES.get(mes, mes)
    return "Todos os períodos"


def period_end_iso(year: str, month: str):
    mes = month_value(month)
    if not year or year == "Todos" or mes == "Todos":
        return None
    try:
        ano = int(year)
        m = int(mes)
        last_day = monthrange(ano, m)[1]
        return f"{ano:04d}-{m:02d}-{last_day:02d}"
    except Exception:
        return None


def format_minutes(total_minutes: int, hours_per_day: int) -> str:
    sign = "-" if total_minutes < 0 else ""
    total = abs(int(total_minutes))
    day_minutes = hours_per_day * 60
    days = total // day_minutes
    rest = total % day_minutes
    hours = rest // 60
    minutes = rest % 60

    parts = []
    if days:
        parts.append(f"{days} dia" + ("s" if days != 1 else ""))
    if hours:
        parts.append(f"{hours} hora" + ("s" if hours != 1 else ""))
    if minutes:
        parts.append(f"{minutes} minuto" + ("s" if minutes != 1 else ""))
    if not parts:
        return "0"
    return sign + " e ".join(parts)


def guess_movimento(categoria: str, lado_doc: str = "credito") -> str:
    c = normalize_spaces(categoria).lower()
    if lado_doc == "debito":
        return "Débito"
    for item in CATEGORIAS_DEBITO:
        if item in c:
            return "Débito"
    for item in CATEGORIAS_CREDITO:
        if item in c:
            return "Crédito"
    return "Crédito"


class PortugueseDateEntry(ctk.CTkFrame):
    """Campo de data com calendário próprio em português e compatível com PyInstaller."""

    MONTH_NAMES = [
        "",
        "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
        "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
    ]
    WEEKDAYS = ["Dom", "Seg", "Ter", "Qua", "Qui", "Sex", "Sáb"]

    def __init__(self, master, textvariable=None, height=30, corner_radius=10, **kwargs):
        super().__init__(master, fg_color="transparent", **kwargs)
        self.variable = textvariable or tk.StringVar()
        self.popup = None
        self.view_year = datetime.now().year
        self.view_month = datetime.now().month

        self.grid_columnconfigure(0, weight=1)
        self.entry = ctk.CTkEntry(
            self,
            textvariable=self.variable,
            height=height,
            corner_radius=corner_radius,
        )
        self.entry.grid(row=0, column=0, sticky="ew")

        self.button = ctk.CTkButton(
            self,
            text="📅",
            width=38,
            height=height,
            corner_radius=corner_radius,
            command=self.open_calendar,
        )
        self.button.grid(row=0, column=1, padx=(5, 0))

    def get(self):
        return self.variable.get()

    def set(self, value):
        self.variable.set(value)

    def open_calendar(self):
        if self.popup is not None and self.popup.winfo_exists():
            self.popup.lift()
            self.popup.focus_force()
            return

        try:
            selected = datetime.strptime(self.variable.get().strip(), "%d/%m/%Y")
        except Exception:
            selected = datetime.now()

        self.view_year = selected.year
        self.view_month = selected.month

        self.popup = tk.Toplevel(self)
        self.popup.title("Selecionar data")
        apply_window_icon(self.popup)
        self.popup.resizable(False, False)
        self.popup.transient(self.winfo_toplevel())
        self.popup.grab_set()
        self.popup.protocol("WM_DELETE_WINDOW", self.close_calendar)

        self.calendar_frame = ttk.Frame(self.popup, padding=8)
        self.calendar_frame.pack(fill="both", expand=True)
        self.draw_calendar()

        self.popup.update_idletasks()
        x = self.winfo_rootx()
        y = self.winfo_rooty() + self.winfo_height() + 4
        pw = self.popup.winfo_reqwidth()
        ph = self.popup.winfo_reqheight()

        # O calendário é um Toplevel Tk puro, portanto suas medidas já estão em
        # pixels físicos. Limitamos também à área útil para não cobrir a barra
        # de tarefas em notebooks ou em escalas altas.
        top = self.winfo_toplevel()
        try:
            if hasattr(top, "get_work_area"):
                work_w, work_h, work_x, work_y = top.get_work_area()
            else:
                work_w, work_h = self.popup.winfo_screenwidth(), self.popup.winfo_screenheight()
                work_x, work_y = 0, 0
        except Exception:
            work_w, work_h = self.popup.winfo_screenwidth(), self.popup.winfo_screenheight()
            work_x, work_y = 0, 0

        x = min(max(work_x, x), max(work_x, work_x + work_w - pw))
        y = min(max(work_y, y), max(work_y, work_y + work_h - ph))
        self.popup.geometry(f"+{x}+{y}")

    def close_calendar(self):
        if self.popup is not None:
            try:
                self.popup.grab_release()
            except Exception:
                pass
            try:
                self.popup.destroy()
            except Exception:
                pass
        self.popup = None

    def change_month(self, delta):
        new_month = self.view_month + delta
        if new_month < 1:
            self.view_month = 12
            self.view_year -= 1
        elif new_month > 12:
            self.view_month = 1
            self.view_year += 1
        else:
            self.view_month = new_month
        self.draw_calendar()

    def draw_calendar(self):
        for widget in self.calendar_frame.winfo_children():
            widget.destroy()

        nav = ttk.Frame(self.calendar_frame)
        nav.grid(row=0, column=0, columnspan=7, sticky="ew", pady=(0, 6))
        nav.grid_columnconfigure(1, weight=1)

        ttk.Button(nav, text="◀", width=4, command=lambda: self.change_month(-1)).grid(row=0, column=0)
        ttk.Label(
            nav,
            text=f"{self.MONTH_NAMES[self.view_month]} de {self.view_year}",
            anchor="center",
            font=("Segoe UI", 10, "bold"),
        ).grid(row=0, column=1, padx=10, sticky="ew")
        ttk.Button(nav, text="▶", width=4, command=lambda: self.change_month(1)).grid(row=0, column=2)

        for col, name in enumerate(self.WEEKDAYS):
            ttk.Label(
                self.calendar_frame,
                text=name,
                anchor="center",
                width=4,
                font=("Segoe UI", 9, "bold"),
            ).grid(row=1, column=col, padx=1, pady=1)

        first_weekday = (datetime(self.view_year, self.view_month, 1).weekday() + 1) % 7
        total_days = monthrange(self.view_year, self.view_month)[1]
        row = 2
        col = first_weekday

        for day in range(1, total_days + 1):
            ttk.Button(
                self.calendar_frame,
                text=str(day),
                width=4,
                command=lambda d=day: self.select_date(d),
            ).grid(row=row, column=col, padx=1, pady=1)
            col += 1
            if col > 6:
                col = 0
                row += 1

        ttk.Button(
            self.calendar_frame,
            text="Hoje",
            command=self.select_today,
        ).grid(row=row + 1, column=0, columnspan=7, sticky="ew", pady=(7, 0))

    def select_date(self, day):
        self.variable.set(f"{day:02d}/{self.view_month:02d}/{self.view_year:04d}")
        self.close_calendar()

    def select_today(self):
        today = datetime.now()
        self.variable.set(today.strftime("%d/%m/%Y"))
        self.close_calendar()



class AutoHideScrollbar(ttk.Scrollbar):
    """Scrollbar que só aparece quando realmente existe conteúdo fora da área visível."""

    def set(self, low, high):
        try:
            low_f = float(low)
            high_f = float(high)
            if low_f <= 0.0 and high_f >= 1.0:
                self.grid_remove()
            else:
                self.grid()
        except Exception:
            pass
        super().set(low, high)

class Database:
    def __init__(self):
        self.backup_before_security_migration()
        self.conn = sqlite3.connect(db_path())
        self.conn.row_factory = sqlite3.Row
        self.create_tables()

    def backup_before_security_migration(self):
        """Cria uma cópia do banco antigo antes de acrescentar Lixeira e auditoria."""
        source = db_path()
        if not source.exists():
            return
        try:
            check_conn = sqlite3.connect(source)
            tables = {
                row[0]
                for row in check_conn.execute(
                    "SELECT name FROM sqlite_master WHERE type='table'"
                ).fetchall()
            }
            if "lancamentos" not in tables:
                check_conn.close()
                return
            columns = {
                row[1]
                for row in check_conn.execute("PRAGMA table_info(lancamentos)").fetchall()
            }
            check_conn.close()
            required_columns = {"excluido", "excluido_em", "motivo_exclusao"}
            if not required_columns.issubset(columns) or "auditoria" not in tables:
                create_database_backup("antes_migracao_seguranca")
        except Exception:
            # A falha em verificar a migração não impede a abertura do sistema.
            try:
                check_conn.close()
            except Exception:
                pass

    def create_tables(self):
        cur = self.conn.cursor()
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS config (
                chave TEXT PRIMARY KEY,
                valor TEXT NOT NULL
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS funcionarios (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nome TEXT NOT NULL UNIQUE,
                ativo INTEGER NOT NULL DEFAULT 1,
                criado_em TEXT NOT NULL,
                jornada_horas INTEGER NOT NULL DEFAULT 12,
                regime TEXT NOT NULL DEFAULT '12x36'
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS lancamentos (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                funcionario_id INTEGER NOT NULL,
                data TEXT NOT NULL,
                movimento TEXT NOT NULL,
                categoria TEXT NOT NULL,
                quantidade_texto TEXT,
                minutos INTEGER NOT NULL,
                observacao TEXT,
                origem TEXT DEFAULT 'Manual',
                criado_em TEXT NOT NULL,
                excluido INTEGER NOT NULL DEFAULT 0,
                excluido_em TEXT,
                motivo_exclusao TEXT,
                FOREIGN KEY(funcionario_id) REFERENCES funcionarios(id)
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS auditoria (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                acao TEXT NOT NULL,
                lancamento_id INTEGER,
                detalhes TEXT,
                criado_em TEXT NOT NULL
            )
            """
        )

        # Migração automática para bancos criados por versões anteriores.
        funcionario_columns = {row[1] for row in cur.execute("PRAGMA table_info(funcionarios)").fetchall()}
        if "jornada_horas" not in funcionario_columns:
            cur.execute("ALTER TABLE funcionarios ADD COLUMN jornada_horas INTEGER NOT NULL DEFAULT 12")
        if "regime" not in funcionario_columns:
            cur.execute("ALTER TABLE funcionarios ADD COLUMN regime TEXT NOT NULL DEFAULT '12x36'")

        lancamento_columns = {row[1] for row in cur.execute("PRAGMA table_info(lancamentos)").fetchall()}
        if "excluido" not in lancamento_columns:
            cur.execute("ALTER TABLE lancamentos ADD COLUMN excluido INTEGER NOT NULL DEFAULT 0")
        if "excluido_em" not in lancamento_columns:
            cur.execute("ALTER TABLE lancamentos ADD COLUMN excluido_em TEXT")
        if "motivo_exclusao" not in lancamento_columns:
            cur.execute("ALTER TABLE lancamentos ADD COLUMN motivo_exclusao TEXT")

        cur.execute("UPDATE funcionarios SET jornada_horas=12 WHERE jornada_horas IS NULL OR jornada_horas NOT IN (6,8,12)")
        cur.execute("UPDATE funcionarios SET regime='12x36' WHERE regime IS NULL OR trim(regime)=''")
        cur.execute("UPDATE lancamentos SET excluido=0 WHERE excluido IS NULL")
        cur.execute("INSERT OR IGNORE INTO config (chave, valor) VALUES ('horas_por_dia', ?)", (str(DEFAULT_HOURS_PER_DAY),))
        cur.execute("CREATE INDEX IF NOT EXISTS idx_lancamentos_ativos ON lancamentos(excluido, funcionario_id, data)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_auditoria_data ON auditoria(criado_em DESC)")
        self.conn.commit()

    def get_config(self, chave: str, default=None):
        row = self.conn.execute("SELECT valor FROM config WHERE chave=?", (chave,)).fetchone()
        return row["valor"] if row else default

    def set_config(self, chave: str, valor):
        self.conn.execute(
            "INSERT OR REPLACE INTO config (chave, valor) VALUES (?, ?)",
            (chave, str(valor)),
        )
        self.conn.commit()

    def delete_config(self, chave: str):
        self.conn.execute("DELETE FROM config WHERE chave=?", (chave,))
        self.conn.commit()

    def get_hours_per_day(self) -> int:
        try:
            return int(self.get_config("horas_por_dia", DEFAULT_HOURS_PER_DAY))
        except Exception:
            return DEFAULT_HOURS_PER_DAY

    def set_hours_per_day(self, value: int):
        self.set_config("horas_por_dia", int(value))

    def get_service_name(self) -> str:
        """Retorna o nome do serviço cadastrado para identificação dos relatórios."""
        return normalize_spaces(self.get_config("nome_servico", ""))

    def set_service_name(self, value: str):
        """Salva o nome do serviço na tabela de configurações do banco atual."""
        value = normalize_spaces(value)
        if not value:
            raise ValueError("Informe o nome do serviço.")
        self.set_config("nome_servico", value)

    # ------------------------------------------------------------------
    # Segurança administrativa
    # ------------------------------------------------------------------
    def admin_password_exists(self) -> bool:
        return bool(self.get_config("admin_password_hash"))

    def set_admin_credentials(self, password: str) -> str:
        recovery_code = generate_recovery_code()
        self.conn.execute(
            "INSERT OR REPLACE INTO config (chave, valor) VALUES ('admin_password_hash', ?)",
            (hash_secret(password),),
        )
        self.conn.execute(
            "INSERT OR REPLACE INTO config (chave, valor) VALUES ('admin_recovery_hash', ?)",
            (hash_secret(normalize_recovery_code(recovery_code)),),
        )
        self.conn.execute(
            "INSERT OR REPLACE INTO config (chave, valor) VALUES ('admin_failed_attempts', '0')"
        )
        self.conn.execute("DELETE FROM config WHERE chave='admin_locked_until'")
        self.conn.commit()
        return recovery_code

    def verify_admin_password(self, password: str) -> bool:
        stored = self.get_config("admin_password_hash", "")
        return bool(stored) and verify_secret(password, stored)

    def verify_recovery_code(self, recovery_code: str) -> bool:
        stored = self.get_config("admin_recovery_hash", "")
        normalized = normalize_recovery_code(recovery_code)
        return bool(stored and normalized) and verify_secret(normalized, stored)

    def admin_lock_remaining_seconds(self) -> int:
        value = self.get_config("admin_locked_until")
        if not value:
            return 0
        try:
            locked_until = datetime.fromisoformat(value)
            remaining = int((locked_until - datetime.now()).total_seconds())
            if remaining <= 0:
                self.reset_admin_attempts()
                return 0
            return remaining
        except Exception:
            self.reset_admin_attempts()
            return 0

    def register_failed_admin_attempt(self) -> tuple[int, bool]:
        try:
            attempts = int(self.get_config("admin_failed_attempts", "0")) + 1
        except Exception:
            attempts = 1

        locked = attempts >= ADMIN_MAX_ATTEMPTS
        if locked:
            locked_until = datetime.now() + timedelta(minutes=ADMIN_LOCK_MINUTES)
            self.conn.execute(
                "INSERT OR REPLACE INTO config (chave, valor) VALUES ('admin_locked_until', ?)",
                (locked_until.isoformat(timespec="seconds"),),
            )
            attempts = 0

        self.conn.execute(
            "INSERT OR REPLACE INTO config (chave, valor) VALUES ('admin_failed_attempts', ?)",
            (str(attempts),),
        )
        self.conn.commit()
        return attempts, locked

    def reset_admin_attempts(self):
        self.conn.execute(
            "INSERT OR REPLACE INTO config (chave, valor) VALUES ('admin_failed_attempts', '0')"
        )
        self.conn.execute("DELETE FROM config WHERE chave='admin_locked_until'")
        self.conn.commit()

    # ------------------------------------------------------------------
    # Funcionários e lançamentos
    # ------------------------------------------------------------------
    def add_employee(self, nome: str, jornada_horas: int = 12, regime: str = "12x36"):
        nome = normalize_spaces(nome)
        if not nome:
            raise ValueError("Informe o nome do funcionário.")
        jornada_horas = int(jornada_horas)
        if jornada_horas not in (6, 8, 12):
            raise ValueError("A jornada deve ser de 6, 8 ou 12 horas.")
        self.conn.execute(
            "INSERT INTO funcionarios (nome, ativo, criado_em, jornada_horas, regime) VALUES (?, 1, ?, ?, ?)",
            (nome, datetime.now().isoformat(timespec="seconds"), jornada_horas, normalize_spaces(regime)),
        )
        self.conn.commit()

    def update_employee(self, employee_id: int, nome: str, ativo: int, jornada_horas: int = 12, regime: str = "12x36"):
        old = self.conn.execute("SELECT * FROM funcionarios WHERE id=?", (employee_id,)).fetchone()
        if not old:
            raise ValueError("O funcionário não foi encontrado.")

        nome = normalize_spaces(nome)
        regime = normalize_spaces(regime)
        jornada_horas = int(jornada_horas)
        ativo = int(ativo)
        if not nome:
            raise ValueError("Informe o nome do funcionário.")
        if jornada_horas not in (6, 8, 12):
            raise ValueError("A jornada deve ser de 6, 8 ou 12 horas.")

        changes = []
        if normalize_spaces(old["nome"]) != nome:
            changes.append(f"Nome: {old['nome']} → {nome}")
        if int(old["ativo"] or 0) != ativo:
            changes.append(f"Ativo: {'Sim' if int(old['ativo'] or 0) else 'Não'} → {'Sim' if ativo else 'Não'}")
        if int(old["jornada_horas"] or DEFAULT_HOURS_PER_DAY) != jornada_horas:
            changes.append(f"Jornada: {old['jornada_horas']}h → {jornada_horas}h")
        if normalize_spaces(old["regime"]) != regime:
            changes.append(f"Regime: {old['regime']} → {regime}")

        if not changes:
            return False

        self.conn.execute(
            "UPDATE funcionarios SET nome=?, ativo=?, jornada_horas=?, regime=? WHERE id=?",
            (nome, ativo, jornada_horas, regime, employee_id),
        )
        self._insert_audit(
            "FUNCIONÁRIO ALTERADO",
            detalhes=f"Funcionário ID: {employee_id} | " + " | ".join(changes),
        )
        self.conn.commit()
        return True

    def employees(self, only_active=False):
        if only_active:
            return self.conn.execute("SELECT * FROM funcionarios WHERE ativo=1 ORDER BY nome").fetchall()
        return self.conn.execute("SELECT * FROM funcionarios ORDER BY nome").fetchall()

    def employee_by_name(self, nome: str):
        return self.conn.execute("SELECT * FROM funcionarios WHERE lower(nome)=lower(?)", (normalize_spaces(nome),)).fetchone()

    def employee_by_id(self, employee_id: int):
        return self.conn.execute("SELECT * FROM funcionarios WHERE id=?", (int(employee_id),)).fetchone()

    def employee_hours(self, employee_id: int) -> int:
        row = self.conn.execute("SELECT jornada_horas FROM funcionarios WHERE id=?", (employee_id,)).fetchone()
        return int(row["jornada_horas"] if row and row["jornada_horas"] else DEFAULT_HOURS_PER_DAY)

    def add_entry(self, funcionario_id, data, movimento, categoria, quantidade_texto, minutos, observacao, origem="Manual"):
        signed = abs(int(minutos)) if movimento == "Crédito" else -abs(int(minutos))
        self.conn.execute(
            """
            INSERT INTO lancamentos
            (funcionario_id, data, movimento, categoria, quantidade_texto, minutos, observacao, origem, criado_em, excluido)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
            """,
            (
                funcionario_id,
                data,
                movimento,
                categoria,
                quantidade_texto,
                signed,
                observacao,
                origem,
                datetime.now().isoformat(timespec="seconds"),
            ),
        )
        self.conn.commit()

    def update_entry(self, entry_id, funcionario_id, data, movimento, categoria, quantidade_texto, minutos, observacao):
        old = self.entry_by_id(entry_id, include_deleted=False)
        if not old:
            raise ValueError("O lançamento não foi encontrado ou não está disponível para alteração.")

        employee = self.conn.execute(
            "SELECT id, nome, jornada_horas FROM funcionarios WHERE id=?",
            (int(funcionario_id),),
        ).fetchone()
        if not employee:
            raise ValueError("O funcionário selecionado não foi encontrado.")

        movimento = normalizar_movimento(movimento)
        categoria = normalize_spaces(categoria)
        quantidade_texto = normalize_spaces(quantidade_texto)
        observacao = str(observacao or "").strip()
        signed = abs(int(minutos)) if movimento == MOVIMENTO_CREDITO else -abs(int(minutos))

        changes = []
        if int(old["funcionario_id"]) != int(funcionario_id):
            changes.append(f"Funcionário: {old['funcionario']} → {employee['nome']}")
        if old["data"] != data:
            changes.append(f"Data: {br_date(old['data'])} → {br_date(data)}")
        if normalize_spaces(old["movimento"]) != movimento:
            changes.append(f"Movimento: {old['movimento']} → {movimento}")
        if normalize_spaces(old["categoria"]) != categoria:
            changes.append(f"Categoria: {old['categoria']} → {categoria}")

        old_qtd = normalize_spaces(old["quantidade_texto"] or "")
        old_minutes = int(old["minutos"] or 0)
        if old_qtd != quantidade_texto or abs(old_minutes) != abs(signed):
            old_hpd = int(old["jornada_horas"] or DEFAULT_HOURS_PER_DAY)
            new_hpd = int(employee["jornada_horas"] or DEFAULT_HOURS_PER_DAY)
            old_display = old_qtd or format_minutes(abs(old_minutes), old_hpd)
            new_display = quantidade_texto or format_minutes(abs(signed), new_hpd)
            if old_display == new_display and abs(old_minutes) != abs(signed):
                old_display = f"{old_display} ({abs(old_minutes)} min)"
                new_display = f"{new_display} ({abs(signed)} min)"
            changes.append(f"Quantidade: {old_display} → {new_display}")

        old_obs = normalize_spaces(old["observacao"] or "")
        new_obs_audit = normalize_spaces(observacao)
        if old_obs != new_obs_audit:
            changes.append(f"Observação: {old_obs or '(vazia)'} → {new_obs_audit or '(vazia)'}")

        if not changes:
            return False

        self.conn.execute(
            """
            UPDATE lancamentos
            SET funcionario_id=?, data=?, movimento=?, categoria=?, quantidade_texto=?, minutos=?, observacao=?
            WHERE id=? AND excluido=0
            """,
            (int(funcionario_id), data, movimento, categoria, quantidade_texto, signed, observacao, entry_id),
        )
        self._insert_audit(
            "LANÇAMENTO ALTERADO",
            entry_id,
            "Alterações: " + " | ".join(changes),
        )
        self.conn.commit()
        return True

    def entry_by_id(self, entry_id: int, include_deleted: bool = True):
        sql = """
            SELECT l.*, f.nome AS funcionario, f.jornada_horas, f.regime
            FROM lancamentos l
            JOIN funcionarios f ON f.id=l.funcionario_id
            WHERE l.id=?
        """
        params = [entry_id]
        if not include_deleted:
            sql += " AND l.excluido=0"
        return self.conn.execute(sql, params).fetchone()

    def entries(self, funcionario_id=None, year=None, month=None, search=""):
        params = []
        where = ["l.excluido=0"]
        if funcionario_id:
            where.append("l.funcionario_id=?")
            params.append(funcionario_id)
        if year and year != "Todos":
            where.append("substr(l.data, 1, 4)=?")
            params.append(str(year))
        month = month_value(month)
        if month and month != "Todos":
            where.append("substr(l.data, 6, 2)=?")
            params.append(month)
        if search:
            where.append("(lower(l.categoria) LIKE ? OR lower(COALESCE(l.observacao,'')) LIKE ?)")
            s = f"%{search.lower()}%"
            params.extend([s, s])
        sql = """
            SELECT l.*, f.nome AS funcionario, f.jornada_horas, f.regime
            FROM lancamentos l
            JOIN funcionarios f ON f.id = l.funcionario_id
            WHERE """ + " AND ".join(where)
        sql += " ORDER BY l.data DESC, l.id DESC"
        return self.conn.execute(sql, params).fetchall()

    def trash_entries(self):
        return self.conn.execute(
            """
            SELECT l.*, f.nome AS funcionario, f.jornada_horas, f.regime
            FROM lancamentos l
            JOIN funcionarios f ON f.id=l.funcionario_id
            WHERE l.excluido=1
            ORDER BY COALESCE(l.excluido_em, l.criado_em) DESC, l.id DESC
            """
        ).fetchall()

    def soft_delete_entry(self, entry_id: int, motivo: str):
        row = self.entry_by_id(entry_id, include_deleted=False)
        if not row:
            raise ValueError("O lançamento não foi encontrado ou já está na Lixeira.")
        motivo = normalize_spaces(motivo)
        if not motivo:
            raise ValueError("Informe o motivo da exclusão.")
        now = datetime.now().isoformat(timespec="seconds")
        self.conn.execute(
            "UPDATE lancamentos SET excluido=1, excluido_em=?, motivo_exclusao=? WHERE id=? AND excluido=0",
            (now, motivo, entry_id),
        )
        details = (
            f"Funcionário: {row['funcionario']} | Data: {br_date(row['data'])} | "
            f"Categoria: {row['categoria']} | Movimento: {row['movimento']} | "
            f"Quantidade: {row['quantidade_texto'] or ''} | Motivo: {motivo}"
        )
        self._insert_audit("LANÇAMENTO MOVIDO PARA A LIXEIRA", entry_id, details)
        self.conn.commit()

    def restore_entry(self, entry_id: int):
        row = self.entry_by_id(entry_id, include_deleted=True)
        if not row or int(row["excluido"] or 0) != 1:
            raise ValueError("O lançamento não foi encontrado na Lixeira.")
        previous_reason = row["motivo_exclusao"] or ""
        self.conn.execute(
            "UPDATE lancamentos SET excluido=0, excluido_em=NULL, motivo_exclusao=NULL WHERE id=? AND excluido=1",
            (entry_id,),
        )
        details = (
            f"Funcionário: {row['funcionario']} | Data: {br_date(row['data'])} | "
            f"Categoria: {row['categoria']} | Motivo anterior: {previous_reason}"
        )
        self._insert_audit("LANÇAMENTO RESTAURADO", entry_id, details)
        self.conn.commit()

    def permanent_delete_entry(self, entry_id: int):
        row = self.entry_by_id(entry_id, include_deleted=True)
        if not row or int(row["excluido"] or 0) != 1:
            raise ValueError("O lançamento não foi encontrado na Lixeira.")
        details = (
            f"Funcionário: {row['funcionario']} | Data: {br_date(row['data'])} | "
            f"Categoria: {row['categoria']} | Motivo da Lixeira: {row['motivo_exclusao'] or ''}"
        )
        self._insert_audit("LANÇAMENTO EXCLUÍDO DEFINITIVAMENTE", entry_id, details)
        self.conn.execute("DELETE FROM lancamentos WHERE id=? AND excluido=1", (entry_id,))
        self.conn.commit()

    def empty_trash(self) -> int:
        rows = self.trash_entries()
        for row in rows:
            details = (
                f"Funcionário: {row['funcionario']} | Data: {br_date(row['data'])} | "
                f"Categoria: {row['categoria']} | Motivo da Lixeira: {row['motivo_exclusao'] or ''}"
            )
            self._insert_audit("LANÇAMENTO EXCLUÍDO AO ESVAZIAR LIXEIRA", row["id"], details)
        self.conn.execute("DELETE FROM lancamentos WHERE excluido=1")
        self.conn.commit()
        return len(rows)

    def saldo(self, funcionario_id=None):
        if funcionario_id:
            row = self.conn.execute(
                "SELECT COALESCE(SUM(minutos),0) AS total FROM lancamentos WHERE funcionario_id=? AND excluido=0",
                (funcionario_id,),
            ).fetchone()
        else:
            row = self.conn.execute(
                "SELECT COALESCE(SUM(minutos),0) AS total FROM lancamentos WHERE excluido=0"
            ).fetchone()
        return int(row["total"] or 0)

    def saldo_ate(self, funcionario_id=None, data_limite=None):
        params = []
        where = ["excluido=0"]
        if funcionario_id:
            where.append("funcionario_id=?")
            params.append(funcionario_id)
        if data_limite:
            where.append("data<=?")
            params.append(data_limite)
        sql = "SELECT COALESCE(SUM(minutos),0) AS total FROM lancamentos WHERE " + " AND ".join(where)
        row = self.conn.execute(sql, params).fetchone()
        return int(row["total"] or 0)

    def totals(self):
        return self.conn.execute(
            """
            SELECT
                COUNT(*) AS lancamentos,
                COALESCE(SUM(CASE WHEN minutos > 0 THEN minutos ELSE 0 END),0) AS creditos,
                COALESCE(SUM(CASE WHEN minutos < 0 THEN minutos ELSE 0 END),0) AS debitos,
                COALESCE(SUM(minutos),0) AS saldo
            FROM lancamentos
            WHERE excluido=0
            """
        ).fetchone()

    def years(self):
        rows = self.conn.execute(
            "SELECT DISTINCT substr(data,1,4) AS ano FROM lancamentos WHERE excluido=0 ORDER BY ano DESC"
        ).fetchall()
        return [r["ano"] for r in rows if r["ano"]]

    # ------------------------------------------------------------------
    # Auditoria
    # ------------------------------------------------------------------
    def _insert_audit(self, acao: str, lancamento_id=None, detalhes: str = ""):
        self.conn.execute(
            "INSERT INTO auditoria (acao, lancamento_id, detalhes, criado_em) VALUES (?, ?, ?, ?)",
            (
                normalize_spaces(acao),
                lancamento_id,
                normalize_spaces(detalhes),
                datetime.now().isoformat(timespec="seconds"),
            ),
        )

    def record_audit(self, acao: str, lancamento_id=None, detalhes: str = ""):
        self._insert_audit(acao, lancamento_id, detalhes)
        self.conn.commit()

    def audit_entries(self, limit: int = 500):
        return self.conn.execute(
            "SELECT * FROM auditoria ORDER BY criado_em DESC, id DESC LIMIT ?",
            (int(limit),),
        ).fetchall()


def enable_dpi_awareness():
    """Compatibilidade mantida: o CustomTkinter já gerencia HighDPI automaticamente.

    Não chamamos APIs de DPI manualmente aqui porque isso pode impedir que o
    CustomTkinter ative o modo Per-Monitor e gerar cálculos inconsistentes entre
    pixels físicos do Windows e as dimensões lógicas usadas por CTk.
    """
    return


class App(ctk.CTk):
    def __init__(self):
        # O próprio CustomTkinter ativa e acompanha o DPI do monitor.
        # É importante inicializá-lo antes de consultar a área útil do Windows.
        super().__init__()
        # Mantém a janela principal escondida enquanto o layout é construído.
        # O splash permanece visível e evita o efeito de uma janela incompleta.
        try:
            self.withdraw()
        except Exception:
            pass

        # Assume o carregamento assim que o Python iniciou. Dessa forma, mesmo
        # que o splash experimental do PyInstaller tenha alguma incompatibilidade
        # visual com Tcl/Tk da máquina, o usuário recebe uma tela de espera real.
        create_app_startup_window(self)
        startup_splash_update("Verificando atualização anterior...")

        # Se esta execução já é a nova versão iniciada pelo atualizador,
        # remove o backup temporário do executável anterior.
        cleanup_previous_executable_backup()

        startup_splash_update("Abrindo banco de dados...")
        self.db = Database()
        self.selected_entry_id = None
        self.selected_employee_id = None
        self._admin_credentials_reset_during_auth = False

        # Controle silencioso de uso por computador. Não aparece na interface.
        self._status_tracker = ProgramStatusTracker()
        self._status_after_id = None
        self._status_closing = False
        self._status_closed = False

        # Estado do atualizador. Toda comunicação entre as threads de rede e o
        # Tk ocorre pela fila para não tocar na interface fora da thread principal.
        self._update_queue = queue.Queue()
        self._update_check_running = False
        self._update_download_running = False
        self._update_poll_after_id = None
        self._update_progress_window = None
        self._update_progress_bar = None
        self._update_progress_label = None

        self.title(APP_NAME)
        apply_window_icon(self)
        self._ctk_image_cache = {}
        self.screen_w = self.winfo_screenwidth()
        self.screen_h = self.winfo_screenheight()
        self.work_w, self.work_h, self.work_x, self.work_y = self.get_work_area()

        # CustomTkinter trabalha com dimensões lógicas para geometry()/minsize(),
        # mas a API do Windows retorna a área útil em pixels físicos. Converter
        # aqui evita que 125%/150% aumente a janela uma segunda vez.
        try:
            self.window_scale = max(0.5, float(self._get_window_scaling()))
        except Exception:
            self.window_scale = 1.0
        self.work_logical_w = max(1, int(self.work_w / self.window_scale))
        self.work_logical_h = max(1, int(self.work_h / self.window_scale))

        # Usar layout compacto como padrão deixa o sistema seguro em notebooks
        # com escala alta e monitores menores.
        self.compact_mode = True
        self.form_width = 320 if self.work_logical_w <= 1400 else 360
        self.header_title_font = 20 if self.compact_mode else 24
        self.header_sub_font = 11 if self.compact_mode else 13
        self.body_font = 9 if self.compact_mode else 10
        self.card_font = 18 if self.compact_mode else 24

        # Começa ocupando a área útil em unidades lógicas; depois tentamos
        # maximizar de verdade pelo Windows. O tamanho mínimo também é limitado
        # ao que realmente cabe no monitor atual.
        self.geometry(f"{self.work_logical_w}x{self.work_logical_h}+{self.work_x}+{self.work_y}")
        self.minsize(
            max(420, min(980, self.work_logical_w - 12)),
            max(360, min(620, self.work_logical_h - 12)),
        )

        # Recupera o último tema escolhido pelo usuário. O JSON fica ao lado do programa.
        self.app_config = load_app_config()
        ctk.set_appearance_mode(self.app_config.get("tema", "dark"))
        ctk.set_default_color_theme("blue")
        self.configure(fg_color=("#f3f4f6", "#0f172a"))

        # Cria o JSON na primeira execução, preservando o tema padrão escuro.
        save_app_config(self.app_config)

        startup_splash_update("Preparando aparência...")
        self.setup_style()
        startup_splash_update("Montando telas do sistema...")
        self.build_ui()
        self.apply_treeview_theme()
        startup_splash_update("Carregando dados...")
        self.refresh_all()

        # Registra que este computador está usando o programa e inicia o heartbeat.
        self._start_program_status_tracking()

        # Em bancos já existentes, o cadastro aparece somente uma vez após a atualização.
        # Nas próximas aberturas, o nome salvo é reutilizado automaticamente.
        self.after(250, self.ensure_service_name)

        # A verificação ocorre em segundo plano para não atrasar a abertura.
        # Falhas de internet ficam silenciosas na verificação automática.
        self.after(1400, lambda: self.check_for_updates(silent=True))

        self.bind("<Configure>", self.on_window_resize)
        self.protocol("WM_DELETE_WINDOW", self.destroy)

        # Várias tentativas porque alguns Windows/Tk ignoram a primeira maximização
        # durante a criação da janela, principalmente com customtkinter e escala do Windows.
        self.after(10, self.maximize_window)
        self.after(150, self.maximize_window)
        self.after(700, self.maximize_window)

        # Recalcula as tabelas depois que a janela e o Tabview já receberam
        # suas dimensões reais. Isso evita que Dashboard/Lançamentos abram
        # estreitos na primeira visualização e só se ajustem depois.
        self.after(40, self.refresh_table_layouts)
        self.after(220, self.refresh_table_layouts)
        self.after(850, self.refresh_table_layouts)

        # A janela só é exibida quando a interface já está pronta. O pequeno
        # atraso permite ao Tk processar o primeiro desenho antes de remover
        # o splash do PyInstaller.
        startup_splash_update("Finalizando abertura...")
        self.after(90, self._finish_startup)

    def _finish_startup(self):
        """Exibe a janela pronta e encerra o splash de inicialização."""
        try:
            self.deiconify()
            self.update_idletasks()
            self.maximize_window()
            self.lift()
        except Exception:
            pass
        startup_splash_update("Pronto.")
        try:
            self.after(60, startup_splash_close)
        except Exception:
            startup_splash_close()

    # ------------------------------------------------------------------
    # Atualização automática via GitHub Releases
    # ------------------------------------------------------------------
    def _set_update_status(self, text: str):
        label = getattr(self, "update_status_label", None)
        if label is not None:
            try:
                label.configure(text=str(text))
            except Exception:
                pass

    def _ensure_update_queue_polling(self):
        if self._update_poll_after_id is not None:
            return
        try:
            self._update_poll_after_id = self.after(120, self._poll_update_queue)
        except Exception:
            self._update_poll_after_id = None

    def _poll_update_queue(self):
        self._update_poll_after_id = None
        while True:
            try:
                event = self._update_queue.get_nowait()
            except queue.Empty:
                break

            kind = event[0]
            if kind == "check_ok":
                self._handle_update_check_success(event[1], event[2])
            elif kind == "check_error":
                self._handle_update_check_error(event[1], event[2])
            elif kind == "download_progress":
                self._handle_update_download_progress(event[1], event[2])
            elif kind == "download_ok":
                self._handle_update_download_success(event[1])
            elif kind == "download_error":
                self._handle_update_download_error(event[1])

        if self._update_check_running or self._update_download_running or not self._update_queue.empty():
            self._ensure_update_queue_polling()

    def check_for_updates(self, silent: bool = False):
        if self._update_check_running or self._update_download_running:
            if not silent:
                self.show_info(
                    "Atualizações",
                    "Já existe uma verificação ou download de atualização em andamento.",
                )
            return

        self._update_check_running = True
        self._set_update_status("Verificando atualizações no GitHub...")
        self._ensure_update_queue_polling()

        def worker():
            try:
                release = get_latest_release()
                self._update_queue.put(("check_ok", release, bool(silent)))
            except Exception as exc:
                message = str(exc) if isinstance(exc, UpdateError) else f"Falha inesperada: {exc}"
                self._update_queue.put(("check_error", message, bool(silent)))

        threading.Thread(
            target=worker,
            name="ControleFolgasUpdateCheck",
            daemon=True,
        ).start()

    def _handle_update_check_success(self, release: ReleaseInfo, silent: bool):
        self._update_check_running = False
        current = current_version_for_update_check(APP_VERSION)

        try:
            newer = is_newer_version(release.version, current)
        except Exception as exc:
            self._set_update_status("Não foi possível comparar as versões.")
            if not silent:
                self.show_error("Atualizações", str(exc))
            return

        if not newer:
            self._set_update_status(f"Versão {APP_VERSION} está atualizada.")
            if not silent:
                self.show_info(
                    "Programa atualizado",
                    f"Você já está usando a versão mais recente do Controle de Folgas: {APP_VERSION}.",
                )
            return

        self._set_update_status(f"Nova versão {release.version} disponível.")
        self._prompt_available_update(release)

    def _handle_update_check_error(self, message: str, silent: bool):
        self._update_check_running = False
        self._set_update_status("Não foi possível verificar atualizações agora.")
        if not silent:
            self.show_error("Não foi possível verificar atualizações", message)

    def _prompt_available_update(self, release: ReleaseInfo):
        notes = normalize_spaces(release.notes)
        if len(notes) > 1100:
            notes = notes[:1097].rstrip() + "..."

        message = (
            f"Existe uma nova versão do Controle de Folgas.\n\n"
            f"Versão instalada: {APP_VERSION}\n"
            f"Nova versão: {release.version}\n\n"
        )
        if notes:
            message += f"Novidades:\n{notes}\n\n"

        if not can_auto_install_update():
            message += (
                "A atualização foi detectada, mas a substituição automática do programa "
                "só funciona ao executar o arquivo .exe no Windows."
            )
            self.show_warning("Nova versão disponível", message)
            return

        message += (
            "Deseja baixar e instalar agora?\n\n"
            "Somente o executável será substituído. O banco de dados, configurações, "
            "backups e arquivos de status permanecerão na pasta atual."
        )
        if self.ask_yes_no("Nova versão disponível", message):
            self._start_update_download(release)

    def _start_update_download(self, release: ReleaseInfo):
        if self._update_download_running:
            return

        self._update_download_running = True
        self._open_update_progress_window(release)
        self._ensure_update_queue_polling()

        def progress(downloaded: int, total: int):
            self._update_queue.put(("download_progress", int(downloaded), int(total or 0)))

        def worker():
            try:
                downloaded = download_update(release, progress_callback=progress)
                self._update_queue.put(("download_ok", downloaded))
            except Exception as exc:
                message = str(exc) if isinstance(exc, UpdateError) else f"Falha inesperada: {exc}"
                self._update_queue.put(("download_error", message))

        threading.Thread(
            target=worker,
            name="ControleFolgasUpdateDownload",
            daemon=True,
        ).start()

    def _open_update_progress_window(self, release: ReleaseInfo):
        window = self.create_modal(
            "Atualizando Controle de Folgas",
            560,
            270,
            resizable=False,
            parent=self,
            grab=True,
        )
        window.grid_columnconfigure(0, weight=1)
        window.protocol("WM_DELETE_WINDOW", lambda: None)

        ctk.CTkLabel(
            window,
            text=f"Baixando a versão {release.version}",
            font=("Segoe UI Semibold", 18),
        ).grid(row=0, column=0, padx=24, pady=(26, 8), sticky="w")

        self._update_progress_label = ctk.CTkLabel(
            window,
            text="Preparando download...",
            font=("Segoe UI", 12),
            text_color=("#475569", "#cbd5e1"),
        )
        self._update_progress_label.grid(row=1, column=0, padx=24, pady=(0, 12), sticky="w")

        self._update_progress_bar = ctk.CTkProgressBar(window, height=18)
        self._update_progress_bar.grid(row=2, column=0, padx=24, pady=(0, 12), sticky="ew")
        self._update_progress_bar.set(0)

        ctk.CTkLabel(
            window,
            text="Após o download, o programa será fechado, o EXE será trocado e a nova versão será aberta automaticamente.",
            justify="left",
            anchor="w",
            wraplength=505,
            font=("Segoe UI", 11),
            text_color=("#64748b", "#94a3b8"),
        ).grid(row=3, column=0, padx=24, pady=(0, 22), sticky="ew")

        self._update_progress_window = window

    @staticmethod
    def _format_update_bytes(value: int) -> str:
        value = max(0, int(value or 0))
        units = ("B", "KB", "MB", "GB")
        number = float(value)
        for unit in units:
            if number < 1024 or unit == units[-1]:
                if unit == "B":
                    return f"{int(number)} {unit}"
                return f"{number:.1f} {unit}"
            number /= 1024
        return f"{value} B"

    def _handle_update_download_progress(self, downloaded: int, total: int):
        label = self._update_progress_label
        bar = self._update_progress_bar
        if label is None or bar is None:
            return

        if total > 0:
            fraction = min(1.0, max(0.0, downloaded / total))
            try:
                bar.set(fraction)
            except Exception:
                pass
            text = (
                f"{self._format_update_bytes(downloaded)} de "
                f"{self._format_update_bytes(total)}  •  {fraction * 100:.0f}%"
            )
        else:
            text = f"Baixados {self._format_update_bytes(downloaded)}"
        try:
            label.configure(text=text)
        except Exception:
            pass

    def _close_update_progress_window(self):
        window = self._update_progress_window
        self._update_progress_window = None
        self._update_progress_bar = None
        self._update_progress_label = None
        if window is not None:
            try:
                window.grab_release()
            except Exception:
                pass
            try:
                window.destroy()
            except Exception:
                pass

    def _handle_update_download_success(self, downloaded: DownloadedUpdate):
        self._update_download_running = False
        self._close_update_progress_window()
        self._set_update_status(f"Versão {downloaded.release.version} baixada. Instalando...")

        try:
            launch_update_installer(downloaded)
        except Exception as exc:
            self._set_update_status("Falha ao iniciar a instalação da atualização.")
            self.show_error("Não foi possível instalar a atualização", str(exc))
            return

        self.show_info(
            "Atualização pronta",
            f"A versão {downloaded.release.version} foi baixada e validada com SHA-256.\n\n"
            "Ao fechar esta mensagem, o Controle de Folgas será reiniciado automaticamente na nova versão.",
        )
        self.destroy()

    def _handle_update_download_error(self, message: str):
        self._update_download_running = False
        self._close_update_progress_window()
        self._set_update_status("Falha ao baixar a atualização.")
        self.show_error("Atualização cancelada", message)

    def _start_program_status_tracking(self):
        """Inicia o status ABERTO e agenda confirmações periódicas sem afetar a interface."""
        try:
            self._status_tracker.start()
        except Exception:
            pass
        self._schedule_program_status_heartbeat()

    def _schedule_program_status_heartbeat(self):
        if self._status_closing:
            return
        try:
            self._status_after_id = self.after(STATUS_HEARTBEAT_MS, self._program_status_heartbeat)
        except Exception:
            self._status_after_id = None

    def _program_status_heartbeat(self):
        self._status_after_id = None
        if self._status_closing:
            return
        try:
            self._status_tracker.heartbeat()
        except Exception:
            pass
        self._schedule_program_status_heartbeat()

    def destroy(self):
        """Ao encerrar normalmente, registra FECHADO e o tempo total da sessão."""
        # Evita que callbacks de ajuste de layout agendados via after_idle
        # tentem acessar Treeviews depois de destruídos no fechamento.
        self._ui_closing = True
        startup_splash_close()
        if not getattr(self, "_status_closed", False):
            self._status_closing = True
            after_id = getattr(self, "_status_after_id", None)
            if after_id:
                try:
                    self.after_cancel(after_id)
                except Exception:
                    pass
                self._status_after_id = None
            try:
                self._status_tracker.close()
            except Exception:
                pass
            self._status_closed = True
        try:
            super().destroy()
        except tk.TclError:
            pass

    def get_work_area(self, window=None):
        """Retorna a área útil física do monitor onde a janela está localizada."""
        if sys.platform.startswith("win"):
            try:
                class RECT(ctypes.Structure):
                    _fields_ = [
                        ("left", ctypes.c_long),
                        ("top", ctypes.c_long),
                        ("right", ctypes.c_long),
                        ("bottom", ctypes.c_long),
                    ]

                class MONITORINFO(ctypes.Structure):
                    _fields_ = [
                        ("cbSize", ctypes.c_ulong),
                        ("rcMonitor", RECT),
                        ("rcWork", RECT),
                        ("dwFlags", ctypes.c_ulong),
                    ]

                target = window if window is not None else self
                hwnd = int(target.winfo_id())
                try:
                    parent_hwnd = ctypes.windll.user32.GetParent(hwnd)
                    while parent_hwnd:
                        hwnd = int(parent_hwnd)
                        parent_hwnd = ctypes.windll.user32.GetParent(hwnd)
                except Exception:
                    pass

                monitor = ctypes.windll.user32.MonitorFromWindow(hwnd, 2)  # MONITOR_DEFAULTTONEAREST
                if monitor:
                    info = MONITORINFO()
                    info.cbSize = ctypes.sizeof(MONITORINFO)
                    if ctypes.windll.user32.GetMonitorInfoW(monitor, ctypes.byref(info)):
                        rect = info.rcWork
                        width = int(rect.right - rect.left)
                        height = int(rect.bottom - rect.top)
                        if width > 0 and height > 0:
                            return width, height, int(rect.left), int(rect.top)
            except Exception:
                pass

            # Fallback para a área útil do monitor principal.
            try:
                class RECT(ctypes.Structure):
                    _fields_ = [
                        ("left", ctypes.c_long),
                        ("top", ctypes.c_long),
                        ("right", ctypes.c_long),
                        ("bottom", ctypes.c_long),
                    ]
                rect = RECT()
                ctypes.windll.user32.SystemParametersInfoW(48, 0, ctypes.byref(rect), 0)
                width = int(rect.right - rect.left)
                height = int(rect.bottom - rect.top)
                if width > 0 and height > 0:
                    return width, height, int(rect.left), int(rect.top)
            except Exception:
                pass

        target = window if window is not None else self
        try:
            return target.winfo_screenwidth(), target.winfo_screenheight(), 0, 0
        except Exception:
            return 1360, 720, 0, 0

    def maximize_window(self):
        """Maximiza no Windows; se falhar, força a geometria da área útil."""
        try:
            self.update_idletasks()
            self.deiconify()
        except Exception:
            pass

        if sys.platform.startswith("win"):
            # 1) Método nativo do Tk.
            try:
                self.state("zoomed")
            except Exception:
                pass

            # 2) Método direto pela API do Windows.
            try:
                hwnd = self.winfo_id()
                # Sobe para o handle pai quando necessário.
                try:
                    parent = ctypes.windll.user32.GetParent(hwnd)
                    while parent:
                        hwnd = parent
                        parent = ctypes.windll.user32.GetParent(hwnd)
                except Exception:
                    pass
                ctypes.windll.user32.ShowWindow(hwnd, 3)  # SW_MAXIMIZE
            except Exception:
                pass

        # 3) Fallback final: área útil inteira. Isso resolve mesmo quando o Windows
        # não aceita o estado zoomed por causa de escala, RDP, VM ou tema.
        try:
            self.update_idletasks()
            work_w, work_h, work_x, work_y = self.get_work_area()
            atual_w, atual_h = self.winfo_width(), self.winfo_height()
            if atual_w < int(work_w * 0.95) or atual_h < int(work_h * 0.90):
                scale = self._ctk_window_scale(self)
                logical_w = max(1, int(work_w / scale))
                logical_h = max(1, int(work_h / scale))
                self.geometry(f"{logical_w}x{logical_h}+{work_x}+{work_y}")
                self.update_idletasks()
        except Exception:
            pass

        # A maximização pode ocorrer depois da criação das tabelas; refaça o
        # cálculo com a geometria final para que a primeira exibição já nasça correta.
        try:
            self.after_idle(self.refresh_table_layouts)
        except Exception:
            pass

    def on_window_resize(self, event=None):
        if getattr(self, "_ui_closing", False):
            return
        if event is not None and event.widget is not self:
            return
        try:
            width = max(self.winfo_width(), 980)
            right_area = max(width - self.form_width - 70, 360)
            wrap_status = max(300, right_area - 60)
            wrap_card = max(220, int((width - 80) / 2) - 60)

            for widget_name in ("label_status_funcionario", "label_saldo_filtro"):
                if hasattr(self, widget_name):
                    getattr(self, widget_name).configure(wraplength=wrap_status)

            for attr in ("card_saldo", "card_creditos", "card_debitos", "card_lancamentos"):
                if hasattr(self, attr):
                    getattr(self, attr).configure(wraplength=wrap_card)

            # As tabelas principais recalculam as larguras para aproveitar a
            # janela disponível sem exibir barras horizontais desnecessárias.
            if hasattr(self, "tree_resumo"):
                self.after_idle(self.adjust_dashboard_columns)
            if hasattr(self, "employee_list_card"):
                self.after_idle(self.resize_employee_list_card)
            if hasattr(self, "tree_entries"):
                self.after_idle(self.adjust_entries_columns)
        except Exception:
            pass

    def _available_table_width(self, tree=None, container=None, minimum=620):
        """Retorna uma largura útil confiável mesmo quando a aba ainda está oculta."""
        # Se a tabela já estiver visível, sua própria largura é a referência mais precisa.
        try:
            width = int(tree.winfo_width()) if tree is not None else 0
            if width > 300:
                return max(int(minimum), width - 6)
        except Exception:
            pass

        # Na primeira abertura de uma aba oculta, o Treeview pode reportar apenas 1 px.
        # O contêiner, Tabview ou janela principal fornecem então a largura real esperada.
        for widget, margin in (
            (container, 24),
            (getattr(self, "tabs", None), 30),
            (self, 48),
        ):
            if widget is None:
                continue
            try:
                width = int(widget.winfo_width())
                if width > 300:
                    return max(int(minimum), width - margin)
            except Exception:
                pass

        return max(int(minimum), int(getattr(self, "work_w", minimum)) - 48)

    def refresh_table_layouts(self):
        """Força Dashboard, Lançamentos e Funcionários a usarem o tamanho real da janela."""
        try:
            self.update_idletasks()
        except Exception:
            pass
        try:
            if hasattr(self, "tree_resumo"):
                self.adjust_dashboard_columns()
        except Exception:
            pass
        try:
            if hasattr(self, "tree_entries"):
                self.adjust_entries_columns()
        except Exception:
            pass
        try:
            if hasattr(self, "employee_list_card"):
                self.resize_employee_list_card()
        except Exception:
            pass

    def on_tab_changed(self, *_args):
        """Reajusta a área escolhida sempre que o usuário troca de aba."""
        # Duas passagens cobrem o momento em que o CustomTkinter mapeia a aba
        # e o instante logo depois, quando todos os widgets já têm largura real.
        self.after(1, self.refresh_table_layouts)
        self.after(70, self.refresh_table_layouts)

    def center_window(self):
        self.maximize_window()

    def setup_style(self):
        self.ttk_style = ttk.Style(self)
        try:
            self.ttk_style.theme_use("clam")
        except Exception:
            pass
        self.apply_treeview_theme()

    def apply_treeview_theme(self):
        """Aplica um visual corporativo limpo às tabelas do sistema."""
        dark = ctk.get_appearance_mode() == "Dark"

        if dark:
            table_bg = "#111827"
            table_fg = "#e5e7eb"
            zebra_bg = "#141e2f"
            heading_bg = "#1b2638"
            heading_fg = "#f8fafc"
            border = "#334155"
            muted_fg = "#7c8aa0"
            inactive_bg = "#151b26"
            positive_bg, positive_fg = "#10271d", "#bbf7d0"
            negative_bg, negative_fg = "#2b171b", "#fecaca"
        else:
            table_bg = "#ffffff"
            table_fg = "#0f172a"
            zebra_bg = "#f8fafc"
            heading_bg = "#eef2f7"
            heading_fg = "#0f172a"
            border = "#cbd5e1"
            muted_fg = "#94a3b8"
            inactive_bg = "#f8fafc"
            positive_bg, positive_fg = "#f0fdf4", "#166534"
            negative_bg, negative_fg = "#fff1f2", "#991b1b"

        # Tabela principal: sem grade pesada, cabeçalho plano, linhas com mais
        # respiro e contraste sutil. A moldura externa é feita pelo CTkFrame.
        self.ttk_style.configure(
            "Professional.Treeview",
            rowheight=34,
            font=("Segoe UI", max(10, self.body_font)),
            background=table_bg,
            fieldbackground=table_bg,
            foreground=table_fg,
            borderwidth=0,
            relief="flat",
        )
        self.ttk_style.configure(
            "Professional.Treeview.Heading",
            font=("Segoe UI Semibold", max(10, self.body_font)),
            background=heading_bg,
            foreground=heading_fg,
            bordercolor=border,
            lightcolor=border,
            darkcolor=border,
            borderwidth=1,
            relief="flat",
            padding=(8, 8),
        )
        self.ttk_style.map(
            "Professional.Treeview",
            background=[("selected", "#2563eb")],
            foreground=[("selected", "#ffffff")],
        )
        self.ttk_style.map(
            "Professional.Treeview.Heading",
            background=[("active", heading_bg)],
            foreground=[("active", heading_fg)],
        )

        # Tabelas administrativas mantêm o mesmo idioma visual, mas um pouco
        # mais compactas para acomodar grande quantidade de informações.
        self.ttk_style.configure(
            "Treeview",
            rowheight=30,
            font=("Segoe UI", max(9, self.body_font)),
            background=table_bg,
            fieldbackground=table_bg,
            foreground=table_fg,
            borderwidth=0,
            relief="flat",
        )
        self.ttk_style.configure(
            "Treeview.Heading",
            font=("Segoe UI Semibold", max(9, self.body_font)),
            background=heading_bg,
            foreground=heading_fg,
            bordercolor=border,
            borderwidth=1,
            relief="flat",
            padding=(6, 7),
        )
        self.ttk_style.map(
            "Treeview",
            background=[("selected", "#2563eb")],
            foreground=[("selected", "#ffffff")],
        )
        self.ttk_style.map(
            "Treeview.Heading",
            background=[("active", heading_bg)],
            foreground=[("active", heading_fg)],
        )

        for name in ("tree_resumo", "tree_entries"):
            tree = getattr(self, name, None)
            if tree is not None:
                tree.configure(style="Professional.Treeview")
                tree.tag_configure("zebra", background=zebra_bg, foreground=table_fg)
                tree.update_idletasks()

        if hasattr(self, "tree_resumo"):
            self.tree_resumo.tag_configure("saldo_positivo", background=positive_bg, foreground=positive_fg)
            self.tree_resumo.tag_configure("saldo_negativo", background=negative_bg, foreground=negative_fg)

        if hasattr(self, "tree_entries"):
            self.tree_entries.tag_configure("movimento_credito", background=positive_bg, foreground=positive_fg)
            self.tree_entries.tag_configure("movimento_debito", background=negative_bg, foreground=negative_fg)


    def get_logo_image(self, width: int, height: int):
        """Carrega o logo 3D uma única vez por tamanho para uso na interface."""
        key = (int(width), int(height))
        cached = self._ctk_image_cache.get(key)
        if cached is not None:
            return cached

        path = main_logo_path()
        if not path.exists():
            return None

        try:
            with Image.open(path) as source:
                pil_image = source.convert("RGBA").copy()
            ctk_image = ctk.CTkImage(
                light_image=pil_image,
                dark_image=pil_image,
                size=key,
            )
            self._ctk_image_cache[key] = ctk_image
            return ctk_image
        except Exception:
            return None

    def main_header_title(self) -> str:
        """Monta o título principal incluindo o nome do serviço, quando já cadastrado."""
        service_name = self.db.get_service_name()
        if service_name:
            return f"Controle de Folgas e Banco de Horas - {service_name}"
        return "Controle de Folgas e Banco de Horas"

    def refresh_header_title(self):
        """Atualiza o cabeçalho principal após cadastrar ou alterar o nome do serviço."""
        if hasattr(self, "header_title_label"):
            self.header_title_label.configure(text=self.main_header_title())

    def build_ui(self):
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)

        header = ctk.CTkFrame(self, corner_radius=0, fg_color=("#ffffff", "#111827"), height=88)
        header.grid(row=0, column=0, sticky="ew")
        header.grid_propagate(False)
        header.grid_columnconfigure(1, weight=1)

        # Logo principal. O objeto CTkImage fica em cache para não ser coletado
        # pelo Python e para poder ser reutilizado em outras janelas.
        self.header_logo_image = self.get_logo_image(78, 78)
        if self.header_logo_image is not None:
            ctk.CTkLabel(
                header,
                text="",
                image=self.header_logo_image,
                width=80,
                height=80,
            ).grid(row=0, column=0, rowspan=2, padx=(12, 4), pady=4, sticky="w")

        self.header_title_label = ctk.CTkLabel(
            header,
            text=self.main_header_title(),
            font=("Segoe UI Semibold", self.header_title_font),
            anchor="w",
        )
        self.header_title_label.grid(row=0, column=1, padx=(8, 18), pady=(16, 0), sticky="ew")
        ctk.CTkLabel(
            header,
            text=f"Cadastro, lançamentos, saldo automático, abatimento de folgas e exportação filtrada • Versão {APP_VERSION}",
            font=("Segoe UI", self.header_sub_font),
            text_color=("#475569", "#cbd5e1"),
        ).grid(row=1, column=1, padx=(8, 18), pady=(0, 16), sticky="w")

        top_actions = ctk.CTkFrame(header, fg_color="transparent")
        top_actions.grid(row=0, column=2, rowspan=2, padx=14, pady=9, sticky="e")
        ctk.CTkButton(top_actions, text="Imprimir para Assinaturas", command=self.open_print_signatures_dialog, height=32, width=150, corner_radius=12).pack(side="left", padx=4)
        ctk.CTkButton(top_actions, text="Exportar Excel", command=self.export_excel, height=32, width=126, corner_radius=12).pack(side="left", padx=4)
        ctk.CTkButton(top_actions, text="Tema", command=self.toggle_theme, height=32, width=74, corner_radius=12).pack(side="left", padx=4)

        self.tabs = ctk.CTkTabview(self, corner_radius=14, segmented_button_selected_color="#2563eb", command=self.on_tab_changed)
        self.tabs.grid(row=1, column=0, padx=8, pady=8, sticky="nsew")
        self.tab_dashboard = self.tabs.add("Dashboard")
        self.tab_lancamentos = self.tabs.add("Lançamentos")
        self.tab_funcionarios = self.tabs.add("Funcionários")
        self.tab_tutoriais = self.tabs.add("Tutoriais")
        self.tab_ajuda = self.tabs.add("Ajuda")

        self.build_dashboard()
        self.build_lancamentos()
        self.build_funcionarios()
        self.build_tutoriais()
        self.build_ajuda()

    def build_dashboard(self):
        tab = self.tab_dashboard
        tab.grid_columnconfigure((0, 1), weight=1, uniform="cards")
        tab.grid_rowconfigure(3, weight=1)

        self.card_saldo = self.create_card(tab, "Saldo geral", "0", 0, 0)
        self.card_creditos = self.create_card(tab, "Créditos", "0", 0, 1)
        self.card_debitos = self.create_card(tab, "Débitos", "0", 1, 0)
        self.card_lancamentos = self.create_card(tab, "Lançamentos", "0", 1, 1)

        filter_frame = ctk.CTkFrame(tab, corner_radius=13)
        filter_frame.grid(row=2, column=0, columnspan=2, padx=6, pady=(4, 6), sticky="ew")
        filter_frame.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(
            filter_frame,
            text="Resumo por funcionário",
            font=("Segoe UI Semibold", 14),
        ).grid(row=0, column=0, padx=12, pady=8, sticky="w")
        ctk.CTkLabel(
            filter_frame,
            text="Verde = Tem horas para tirar  •  Vermelho = Está devendo horas",
            font=("Segoe UI", 10),
            text_color=("#64748b", "#94a3b8"),
        ).grid(row=0, column=1, padx=8, pady=8, sticky="w")
        ctk.CTkButton(
            filter_frame,
            text="Atualizar",
            command=self.refresh_dashboard,
            height=30,
            width=120,
            corner_radius=11,
        ).grid(row=0, column=2, padx=12, pady=8, sticky="e")

        tree_frame = ctk.CTkFrame(
            tab,
            corner_radius=13,
            border_width=1,
            border_color=("#d8e0ea", "#334155"),
            fg_color=("#ffffff", "#111827"),
        )
        tree_frame.grid(row=3, column=0, columnspan=2, padx=6, pady=4, sticky="nsew")
        tree_frame.grid_rowconfigure(0, weight=1)
        tree_frame.grid_columnconfigure(0, weight=1)
        self.dashboard_table_frame = tree_frame
        tree_frame.bind("<Configure>", lambda _event: self.after_idle(self.adjust_dashboard_columns))

        columns = ("funcionario", "status", "saldo", "creditos", "debitos", "lancamentos")
        self.tree_resumo = ttk.Treeview(
            tree_frame,
            columns=columns,
            show="headings",
            style="Professional.Treeview",
        )
        headers = [
            ("funcionario", "Funcionário", "w"),
            ("status", "Situação atual", "center"),
            ("saldo", "Saldo", "center"),
            ("creditos", "Créditos", "center"),
            ("debitos", "Débitos", "center"),
            ("lancamentos", "Lançamentos", "center"),
        ]
        for col, title, anchor in headers:
            self.tree_resumo.heading(col, text=title, anchor="center")
            self.tree_resumo.column(col, width=120, minwidth=60, anchor=anchor, stretch=False)

        self.tree_resumo.grid(row=0, column=0, padx=1, pady=1, sticky="nsew")
        yscroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree_resumo.yview)
        yscroll.grid(row=0, column=1, sticky="ns", pady=1)
        self.tree_resumo.configure(yscrollcommand=yscroll.set)
        self.after_idle(self.adjust_dashboard_columns)

    def adjust_dashboard_columns(self):
        """Distribui a largura do Dashboard proporcionalmente ao espaço disponível."""
        tree = getattr(self, "tree_resumo", None)
        if tree is None:
            return
        try:
            available = self._available_table_width(
                tree=tree,
                container=getattr(self, "dashboard_table_frame", None),
                minimum=620,
            )
            proportions = {
                "funcionario": 0.28,
                "status": 0.19,
                "saldo": 0.17,
                "creditos": 0.14,
                "debitos": 0.14,
                "lancamentos": 0.08,
            }
            for col, proportion in proportions.items():
                tree.column(col, width=max(70, int(available * proportion)), stretch=False)
        except Exception:
            pass

    def create_card(self, parent, title, value, row, col):
        frame = ctk.CTkFrame(parent, corner_radius=13, height=74)
        frame.grid(row=row, column=col, padx=6, pady=4, sticky="ew")
        frame.grid_propagate(False)
        ctk.CTkLabel(
            frame,
            text=title,
            font=("Segoe UI", 11),
            text_color=("#64748b", "#cbd5e1"),
        ).pack(anchor="w", padx=12, pady=(8, 1))
        label_value = ctk.CTkLabel(
            frame,
            text=value,
            font=("Segoe UI Semibold", self.card_font),
            justify="left",
            anchor="w",
            wraplength=420,
        )
        label_value.pack(anchor="w", padx=12, pady=(0, 7), fill="x")
        return label_value

    @staticmethod
    def validate_numeric_entry(value: str) -> bool:
        """Permite apenas números inteiros ou campo vazio em Dias, Horas e Minutos."""
        return value == "" or value.isdigit()

    def build_lancamentos(self):
        """Tela de consulta dos lançamentos. Cadastro e edição ficam em janelas próprias."""
        tab = self.tab_lancamentos
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(2, weight=1)

        actions = ctk.CTkFrame(tab, corner_radius=13)
        actions.grid(row=0, column=0, padx=6, pady=(6, 3), sticky="ew")
        actions.grid_columnconfigure(4, weight=1)

        ctk.CTkLabel(actions, text="Lançamentos", font=("Segoe UI Semibold", 18)).grid(
            row=0, column=0, padx=(12, 14), pady=10, sticky="w"
        )
        ctk.CTkButton(
            actions, text="+ Novo lançamento", command=self.open_new_entry,
            height=34, width=150, corner_radius=11,
        ).grid(row=0, column=1, padx=4, pady=9)
        ctk.CTkButton(
            actions, text="Editar selecionado", command=self.edit_selected_entry,
            height=34, width=145, corner_radius=11,
            fg_color="#0f766e", hover_color="#115e59",
        ).grid(row=0, column=2, padx=4, pady=9)
        ctk.CTkButton(
            actions, text="Enviar para a Lixeira", command=self.admin_delete_selected_entry,
            height=34, width=155, corner_radius=11,
            fg_color="#b45309", hover_color="#92400e",
        ).grid(row=0, column=3, padx=4, pady=9)
        ctk.CTkButton(
            actions, text="Opções administrativas", command=self.open_admin_options,
            height=34, width=165, corner_radius=11,
            fg_color="#475569", hover_color="#334155",
        ).grid(row=0, column=5, padx=(8, 12), pady=9, sticky="e")

        top = ctk.CTkFrame(tab, corner_radius=13)
        top.grid(row=1, column=0, padx=6, pady=3, sticky="ew")
        top.grid_columnconfigure(1, weight=1)
        top.grid_columnconfigure(3, weight=1)

        ctk.CTkLabel(top, text="Filtros", font=("Segoe UI Semibold", 14)).grid(
            row=0, column=0, padx=(10, 4), pady=(22, 2), sticky="sw"
        )

        frame_filter_func = ctk.CTkFrame(top, fg_color="transparent")
        frame_filter_func.grid(row=0, column=1, padx=4, pady=(4, 2), sticky="ew")
        frame_filter_func.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(frame_filter_func, text="Funcionário", font=("Segoe UI", 10)).grid(row=0, column=0, pady=(0, 2), sticky="w")
        self.combo_filter_func = ctk.CTkComboBox(
            frame_filter_func, values=["Todos"], height=28, width=180, corner_radius=10,
            command=lambda _: self.refresh_entries(),
        )
        self.combo_filter_func.set("Todos")
        self.combo_filter_func.grid(row=1, column=0, sticky="ew")

        frame_filter_year = ctk.CTkFrame(top, fg_color="transparent")
        frame_filter_year.grid(row=0, column=2, padx=4, pady=(4, 2), sticky="w")
        ctk.CTkLabel(frame_filter_year, text="Ano", font=("Segoe UI", 10)).grid(row=0, column=0, pady=(0, 2), sticky="w")
        self.combo_filter_year = ctk.CTkComboBox(
            frame_filter_year, values=["Todos"], height=28, width=90, corner_radius=10,
            command=lambda _: self.refresh_entries(),
        )
        self.combo_filter_year.set("Todos")
        self.combo_filter_year.grid(row=1, column=0, sticky="w")

        frame_filter_month = ctk.CTkFrame(top, fg_color="transparent")
        frame_filter_month.grid(row=0, column=3, padx=4, pady=(4, 2), sticky="ew")
        frame_filter_month.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(frame_filter_month, text="Mês", font=("Segoe UI", 10)).grid(row=0, column=0, pady=(0, 2), sticky="w")
        self.combo_filter_month = ctk.CTkComboBox(
            frame_filter_month, values=MESES_FILTRO, height=28, width=125, corner_radius=10,
            command=lambda _: self.refresh_entries(),
        )
        self.combo_filter_month.set("Todos")
        self.combo_filter_month.grid(row=1, column=0, sticky="ew")

        ctk.CTkButton(top, text="Buscar", command=self.refresh_entries, height=28, width=105, corner_radius=10).grid(
            row=0, column=4, padx=(4, 10), pady=(22, 2), sticky="se"
        )
        self.entry_search = ctk.CTkEntry(top, placeholder_text="Buscar em categoria ou observação", height=28, corner_radius=10)
        self.entry_search.grid(row=1, column=0, columnspan=5, padx=10, pady=(3, 5), sticky="ew")
        self.entry_search.bind("<Return>", lambda _event: self.refresh_entries())

        self.status_box = ctk.CTkFrame(top, corner_radius=12, fg_color="#dbeafe")
        self.status_box.grid(row=2, column=0, columnspan=5, padx=10, pady=(3, 5), sticky="ew")
        self.status_box.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            self.status_box, text="STATUS DO BANCO DE HORAS", font=("Segoe UI Semibold", 10), text_color="#475569"
        ).grid(row=0, column=0, padx=12, pady=(6, 0), sticky="w")
        self.label_status_funcionario = ctk.CTkLabel(
            self.status_box, text="Selecione um funcionário para ver o saldo",
            font=("Segoe UI Semibold", 16), text_color="#1e40af",
            justify="left", anchor="w", wraplength=720,
        )
        self.label_status_funcionario.grid(row=1, column=0, padx=12, pady=(0, 8), sticky="ew")

        self.label_saldo_filtro = ctk.CTkLabel(
            top, text="Saldo filtrado: 0 | Créditos: 0 | Débitos abatidos: 0",
            font=("Segoe UI Semibold", 11), text_color=("#334155", "#dbeafe"),
            justify="left", anchor="w", wraplength=900,
        )
        self.label_saldo_filtro.grid(row=3, column=0, columnspan=5, padx=10, pady=(0, 6), sticky="ew")

        list_frame = ctk.CTkFrame(
            tab, corner_radius=13, border_width=1,
            border_color=("#d8e0ea", "#334155"), fg_color=("#ffffff", "#111827"),
        )
        list_frame.grid(row=2, column=0, padx=6, pady=(3, 6), sticky="nsew")
        list_frame.grid_rowconfigure(0, weight=1)
        list_frame.grid_columnconfigure(0, weight=1)
        self.entries_table_frame = list_frame
        list_frame.bind("<Configure>", lambda _event: self.after_idle(self.adjust_entries_columns))

        columns = ("id", "data", "funcionario", "movimento", "categoria", "quantidade", "valor", "saldo_funcionario", "observacao")
        self.tree_entries = ttk.Treeview(
            list_frame, columns=columns,
            displaycolumns=("data", "funcionario", "movimento", "categoria", "quantidade", "valor", "saldo_funcionario", "observacao"),
            show="headings", style="Professional.Treeview",
        )
        headers = [
            ("id", "ID", 0, "center"),
            ("data", "Data", 90, "center"),
            ("funcionario", "Funcionário", 180, "w"),
            ("movimento", "Tipo", 90, "center"),
            ("categoria", "Categoria", 145, "w"),
            ("quantidade", "Quantidade", 110, "center"),
            ("valor", "Valor", 115, "center"),
            ("saldo_funcionario", "Saldo", 125, "center"),
            ("observacao", "Observação", 260, "w"),
        ]
        for col, title, width, anchor in headers:
            self.tree_entries.heading(col, text=title, anchor="center")
            self.tree_entries.column(
                col, width=width, minwidth=0 if col == "id" else 55,
                anchor=anchor, stretch=False,
            )
        self.tree_entries.grid(row=0, column=0, padx=1, pady=1, sticky="nsew")
        self.tree_entries.bind("<<TreeviewSelect>>", self.on_entry_select)
        self.tree_entries.bind("<Double-1>", lambda _event: self.edit_selected_entry())
        yscroll = ttk.Scrollbar(list_frame, orient="vertical", command=self.tree_entries.yview)
        yscroll.grid(row=0, column=1, sticky="ns", pady=1)
        xscroll = AutoHideScrollbar(list_frame, orient="horizontal", command=self.tree_entries.xview)
        xscroll.grid(row=1, column=0, sticky="ew", padx=1)
        self.tree_entries.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)

    def adjust_entries_columns(self):
        """Ajusta Lançamentos ao conteúdo e ocupa a largura disponível quando possível."""
        if getattr(self, "_ui_closing", False):
            return
        tree = getattr(self, "tree_entries", None)
        if tree is None:
            return
        # Um <Configure> pode ter enfileirado este callback pelo after_idle
        # imediatamente antes de a janela ser fechada. O objeto Python ainda
        # existe, mas o widget Tcl/Tk já pode ter sido destruído.
        try:
            if not tree.winfo_exists():
                return
        except tk.TclError:
            return
        try:
            data_font = tkfont.Font(family="Segoe UI", size=max(10, self.body_font))
            header_font = tkfont.Font(family="Segoe UI Semibold", size=max(10, self.body_font))
        except Exception:
            return

        # O ID permanece no Treeview apenas para uso interno, sem ser mostrado.
        try:
            tree.column("id", width=0, minwidth=0, stretch=False)
        except tk.TclError:
            # Treeview destruído durante o encerramento: não é falha de dados.
            return

        limits = {
            "data": (82, 102),
            "funcionario": (145, 300),
            "movimento": (72, 110),
            "categoria": (115, 230),
            "quantidade": (90, 180),
            "valor": (95, 190),
            "saldo_funcionario": (105, 230),
            "observacao": (170, 480),
        }
        padding = 26
        items = tree.get_children("")
        widths = {}

        for column in tree["displaycolumns"]:
            min_width, max_width = limits[column]
            measured = header_font.measure(str(tree.heading(column, option="text") or "")) + padding
            # Uma amostra suficientemente grande mantém o cálculo rápido mesmo
            # em bancos com muitos lançamentos.
            for item in items[:250]:
                value = str(tree.set(item, column) or "")
                measured = max(measured, data_font.measure(value) + padding)
                if measured >= max_width:
                    measured = max_width
                    break
            widths[column] = max(min_width, min(measured, max_width))

        available = self._available_table_width(
            tree=tree,
            container=getattr(self, "entries_table_frame", None),
            minimum=880,
        )
        total = sum(widths.values())

        # Em telas largas, o espaço restante é distribuído nas colunas textuais
        # para não deixar uma área vazia no fim da tabela.
        if available > total:
            extra = available - total
            flex = (("funcionario", 0.25), ("categoria", 0.20), ("observacao", 0.55))
            for col, weight in flex:
                min_w, max_w = limits[col]
                widths[col] = min(max_w, widths[col] + int(extra * weight))

            # Se ainda sobrar espaço depois dos limites, deixe Observação absorver.
            remaining = available - sum(widths.values())
            if remaining > 0:
                widths["observacao"] += remaining

        for col, width in widths.items():
            tree.column(col, width=max(1, int(width)), minwidth=limits[col][0], stretch=False)

    def add_form_field(self, parent, label, widget):
        ctk.CTkLabel(parent, text=label, font=("Segoe UI", 11)).pack(anchor="w", padx=12, pady=(5, 1))
        widget.pack(fill="x", padx=12, pady=(0, 2))

    def build_funcionarios(self):
        """Tela profissional de consulta de funcionários, sem dependência do ttk.Treeview."""
        tab = self.tab_funcionarios
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(1, weight=1)

        # Barra superior de ações.
        actions = ctk.CTkFrame(tab, corner_radius=13)
        actions.grid(row=0, column=0, padx=6, pady=(6, 3), sticky="ew")
        actions.grid_columnconfigure(3, weight=1)

        title_frame = ctk.CTkFrame(actions, fg_color="transparent")
        title_frame.grid(row=0, column=0, padx=(14, 16), pady=9, sticky="w")
        ctk.CTkLabel(
            title_frame, text="Funcionários", font=("Segoe UI Semibold", 18), anchor="w"
        ).pack(anchor="w")
        self.employee_count_label = ctk.CTkLabel(
            title_frame,
            text="",
            font=("Segoe UI", 11),
            text_color=("#64748b", "#94a3b8"),
            anchor="w",
        )
        self.employee_count_label.pack(anchor="w", pady=(1, 0))

        ctk.CTkButton(
            actions, text="+ Novo funcionário", command=self.open_new_employee,
            height=36, width=155, corner_radius=11,
        ).grid(row=0, column=1, padx=4, pady=9)
        ctk.CTkButton(
            actions, text="Editar funcionário selecionado", command=self.edit_selected_employee,
            height=36, width=215, corner_radius=11,
            fg_color="#0f766e", hover_color="#115e59",
        ).grid(row=0, column=2, padx=4, pady=9)

        search_frame = ctk.CTkFrame(actions, fg_color="transparent")
        search_frame.grid(row=0, column=4, padx=(14, 14), pady=7, sticky="e")
        ctk.CTkLabel(
            search_frame, text="Buscar funcionário", font=("Segoe UI", 11),
            text_color=("#475569", "#cbd5e1"),
        ).pack(anchor="w", pady=(0, 2))
        self.emp_search_var = tk.StringVar()
        self.emp_search_entry = ctk.CTkEntry(
            search_frame, textvariable=self.emp_search_var,
            placeholder_text="Digite um nome...", width=260, height=31, corner_radius=10,
        )
        self.emp_search_entry.pack(fill="x")
        self.emp_search_var.trace_add("write", lambda *_args: self.refresh_employees())

        # A listagem fica centralizada e possui largura máxima. Em monitores grandes
        # ela não fica excessivamente esticada; em telas menores reduz automaticamente.
        self.employee_list_host = ctk.CTkFrame(tab, fg_color="transparent")
        self.employee_list_host.grid(row=1, column=0, padx=6, pady=(3, 6), sticky="nsew")

        self.employee_list_card = ctk.CTkFrame(
            self.employee_list_host,
            corner_radius=14,
            border_width=1,
            border_color=("#d8e0ea", "#334155"),
            fg_color=("#ffffff", "#111827"),
        )
        self.employee_list_card.place(relx=0.5, rely=0, anchor="n", relheight=1.0)
        self.employee_list_card.grid_columnconfigure(0, weight=1)
        self.employee_list_card.grid_rowconfigure(1, weight=1)

        # Cabeçalho fixo: permanece visível enquanto a lista rola.
        self.employee_list_header = ctk.CTkFrame(
            self.employee_list_card,
            height=46,
            corner_radius=10,
            fg_color=("#eef2f7", "#1b2638"),
        )
        self.employee_list_header.grid(row=0, column=0, padx=10, pady=(10, 4), sticky="ew")
        self.employee_list_header.grid_propagate(False)
        self._configure_employee_grid(self.employee_list_header)

        header_font = ("Segoe UI Semibold", 13)
        ctk.CTkLabel(
            self.employee_list_header, text="FUNCIONÁRIO", font=header_font,
            anchor="w", text_color=("#334155", "#f8fafc"),
        ).grid(row=0, column=0, padx=(18, 8), pady=11, sticky="ew")
        ctk.CTkLabel(
            self.employee_list_header, text="REGIME DE TRABALHO", font=header_font,
            anchor="w", text_color=("#334155", "#f8fafc"),
        ).grid(row=0, column=1, padx=8, pady=11, sticky="ew")
        ctk.CTkLabel(
            self.employee_list_header, text="JORNADA", font=header_font,
            anchor="center", text_color=("#334155", "#f8fafc"),
        ).grid(row=0, column=2, padx=8, pady=11, sticky="ew")
        ctk.CTkLabel(
            self.employee_list_header, text="SITUAÇÃO", font=header_font,
            anchor="center", text_color=("#334155", "#f8fafc"),
        ).grid(row=0, column=3, padx=(8, 18), pady=11, sticky="ew")

        self.employee_scroll = ctk.CTkScrollableFrame(
            self.employee_list_card,
            corner_radius=0,
            fg_color=("#ffffff", "#111827"),
            scrollbar_button_color=("#cbd5e1", "#334155"),
            scrollbar_button_hover_color=("#94a3b8", "#475569"),
        )
        self.employee_scroll.grid(row=1, column=0, padx=(2, 3), pady=(0, 3), sticky="nsew")
        self.employee_scroll.grid_columnconfigure(0, weight=1)

        self.employee_row_widgets = {}
        self.employee_empty_label = None

        # O host informa sua largura real por Configure; não dependemos mais da
        # largura inicial de uma aba ainda oculta, como ocorria com o Treeview.
        self.employee_list_host.bind("<Configure>", self.resize_employee_list_card)
        self.after_idle(self.resize_employee_list_card)

    @staticmethod
    def _configure_employee_grid(frame):
        """Distribuição visual das quatro informações da listagem de funcionários."""
        frame.grid_columnconfigure(0, weight=42, minsize=260)
        frame.grid_columnconfigure(1, weight=33, minsize=230)
        frame.grid_columnconfigure(2, weight=10, minsize=95)
        frame.grid_columnconfigure(3, weight=15, minsize=120)

    def resize_employee_list_card(self, event=None):
        """Mantém a listagem centralizada, responsiva e sem ficar exageradamente larga."""
        host = getattr(self, "employee_list_host", None)
        card = getattr(self, "employee_list_card", None)
        if host is None or card is None:
            return
        try:
            available = int(event.width) if event is not None and getattr(event, "widget", None) is host else int(host.winfo_width())
            if available <= 1:
                self.after(80, self.resize_employee_list_card)
                return
            # 1320 px dá boa leitura em Full HD/2K, mas a lista reduz em notebooks.
            target = min(1320, max(760, available - 28))
            target = min(target, max(1, available - 8))
            card.place_configure(width=target)
        except Exception:
            pass

    def _employee_row_colors(self, active=True, zebra=False):
        if not active:
            return ("#f8fafc", "#151b26"), ("#94a3b8", "#7c8aa0")
        if zebra:
            return ("#f8fafc", "#141e2f"), ("#0f172a", "#e5e7eb")
        return ("#ffffff", "#111827"), ("#0f172a", "#e5e7eb")

    def _bind_employee_row_widget(self, widget, employee_id):
        widget.bind("<Button-1>", lambda _event, eid=employee_id: self.select_employee_card(eid))
        widget.bind("<Double-1>", lambda _event, eid=employee_id: self.open_employee_from_double_click(eid))

    def select_employee_card(self, employee_id):
        """Seleciona visualmente uma linha da lista profissional de funcionários."""
        self.selected_employee_id = int(employee_id) if employee_id is not None else None
        rows = getattr(self, "employee_row_widgets", {})
        for eid, data in rows.items():
            selected = eid == self.selected_employee_id
            data["row"].configure(
                fg_color=("#eaf2ff", "#1e3655") if selected else data["base_bg"],
                border_width=1 if selected else 0,
                border_color="#3b82f6" if selected else ("#ffffff", "#111827"),
            )
            data["accent"].configure(fg_color="#3b82f6" if selected else "transparent")

    def open_employee_from_double_click(self, employee_id):
        self.select_employee_card(employee_id)
        self.edit_selected_employee()

    def on_regime_change(self, value):
        """Compatibilidade: retorna a jornada automática para os regimes padronizados."""
        return self.regime_hours(value)

    def open_tutorial_link(self, url: str):
        """Abre um tutorial no navegador padrão do Windows."""
        url = str(url or "").strip()
        if not url.lower().startswith(("https://", "http://")):
            self.show_error(
                "Link indisponível",
                "O endereço deste tutorial não está configurado corretamente.",
            )
            return
        try:
            opened = webbrowser.open_new_tab(url)
            if opened is False:
                raise RuntimeError("O Windows não informou um navegador disponível para abrir o link.")
        except Exception as exc:
            self.show_error(
                "Não foi possível abrir o tutorial",
                f"Não foi possível abrir o navegador padrão.\n\nDetalhes: {exc}",
            )

    def build_tutoriais(self):
        """Central de treinamento com acesso direto aos vídeos publicados no YouTube."""
        tab = self.tab_tutoriais
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(0, weight=1)

        scroll = ctk.CTkScrollableFrame(tab, corner_radius=13)
        scroll.grid(row=0, column=0, padx=6, pady=6, sticky="nsew")
        scroll.grid_columnconfigure(0, weight=1)

        header = ctk.CTkFrame(
            scroll,
            corner_radius=14,
            border_width=1,
            border_color=("#d8e0ea", "#334155"),
            fg_color=("#ffffff", "#111827"),
        )
        header.grid(row=0, column=0, padx=4, pady=(4, 10), sticky="ew")
        header.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            header,
            text="Tutoriais em Vídeo",
            font=("Segoe UI Semibold", 22),
            anchor="w",
        ).grid(row=0, column=0, padx=18, pady=(15, 2), sticky="ew")
        ctk.CTkLabel(
            header,
            text=(
                "Central de treinamento do Controle de Folgas e Banco de Horas. "
                "Clique no vídeo desejado para abri-lo no navegador padrão do Windows."
            ),
            font=("Segoe UI", 11),
            text_color=("#475569", "#cbd5e1"),
            justify="left",
            anchor="w",
            wraplength=920,
        ).grid(row=1, column=0, padx=18, pady=(0, 12), sticky="ew")

        if PLAYLIST_YOUTUBE_URL.strip():
            ctk.CTkButton(
                header,
                text="▶  Assistir à playlist completa",
                command=lambda: self.open_tutorial_link(PLAYLIST_YOUTUBE_URL),
                height=36,
                corner_radius=11,
            ).grid(row=2, column=0, padx=18, pady=(0, 15), sticky="w")

        cards = ctk.CTkFrame(scroll, fg_color="transparent")
        cards.grid(row=1, column=0, padx=4, pady=(0, 8), sticky="ew")
        cards.grid_columnconfigure((0, 1), weight=1, uniform="tutorial_cards")

        for index, tutorial in enumerate(TUTORIAIS_YOUTUBE):
            row = index // 2
            col = index % 2
            card = ctk.CTkFrame(
                cards,
                corner_radius=14,
                border_width=1,
                border_color=("#d8e0ea", "#334155"),
                fg_color=("#ffffff", "#111827"),
            )
            card.grid(row=row, column=col, padx=5, pady=5, sticky="nsew")
            card.grid_columnconfigure(0, weight=1)

            ctk.CTkLabel(
                card,
                text=tutorial["numero"].upper(),
                font=("Segoe UI Semibold", 10),
                text_color=("#2563eb", "#60a5fa"),
                anchor="w",
            ).grid(row=0, column=0, padx=16, pady=(14, 2), sticky="ew")

            ctk.CTkLabel(
                card,
                text=tutorial["titulo"],
                font=("Segoe UI Semibold", 16),
                anchor="w",
                justify="left",
                wraplength=480,
            ).grid(row=1, column=0, padx=16, pady=(0, 5), sticky="ew")

            ctk.CTkLabel(
                card,
                text=tutorial["descricao"],
                font=("Segoe UI", 11),
                text_color=("#64748b", "#94a3b8"),
                anchor="w",
                justify="left",
                wraplength=480,
            ).grid(row=2, column=0, padx=16, pady=(0, 12), sticky="ew")

            button_frame = ctk.CTkFrame(card, fg_color="transparent")
            button_frame.grid(row=3, column=0, padx=16, pady=(0, 15), sticky="ew")
            videos = tutorial["videos"]
            for button_col in range(len(videos)):
                button_frame.grid_columnconfigure(button_col, weight=1)

            for button_col, (button_text, video_url) in enumerate(videos):
                ctk.CTkButton(
                    button_frame,
                    text=f"▶  {button_text}",
                    command=lambda url=video_url: self.open_tutorial_link(url),
                    height=36,
                    corner_radius=10,
                ).grid(
                    row=0,
                    column=button_col,
                    padx=(0 if button_col == 0 else 4, 0 if button_col == len(videos) - 1 else 4),
                    sticky="ew",
                )

        ctk.CTkLabel(
            scroll,
            text="Os vídeos são abertos no YouTube e requerem conexão com a internet.",
            font=("Segoe UI", 10),
            text_color=("#64748b", "#94a3b8"),
        ).grid(row=2, column=0, padx=8, pady=(4, 12), sticky="w")

    def build_ajuda(self):
        tab = self.tab_ajuda
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(0, weight=1)
        frame = ctk.CTkScrollableFrame(tab, corner_radius=13)
        frame.grid(row=0, column=0, padx=6, pady=6, sticky="nsew")
        manual = f"""MANUAL COMPLETO DO SISTEMA
Versão atual: {APP_VERSION}

1. DASHBOARD
Apresenta o saldo geral, créditos, débitos, quantidade de lançamentos e o resumo individual de cada funcionário. Registros enviados para a Lixeira não entram nos cálculos.

2. CADASTRO DE FUNCIONÁRIOS
Na aba Funcionários, use a lista central para localizar e selecionar pessoas. Clique em + Novo funcionário para cadastrar. Informe o nome e escolha o regime de trabalho. Nos regimes padronizados, a jornada diária é definida automaticamente; somente a opção Outro pede a escolha manual de 6, 8 ou 12 horas. Todo novo funcionário é cadastrado como Ativo. Ao concluir, o sistema mostra uma confirmação com OK.

3. JORNADA INDIVIDUAL
O campo Dias utiliza a jornada do funcionário selecionado. Para um funcionário de 6 horas, 1 dia equivale a 6 horas; para 8 horas, equivale a 8; e para 12x36, equivale a 12 horas. Valores informados em Horas e Minutos são sempre somados diretamente.

4. LANÇAMENTOS
Na aba Lançamentos, clique em + Novo lançamento. Selecione o funcionário, a data e obrigatoriamente uma categoria válida. Enquanto nenhuma categoria for escolhida, o sistema não presume crédito nem débito. Na maioria das categorias, o sistema determina automaticamente se o lançamento aumenta ou diminui o saldo e mostra essa informação de forma visual. Nas categorias Banco de horas, Exame médico, Troca de plantão e Outro, escolha manualmente se o lançamento aumenta ou diminui o saldo. Preencha Dias, Horas e Minutos somente com números. A Observação é obrigatória. Ao concluir, o sistema mostra uma confirmação com OK.

5. ALTERAR UM LANÇAMENTO
Selecione o registro na tabela e clique em Editar selecionado ou dê duplo clique na linha. A edição abre em uma janela própria, separada do cadastro. Ajuste os campos e clique em Salvar alterações. Antes da alteração, o sistema cria automaticamente uma cópia do banco na pasta Backups e registra na auditoria somente os campos que realmente mudaram, mostrando o valor anterior e o novo.

6. OPÇÕES ADMINISTRATIVAS
Na aba Lançamentos, o botão Opções administrativas reúne as funções protegidas. Para facilitar o uso, o botão Enviar para a Lixeira também fica disponível na barra de ações acima da tabela, sem retirar a exigência da senha administrativa.

7. PRIMEIRA SENHA ADMINISTRATIVA
No primeiro acesso às opções administrativas, o sistema solicita a criação de uma senha com pelo menos 8 caracteres e gera um código de recuperação. O código é mostrado somente naquele momento e deve ser guardado em local seguro.

8. EXCLUIR UM LANÇAMENTO
Selecione o lançamento e clique em Enviar para a Lixeira, diretamente na tela de Lançamentos ou pelas Opções administrativas. A senha administrativa, o motivo da exclusão e uma confirmação detalhada são obrigatórios. Um backup é criado antes da operação.

9. LIXEIRA
Os registros enviados para a Lixeira deixam de aparecer nas telas, nos saldos e nos relatórios, mas permanecem no banco de dados. A senha administrativa pode restaurar um registro. Para excluir definitivamente um item ou esvaziar a Lixeira, é necessária uma resposta assinada pelo suporte para a operação específica, sem uso de senha mestra. Antes de qualquer exclusão definitiva é criado um backup. Dê dois cliques em uma linha da Lixeira para abrir todos os detalhes do registro.

10. ALTERAR A SENHA
Use Opções administrativas > Alterar senha administrativa. A senha atual é solicitada. Depois da alteração, um novo código de recuperação é gerado e o código anterior deixa de funcionar.

11. RECUPERAR O ACESSO
Use Opções administrativas > Recuperar acesso e informe o código local, ou, na solicitação de senha administrativa, escolha Solicitar recuperação ao suporte. Nesse caso, copie a solicitação exibida e obtenha uma resposta assinada para redefinir a senha. Uma nova senha invalida o código de recuperação anterior.

12. BLOQUEIO DE SEGURANÇA
Após 5 tentativas de senha incorreta, o acesso por senha fica bloqueado por 5 minutos. O código de recuperação continua disponível para o responsável autorizado.

13. AUDITORIA
A auditoria registra data, horário, ação, identificação do lançamento e detalhes. Alterações de lançamentos e funcionários também são registradas automaticamente, mostrando os valores anteriores e novos. Dê dois cliques em qualquer linha para abrir os detalhes completos do registro de auditoria.

14. ALTERAR UM FUNCIONÁRIO
Selecione o funcionário na lista e clique em Editar funcionário selecionado ou dê duplo clique na linha. A edição abre em uma janela própria. Ajuste nome, regime, jornada quando aplicável ou situação Ativo/Inativo e clique em Salvar alterações. Um backup é criado antes de alterar um cadastro existente.

15. FILTROS E PESQUISA
Na aba Lançamentos, filtre por funcionário, ano e mês. O campo de busca pesquisa categoria e observação.

16. SALDOS
Saldo positivo indica que tem horas para tirar. Saldo negativo indica que está devendo horas. A conversão exibida em dias respeita a jornada individual do funcionário sempre que houver um funcionário específico.

17. EXPORTAÇÃO PARA EXCEL
O botão Exportar Excel gera as planilhas Lançamentos, Resumo e as páginas individuais de assinatura. Os filtros aplicados na tela também são usados no relatório. Itens da Lixeira não são exportados.

18. PLANILHA DE ASSINATURA
A página individual contém resumo, movimentos do período e dois espaços separados: assinatura do funcionário e assinatura do responsável. As linhas de grade são ocultadas para melhorar a impressão.

19. IMPRESSÃO DIRETA DAS ASSINATURAS
O botão Imprimir para Assinaturas permite escolher um funcionário ou todos os funcionários ativos, além do mês, do ano e de 1 ou 2 cópias por funcionário. Cada funcionário é enviado separadamente à impressora, de modo que a paginação recomeça corretamente em cada relatório (por exemplo, Página 1 de 1 ou Página 1 de 2 / Página 2 de 2). Quando são escolhidas 2 cópias, o sistema envia duas impressões independentes para evitar que drivers de impressora ignorem a quantidade solicitada.

20. TEMA E CORES DO DASHBOARD
O botão Tema alterna toda a interface entre claro e escuro. A escolha é salva automaticamente no arquivo controle_folgas_config.json, na mesma pasta do programa, e é restaurada na próxima abertura. No Dashboard, funcionários com horas para tirar aparecem em verde, funcionários devendo horas aparecem em vermelho e funcionários zerados mantêm a aparência normal da tabela. As tonalidades são adaptadas automaticamente aos temas claro e escuro.

21. BACKUPS
Os backups SQLite consistentes são armazenados na pasta Backups, ao lado do programa. São criados automaticamente antes de alterar lançamentos, alterar funcionários e executar ações na Lixeira. O sistema mantém os 50 arquivos mais recentes.

22. BANCO DE DADOS
Os dados, hashes da senha e do código de recuperação, solicitações temporárias de suporte, Lixeira e auditoria ficam no arquivo controle_folgas.db. A chave PRIVADA de suporte não é distribuída com o programa. Não mova ou apague o banco sem possuir uma cópia de segurança.

23. ATUALIZAÇÕES AUTOMÁTICAS
Ao iniciar, o programa consulta silenciosamente a versão mais recente publicada no GitHub Releases. Quando existe uma versão nova, é oferecida a opção de baixar e instalar. O arquivo baixado é validado por SHA-256 antes da instalação. A atualização substitui somente o executável; controle_folgas.db, controle_folgas_config.json, Backups e Status_Programa permanecem na mesma pasta. Também é possível usar o botão Verificar atualizações nesta aba.
"""
        ctk.CTkLabel(frame, text="Ajuda e Manual do Sistema", font=("Segoe UI Semibold", 20)).pack(anchor="w", padx=16, pady=(14, 5))

        update_card = ctk.CTkFrame(
            frame,
            corner_radius=13,
            border_width=1,
            border_color=("#d8e0ea", "#334155"),
            fg_color=("#ffffff", "#111827"),
        )
        update_card.pack(fill="x", padx=16, pady=(4, 12))
        update_card.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            update_card,
            text="Atualizações do programa",
            font=("Segoe UI Semibold", 14),
            anchor="w",
        ).grid(row=0, column=0, padx=14, pady=(12, 2), sticky="ew")
        self.update_status_label = ctk.CTkLabel(
            update_card,
            text=f"Versão instalada: {APP_VERSION} • verificação automática ao iniciar",
            font=("Segoe UI", 11),
            text_color=("#64748b", "#94a3b8"),
            anchor="w",
        )
        self.update_status_label.grid(row=1, column=0, padx=14, pady=(0, 12), sticky="ew")
        ctk.CTkButton(
            update_card,
            text="Verificar atualizações",
            command=lambda: self.check_for_updates(silent=False),
            width=165,
            height=34,
            corner_radius=10,
        ).grid(row=0, column=1, rowspan=2, padx=14, pady=12, sticky="e")
        ctk.CTkLabel(
            frame,
            text="Prefere assistir? Acesse a aba Tutoriais para abrir os vídeos passo a passo no YouTube.",
            justify="left",
            anchor="w",
            wraplength=1050,
            font=("Segoe UI Semibold", 12),
            text_color=("#2563eb", "#60a5fa"),
        ).pack(fill="x", padx=16, pady=(0, 10))
        ctk.CTkLabel(frame, text=manual, justify="left", anchor="w", wraplength=1050, font=("Segoe UI", 13)).pack(fill="x", padx=16, pady=(0, 16))

    def refresh_all(self):
        self.refresh_employees()
        self.refresh_entries()
        self.refresh_dashboard()

    def refresh_employees(self):
        employees = list(self.db.employees())
        all_names = ["Todos"] + [e["nome"] for e in employees]

        if hasattr(self, "combo_filter_func"):
            self.combo_filter_func.configure(values=all_names)
            if self.combo_filter_func.get() not in all_names:
                self.combo_filter_func.set("Todos")

        search = ""
        if hasattr(self, "emp_search_var"):
            search = normalize_spaces(self.emp_search_var.get()).lower()
        visible_employees = [
            e for e in employees
            if not search or search in normalize_spaces(e["nome"]).lower()
        ]

        if hasattr(self, "employee_count_label"):
            active_count = sum(1 for e in employees if int(e["ativo"] or 0))
            inactive_count = len(employees) - active_count
            if search:
                found = len(visible_employees)
                word = "encontrado" if found == 1 else "encontrados"
                self.employee_count_label.configure(
                    text=f"{found} {word} de {len(employees)} cadastrados"
                )
            else:
                self.employee_count_label.configure(
                    text=f"{len(employees)} cadastrados • {active_count} ativos • {inactive_count} inativos"
                )

        scroll = getattr(self, "employee_scroll", None)
        if scroll is None:
            return

        previous_selection = self.selected_employee_id
        visible_ids = {int(e["id"]) for e in visible_employees}
        if previous_selection not in visible_ids:
            self.selected_employee_id = None

        for child in scroll.winfo_children():
            child.destroy()
        self.employee_row_widgets = {}

        if not visible_employees:
            self.employee_empty_label = ctk.CTkLabel(
                scroll,
                text="Nenhum funcionário encontrado." if search else "Nenhum funcionário cadastrado.",
                font=("Segoe UI", 13),
                text_color=("#64748b", "#94a3b8"),
            )
            self.employee_empty_label.grid(row=0, column=0, padx=20, pady=48)
            return

        for index, employee in enumerate(visible_employees):
            employee_id = int(employee["id"])
            active = bool(int(employee["ativo"] or 0))
            base_bg, text_color = self._employee_row_colors(active=active, zebra=bool(index % 2))

            row = ctk.CTkFrame(
                scroll,
                height=56,
                corner_radius=9,
                fg_color=base_bg,
                border_width=0,
            )
            row.grid(row=index, column=0, padx=(7, 4), pady=2, sticky="ew")
            row.grid_propagate(False)
            row.grid_columnconfigure(0, minsize=5)
            row.grid_columnconfigure(1, weight=1)

            accent = ctk.CTkFrame(row, width=4, corner_radius=2, fg_color="transparent")
            accent.grid(row=0, column=0, padx=(4, 0), pady=7, sticky="ns")

            content = ctk.CTkFrame(row, fg_color="transparent")
            content.grid(row=0, column=1, padx=(7, 6), sticky="nsew")
            self._configure_employee_grid(content)

            name_label = ctk.CTkLabel(
                content,
                text=employee["nome"],
                font=("Segoe UI Semibold", 14),
                text_color=text_color,
                anchor="w",
            )
            name_label.grid(row=0, column=0, padx=(8, 8), pady=10, sticky="ew")

            regime_label = ctk.CTkLabel(
                content,
                text=employee["regime"],
                font=("Segoe UI", 13),
                text_color=text_color,
                anchor="w",
            )
            regime_label.grid(row=0, column=1, padx=8, pady=10, sticky="ew")

            jornada_label = ctk.CTkLabel(
                content,
                text=f"{int(employee['jornada_horas'])}h",
                font=("Segoe UI Semibold", 13),
                text_color=text_color,
                anchor="center",
            )
            jornada_label.grid(row=0, column=2, padx=8, pady=10, sticky="ew")

            if active:
                status_bg = ("#e8f1ff", "#1d3655")
                status_fg = ("#1d4ed8", "#bfdbfe")
                status_text = "Ativo"
            else:
                status_bg = ("#eef2f7", "#222b38")
                status_fg = ("#64748b", "#94a3b8")
                status_text = "Inativo"

            status_label = ctk.CTkLabel(
                content,
                text=status_text,
                width=88,
                height=26,
                corner_radius=13,
                fg_color=status_bg,
                text_color=status_fg,
                font=("Segoe UI Semibold", 13),
                anchor="center",
            )
            status_label.grid(row=0, column=3, padx=(8, 14), pady=8)

            self.employee_row_widgets[employee_id] = {
                "row": row,
                "accent": accent,
                "base_bg": base_bg,
            }

            for widget in (row, accent, content, name_label, regime_label, jornada_label, status_label):
                self._bind_employee_row_widget(widget, employee_id)

        if self.selected_employee_id in self.employee_row_widgets:
            self.select_employee_card(self.selected_employee_id)

    def get_current_filters(self):
        employee_id = None
        selected = self.combo_filter_func.get() if hasattr(self, "combo_filter_func") else "Todos"
        if selected and selected != "Todos":
            emp = self.db.employee_by_name(selected)
            if emp:
                employee_id = emp["id"]
        year = self.combo_filter_year.get() if hasattr(self, "combo_filter_year") else "Todos"
        year = year or "Todos"
        month = self.combo_filter_month.get() if hasattr(self, "combo_filter_month") else "Todos"
        month = month_value(month)
        search = self.entry_search.get().strip() if hasattr(self, "entry_search") else ""
        return employee_id, selected or "Todos", year, month, search

    def employee_running_balances(self, rows):
        """Retorna o saldo real do funcionário após cada lançamento exibido."""
        if not rows:
            return {}
        employee_ids = sorted({r["funcionario_id"] for r in rows})
        balances = {}
        for emp_id in employee_ids:
            emp_rows = self.db.conn.execute(
                """
                SELECT id, minutos
                FROM lancamentos
                WHERE funcionario_id=? AND excluido=0
                ORDER BY data ASC, id ASC
                """,
                (emp_id,),
            ).fetchall()
            saldo = 0
            hpd = self.db.employee_hours(emp_id)
            for item in emp_rows:
                saldo += int(item["minutos"] or 0)
                balances[item["id"]] = format_minutes(saldo, hpd)
        return balances

    def status_saldo(self, saldo_minutos: int, nome: str, hours_per_day: int = DEFAULT_HOURS_PER_DAY):
        hpd = hours_per_day
        saldo_minutos = int(saldo_minutos or 0)
        nome = nome or "Todos os funcionários"
        if saldo_minutos > 0:
            return (
                f"✅ {nome} tem {format_minutes(saldo_minutos, hpd)} para tirar de folga",
                "#166534",
                "#dcfce7",
            )
        if saldo_minutos < 0:
            return (
                f"⚠️ {nome} está devendo {format_minutes(abs(saldo_minutos), hpd)}",
                "#991b1b",
                "#fee2e2",
            )
        return (
            f"✅ {nome} está com o banco de horas zerado",
            "#1e40af",
            "#dbeafe",
        )

    def atualizar_painel_status(self, employee_id, selected_name):
        if not hasattr(self, "label_status_funcionario"):
            return
        nome = selected_name if selected_name and selected_name != "Todos" else "Todos os funcionários"
        saldo_real = self.db.saldo(employee_id) if employee_id else self.db.saldo()
        hpd = self.db.employee_hours(employee_id) if employee_id else DEFAULT_HOURS_PER_DAY
        texto, cor_texto, cor_fundo = self.status_saldo(saldo_real, nome, hpd)
        self.status_box.configure(fg_color=cor_fundo)
        self.label_status_funcionario.configure(text=texto, text_color=cor_texto)

    def refresh_entries(self):
        self.selected_entry_id = None
        for item in self.tree_entries.get_children():
            self.tree_entries.delete(item)
        employee_id, selected_name, year, month, search = self.get_current_filters()
        rows = self.db.entries(funcionario_id=employee_id, year=year, month=month, search=search)
        balances = self.employee_running_balances(rows)
        hpd = self.db.employee_hours(employee_id) if employee_id else DEFAULT_HOURS_PER_DAY
        creditos = sum(int(r["minutos"] or 0) for r in rows if r["minutos"] > 0)
        debitos = sum(int(r["minutos"] or 0) for r in rows if r["minutos"] < 0)
        saldo = creditos + debitos
        self.atualizar_painel_status(employee_id, selected_name)
        if hasattr(self, "label_saldo_filtro"):
            filtro_nome = selected_name if selected_name != "Todos" else "todos os funcionários"
            self.label_saldo_filtro.configure(
                text=(
                    f"Saldo filtrado ({filtro_nome} - {period_label(year, month)}): {format_minutes(saldo, hpd)} | "
                    f"Horas extras/créditos: {format_minutes(creditos, hpd)} | "
                    f"Folgas/débitos abatidos: {format_minutes(debitos, hpd)}"
                )
            )
        for r in rows:
            movimento = normalize_spaces(r["movimento"])
            row_tags = ("movimento_debito",) if movimento.startswith(MOVIMENTO_DEBITO) else ("movimento_credito",)
            self.tree_entries.insert(
                "",
                "end",
                values=(
                    r["id"],
                    br_date(r["data"]),
                    r["funcionario"],
                    r["movimento"],
                    r["categoria"],
                    r["quantidade_texto"] or "",
                    format_minutes(r["minutos"], int(r["jornada_horas"] if "jornada_horas" in r.keys() else self.db.employee_hours(r["funcionario_id"]))),
                    balances.get(r["id"], ""),
                    r["observacao"] or "",
                ),
                tags=row_tags,
            )
        self.adjust_entries_columns()

        years = ["Todos"] + self.db.years()
        current = self.combo_filter_year.get()
        self.combo_filter_year.configure(values=years)
        if current not in years:
            self.combo_filter_year.set("Todos")

    def refresh_dashboard(self):
        hpd = DEFAULT_HOURS_PER_DAY
        totals = self.db.totals()
        self.card_saldo.configure(text=format_minutes(totals["saldo"], hpd))
        self.card_creditos.configure(text=format_minutes(totals["creditos"], hpd))
        self.card_debitos.configure(text=format_minutes(totals["debitos"], hpd))
        self.card_lancamentos.configure(text=str(totals["lancamentos"]))

        for item in self.tree_resumo.get_children():
            self.tree_resumo.delete(item)

        for index, emp in enumerate(self.db.employees()):
            rows = self.db.entries(funcionario_id=emp["id"])
            creditos = sum(int(r["minutos"] or 0) for r in rows if r["minutos"] > 0)
            debitos = sum(int(r["minutos"] or 0) for r in rows if r["minutos"] < 0)
            saldo = creditos + debitos
            hpd_emp = int(emp["jornada_horas"] or DEFAULT_HOURS_PER_DAY)

            if saldo > 0:
                status = "Tem horas para tirar"
                saldo_txt = "+" + format_minutes(saldo, hpd_emp)
                row_tags = ("saldo_positivo",)
            elif saldo < 0:
                status = "Está devendo horas"
                saldo_txt = format_minutes(saldo, hpd_emp)
                row_tags = ("saldo_negativo",)
            else:
                status = "Zerado"
                saldo_txt = "0"
                row_tags = ("zebra",) if index % 2 else ()

            self.tree_resumo.insert(
                "", "end",
                values=(
                    emp["nome"],
                    status,
                    saldo_txt,
                    format_minutes(creditos, hpd_emp),
                    format_minutes(debitos, hpd_emp),
                    len(rows),
                ),
                tags=row_tags,
            )
        self.after_idle(self.adjust_dashboard_columns)

    @staticmethod
    def regime_hours(regime: str):
        mapping = {
            "12x36": 12,
            "Segunda a sexta - 8h": 8,
            "Segunda a sexta - 6h": 6,
            "Segunda a sábado - 8h": 8,
            "Segunda a sábado - 6h": 6,
        }
        return mapping.get(normalize_spaces(regime))

    @staticmethod
    def _ctk_window_scale(window) -> float:
        """Retorna o fator de escala da janela CTk (1.0, 1.25, 1.5...)."""
        try:
            value = float(window._get_window_scaling())
            return value if value > 0 else 1.0
        except Exception:
            return 1.0

    @staticmethod
    def _ctk_widget_scale(widget) -> float:
        """Retorna o fator de escala do widget CTk."""
        try:
            value = float(widget._get_widget_scaling())
            return value if value > 0 else 1.0
        except Exception:
            return 1.0

    def logical_work_area(self, window=None):
        """Converte a área útil física do Windows para unidades lógicas do CTk.

        CustomTkinter multiplica largura/altura passadas para geometry() pelo
        fator de DPI. A API do Windows, por outro lado, informa pixels físicos.
        Trabalhar sempre nesta conversão evita o efeito de escala duplicada.
        """
        target = window if window is not None else self
        work_w, work_h, work_x, work_y = self.get_work_area(target)
        scale = self._ctk_window_scale(target)
        logical_w = max(1, int(work_w / scale))
        logical_h = max(1, int(work_h / scale))
        return logical_w, logical_h, work_x, work_y, scale

    def modal_limits(self, window, margin_x: int = 44, margin_y: int = 52):
        """Retorna o máximo lógico que a janela pode ocupar sem sair da tela."""
        work_w, work_h, _work_x, _work_y = self.get_work_area(window)
        scale = self._ctk_window_scale(window)
        max_w = max(220, int(max(220, work_w - margin_x) / scale))
        max_h = max(220, int(max(220, work_h - margin_y) / scale))
        return max_w, max_h

    def responsive_minsize(self, window, width: int, height: int, floor_w: int = 320, floor_h: int = 240):
        """Aplica minsize sem deixar o mínimo ultrapassar o monitor atual."""
        try:
            max_w, max_h = self.modal_limits(window)
            min_w = min(int(width), max_w)
            min_h = min(int(height), max_h)
            min_w = max(min(floor_w, max_w), min_w)
            min_h = max(min(floor_h, max_h), min_h)
            window.minsize(min_w, min_h)
        except Exception:
            pass

    def center_modal(self, window, width: int, height: int):
        """Centraliza corretamente uma CTkToplevel usando tamanho lógico e posição física."""
        try:
            window.update_idletasks()
        except Exception:
            pass

        scale = self._ctk_window_scale(window)
        physical_w = max(1, int(round(width * scale)))
        physical_h = max(1, int(round(height * scale)))
        work_w, work_h, work_x, work_y = self.get_work_area(window)

        try:
            parent = window.master if window.master is not None else self
            parent_x = int(parent.winfo_rootx())
            parent_y = int(parent.winfo_rooty())
            parent_w = int(parent.winfo_width())
            parent_h = int(parent.winfo_height())
            x = parent_x + (parent_w - physical_w) // 2
            y = parent_y + (parent_h - physical_h) // 2
        except Exception:
            x = work_x + (work_w - physical_w) // 2
            y = work_y + (work_h - physical_h) // 2

        # Nunca permitir que a janela seja centralizada para fora da área útil.
        x = max(work_x, min(x, work_x + max(0, work_w - physical_w)))
        y = max(work_y, min(y, work_y + max(0, work_h - physical_h)))
        window.geometry(f"{int(width)}x{int(height)}+{int(x)}+{int(y)}")

    def resize_modal_to_work_area(self, window, width: int, height: int, margin_x: int = 44, margin_y: int = 52):
        """Limita a janela à área útil considerando corretamente o DPI do CTk."""
        try:
            max_w, max_h = self.modal_limits(window, margin_x=margin_x, margin_y=margin_y)
            final_w = max(220, min(int(width), max_w))
            final_h = max(220, min(int(height), max_h))
            self.center_modal(window, final_w, final_h)
            try:
                window.maxsize(max_w, max_h)
            except Exception:
                pass
            return final_w, final_h
        except Exception:
            self.center_modal(window, int(width), int(height))
            return int(width), int(height)

    def keep_window_inside_work_area(self, window):
        """Revalida posição/tamanho após o Windows concluir o cálculo de DPI."""
        try:
            if not window.winfo_exists():
                return
            geometry = window.geometry()
            match = re.match(r"^(\d+)x(\d+)", geometry)
            if not match:
                return
            width, height = map(int, match.groups())
            self.resize_modal_to_work_area(window, width, height)
        except Exception:
            pass

    def fit_scrollable_form_modal(
        self,
        window,
        scroll_frame,
        content_frame,
        preferred_width: int,
        min_body_height: int = 180,
        max_body_height: int = 520,
    ):
        """Dimensiona formulários roláveis mantendo cabeçalho e botões visíveis.

        winfo_reqheight()/winfo_reqwidth() retornam pixels já escalados. Antes de
        reaproveitar essas medidas em height=/geometry(), convertemos novamente
        para unidades lógicas. Isso é essencial em 125%, 150% e escalas maiores.
        """
        try:
            window.update_idletasks()

            widget_scale = self._ctk_widget_scale(content_frame)
            window_scale = self._ctk_window_scale(window)

            content_height = max(
                min_body_height,
                int(content_frame.winfo_reqheight() / widget_scale) + 12,
            )
            desired_body = min(content_height, max_body_height)
            scroll_frame.configure(height=desired_body)
            window.update_idletasks()

            _max_w, max_window_h = self.modal_limits(window)
            requested_h = int(window.winfo_reqheight() / window_scale)

            if requested_h > max_window_h:
                overflow = requested_h - max_window_h
                reduced_body = max(110, desired_body - overflow - 10)
                scroll_frame.configure(height=reduced_body)
                window.update_idletasks()

            requested_w = max(
                preferred_width,
                int(window.winfo_reqwidth() / window_scale),
            )
            requested_h = int(window.winfo_reqheight() / window_scale)
            self.resize_modal_to_work_area(window, requested_w, requested_h)
        except Exception:
            self.resize_modal_to_work_area(window, preferred_width, 620)

    def _show_branded_message(self, title, message, kind="info", parent=None, yes_no=False):
        """Substitui messagebox por uma caixa CTk com logo e ícone do programa."""
        parent = parent if parent is not None else self
        message = str(message or "")

        # Se a chamada vier de outra janela modal, restaura o grab dela quando
        # esta mensagem for fechada para não deixar a interface sem modal ativo.
        previous_grab = None
        try:
            previous_grab = self.grab_current()
        except Exception:
            pass

        estimated_lines = sum(
            max(1, (len(line) + 67) // 68)
            for line in (message.splitlines() or [""])
        )
        width = 620
        height = min(540, max(285, 225 + estimated_lines * 20))
        window = self.create_modal(
            title,
            width,
            height,
            resizable=False,
            parent=parent,
            grab=True,
        )
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(0, weight=1)

        kind_config = {
            "info": ("Informação", "#2563eb", "#1d4ed8"),
            "warning": ("Atenção", "#d97706", "#b45309"),
            "error": ("Erro", "#dc2626", "#b91c1c"),
            "question": ("Confirmação", "#0f766e", "#115e59"),
        }
        heading, accent, hover = kind_config.get(kind, kind_config["info"])

        card = ctk.CTkFrame(
            window,
            corner_radius=16,
            border_width=1,
            border_color=("#d8e0ea", "#334155"),
            fg_color=("#ffffff", "#111827"),
        )
        card.grid(row=0, column=0, padx=18, pady=(18, 10), sticky="nsew")
        card.grid_columnconfigure(1, weight=1)
        card.grid_rowconfigure(1, weight=1)

        logo_image = self.get_logo_image(92, 92)
        if logo_image is not None:
            ctk.CTkLabel(
                card,
                text="",
                image=logo_image,
                width=96,
                height=96,
            ).grid(row=0, column=0, rowspan=2, padx=(18, 12), pady=18, sticky="n")

        ctk.CTkLabel(
            card,
            text=heading,
            font=("Segoe UI Semibold", 18),
            text_color=accent,
            anchor="w",
        ).grid(row=0, column=1, padx=(4, 20), pady=(22, 7), sticky="ew")

        ctk.CTkLabel(
            card,
            text=message,
            font=("Segoe UI", 12),
            justify="left",
            anchor="nw",
            wraplength=445,
            text_color=("#334155", "#e5e7eb"),
        ).grid(row=1, column=1, padx=(4, 20), pady=(0, 20), sticky="nsew")

        result = {"value": False}

        def close(value=False):
            result["value"] = bool(value)
            try:
                window.grab_release()
            except Exception:
                pass
            window.destroy()

        buttons = ctk.CTkFrame(window, fg_color="transparent")
        buttons.grid(row=1, column=0, padx=18, pady=(0, 18), sticky="ew")

        if yes_no:
            buttons.grid_columnconfigure((0, 1), weight=1)
            ctk.CTkButton(
                buttons,
                text="Não",
                command=lambda: close(False),
                height=38,
                fg_color="#475569",
                hover_color="#334155",
            ).grid(row=0, column=0, padx=(0, 6), sticky="ew")
            ctk.CTkButton(
                buttons,
                text="Sim",
                command=lambda: close(True),
                height=38,
                fg_color=accent,
                hover_color=hover,
            ).grid(row=0, column=1, padx=(6, 0), sticky="ew")
            window.bind("<Return>", lambda _event: close(True))
        else:
            buttons.grid_columnconfigure(0, weight=1)
            ctk.CTkButton(
                buttons,
                text="OK",
                command=lambda: close(True),
                height=38,
                fg_color=accent,
                hover_color=hover,
            ).grid(row=0, column=0, sticky="ew")
            window.bind("<Return>", lambda _event: close(True))

        window.bind("<Escape>", lambda _event: close(False))
        window.protocol("WM_DELETE_WINDOW", lambda: close(False))
        self.wait_window(window)

        if previous_grab is not None:
            try:
                if previous_grab.winfo_exists():
                    previous_grab.grab_set()
                    previous_grab.focus_force()
            except Exception:
                pass

        return result["value"]

    def show_info(self, title, message, parent=None):
        self._show_branded_message(title, message, kind="info", parent=parent)
        return "ok"

    def show_warning(self, title, message, parent=None):
        self._show_branded_message(title, message, kind="warning", parent=parent)
        return "ok"

    def show_error(self, title, message, parent=None):
        self._show_branded_message(title, message, kind="error", parent=parent)
        return "ok"

    def ask_yes_no(self, title, message, parent=None):
        return self._show_branded_message(
            title,
            message,
            kind="question",
            parent=parent,
            yes_no=True,
        )

    def create_modal(
        self,
        title: str,
        width: int,
        height: int,
        resizable=False,
        parent=None,
        grab=True,
    ):
        """Cria CTkToplevel responsiva e segura para qualquer escala do Windows."""
        parent = parent if parent is not None else self
        window = ctk.CTkToplevel(parent)
        window.title(title)
        apply_window_icon(window)
        window.transient(parent)
        window.resizable(bool(resizable), bool(resizable))
        self.resize_modal_to_work_area(window, width, height)
        if grab:
            window.grab_set()
        try:
            window.focus_force()
        except Exception:
            pass

        # A escala de um CTkToplevel pode ser atualizada alguns milissegundos
        # depois da criação (principalmente RDP, múltiplos monitores e 125/150%).
        window.after(80, lambda w=window: self.keep_window_inside_work_area(w))
        window.after(260, lambda w=window: self.keep_window_inside_work_area(w))
        return window

    def open_new_employee(self):
        self.open_employee_form(employee_id=None)

    def edit_selected_employee(self):
        if not self.selected_employee_id:
            self.show_warning(
                "Selecione um funcionário",
                "Selecione na lista o funcionário que deseja editar.",
            )
            return
        self.open_employee_form(employee_id=self.selected_employee_id)

    def open_employee_form(self, employee_id=None):
        editing = employee_id is not None
        employee = self.db.employee_by_id(employee_id) if editing else None
        if editing and not employee:
            self.show_warning("Funcionário não encontrado", "O funcionário selecionado não está mais disponível.")
            self.refresh_employees()
            return

        title = "Editar funcionário" if editing else "Novo funcionário"
        window = self.create_modal(title, 580, 600)
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(2, weight=1)

        ctk.CTkLabel(
            window,
            text=title if not editing else f"Editar funcionário — {employee['nome']}",
            font=("Segoe UI Semibold", 21),
        ).grid(row=0, column=0, padx=26, pady=(22, 5), sticky="w")
        ctk.CTkLabel(
            window,
            text=(
                "Preencha somente as informações abaixo. A jornada é definida automaticamente pelo regime."
                if not editing else
                "Altere apenas o que for necessário e clique em Salvar alterações."
            ),
            font=("Segoe UI", 11),
            text_color=("#64748b", "#94a3b8"),
            wraplength=520,
            justify="left",
        ).grid(row=1, column=0, padx=26, pady=(0, 12), sticky="w")

        # Corpo rolável + botões externos: mesmo em resoluções menores ou com
        # escala de 125%/150%, Cancelar e Salvar nunca ficam escondidos.
        form_scroll = ctk.CTkScrollableFrame(window, corner_radius=13, height=300)
        form_scroll.grid(row=2, column=0, padx=26, sticky="nsew")
        form_scroll.grid_columnconfigure(0, weight=1)
        form = ctk.CTkFrame(form_scroll, fg_color="transparent")
        form.grid(row=0, column=0, sticky="ew")
        form.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(form, text="Nome do funcionário", font=("Segoe UI Semibold", 11)).grid(row=0, column=0, padx=12, pady=(10, 3), sticky="w")
        name_entry = ctk.CTkEntry(form, height=36, corner_radius=10, placeholder_text="Digite o nome completo")
        name_entry.grid(row=1, column=0, padx=12, sticky="ew")

        ctk.CTkLabel(form, text="Regime de trabalho", font=("Segoe UI Semibold", 11)).grid(row=2, column=0, padx=12, pady=(13, 3), sticky="w")
        regime_combo = ctk.CTkComboBox(
            form,
            values=["12x36", "Segunda a sexta - 8h", "Segunda a sexta - 6h", "Segunda a sábado - 8h", "Segunda a sábado - 6h", "Outro"],
            height=36, corner_radius=10, state="readonly",
        )
        regime_combo.grid(row=3, column=0, padx=12, sticky="ew")

        jornada_frame = ctk.CTkFrame(form, fg_color="transparent")
        jornada_frame.grid(row=4, column=0, padx=12, pady=(10, 6), sticky="ew")
        jornada_frame.grid_columnconfigure(0, weight=1)
        jornada_info = ctk.CTkLabel(
            jornada_frame, text="", font=("Segoe UI Semibold", 12),
            anchor="w", justify="left",
        )
        jornada_info.grid(row=0, column=0, sticky="w")
        jornada_combo = ctk.CTkComboBox(
            jornada_frame, values=["6", "8", "12"],
            height=34, width=150, state="readonly",
        )

        status_choice = None
        status_frame = ctk.CTkFrame(form, fg_color="transparent")
        if editing:
            status_frame.grid(row=5, column=0, padx=12, pady=(7, 10), sticky="ew")
            ctk.CTkLabel(status_frame, text="Situação do funcionário", font=("Segoe UI Semibold", 11)).pack(anchor="w", pady=(0, 4))
            status_choice = ctk.CTkSegmentedButton(status_frame, values=["Ativo", "Inativo"], height=34)
            status_choice.pack(fill="x")
            status_choice.set("Ativo" if employee["ativo"] else "Inativo")
            ctk.CTkLabel(
                status_frame,
                text="Inativo preserva todo o histórico, mas deixa o funcionário fora dos novos lançamentos.",
                font=("Segoe UI", 10), text_color=("#64748b", "#94a3b8"),
                wraplength=480, justify="left",
            ).pack(anchor="w", pady=(4, 0))

        def refit_window():
            def apply_fit():
                try:
                    if window.winfo_exists():
                        self.fit_scrollable_form_modal(
                            window, form_scroll, form, preferred_width=580,
                            min_body_height=190, max_body_height=430,
                        )
                except Exception:
                    pass
            window.after_idle(apply_fit)
            window.after(180, apply_fit)

        def refresh_jornada(_value=None):
            regime = regime_combo.get().strip()
            auto_hours = self.regime_hours(regime)
            if auto_hours:
                jornada_combo.grid_remove()
                jornada_combo.set(str(auto_hours))
                jornada_info.configure(
                    text=f"Jornada automática: 1 dia = {auto_hours} horas",
                    text_color=("#166534", "#86efac"),
                )
            else:
                jornada_info.configure(
                    text="Escolha quantas horas equivalem a 1 dia:",
                    text_color=("#334155", "#cbd5e1"),
                )
                jornada_combo.grid(row=1, column=0, pady=(5, 0), sticky="w")
            refit_window()

        regime_combo.configure(command=refresh_jornada)

        if editing:
            name_entry.insert(0, employee["nome"])
            current_regime = employee["regime"] if employee["regime"] in regime_combo.cget("values") else "Outro"
            regime_combo.set(current_regime)
            jornada_combo.set(str(employee["jornada_horas"]))
        else:
            regime_combo.set("12x36")
            jornada_combo.set("12")

        buttons = ctk.CTkFrame(window, fg_color="transparent")
        buttons.grid(row=3, column=0, padx=26, pady=(14, 20), sticky="ew")
        buttons.grid_columnconfigure((0, 1), weight=1)
        ctk.CTkButton(
            buttons, text="Cancelar", command=window.destroy, height=38,
            fg_color="#475569", hover_color="#334155",
        ).grid(row=0, column=0, padx=(0, 6), sticky="ew")

        def submit():
            try:
                nome = normalize_spaces(name_entry.get())
                if not nome:
                    raise ValueError("Informe o nome do funcionário.")
                regime = normalize_spaces(regime_combo.get()) or "Outro"
                auto_hours = self.regime_hours(regime)
                jornada = auto_hours if auto_hours else int(jornada_combo.get())

                if editing:
                    create_database_backup("alteracao_funcionario")
                    changed = self.db.update_employee(
                        employee_id, nome,
                        1 if status_choice.get() == "Ativo" else 0,
                        jornada, regime,
                    )
                    self.refresh_all()
                    window.destroy()
                    if changed:
                        self.show_info("Alterações salvas", "Alterações do funcionário salvas com sucesso.", parent=self)
                    else:
                        self.show_info("Nenhuma alteração", "Nenhuma alteração foi realizada.", parent=self)
                else:
                    self.db.add_employee(nome, jornada, regime)
                    self.refresh_all()
                    window.destroy()
                    self.show_info("Cadastro concluído", "Funcionário cadastrado com sucesso.", parent=self)
            except sqlite3.IntegrityError:
                self.show_error("Funcionário já cadastrado", "Já existe um funcionário com esse nome.", parent=window)
            except Exception as exc:
                self.show_error("Não foi possível salvar", str(exc), parent=window)

        ctk.CTkButton(
            buttons,
            text="Salvar alterações" if editing else "Cadastrar funcionário",
            command=submit, height=38,
        ).grid(row=0, column=1, padx=(6, 0), sticky="ew")

        refresh_jornada()
        refit_window()
        name_entry.focus_set()
        window.bind("<Return>", lambda _event: submit())

    def on_employee_select(self, _event=None):
        """Compatibilidade com versões anteriores; a seleção atual usa a lista CustomTkinter."""
        return self.selected_employee_id

    def open_new_entry(self):
        self.open_entry_form(entry_id=None)

    def edit_selected_entry(self):
        if not self.selected_entry_id:
            self.show_warning(
                "Selecione um lançamento",
                "Selecione na tabela o lançamento que deseja editar.",
            )
            return
        self.open_entry_form(entry_id=self.selected_entry_id)

    def open_entry_form(self, entry_id=None):
        editing = entry_id is not None
        row = self.db.entry_by_id(entry_id, include_deleted=False) if editing else None
        if editing and not row:
            self.show_warning("Lançamento não encontrado", "O lançamento selecionado não está mais disponível.")
            self.refresh_entries()
            return

        active_employees = list(self.db.employees(only_active=True))
        employee_names = [e["nome"] for e in active_employees]
        if editing and row["funcionario"] not in employee_names:
            employee_names.append(row["funcionario"])
        if not employee_names:
            self.show_warning(
                "Nenhum funcionário ativo",
                "Cadastre ou reative um funcionário antes de criar um lançamento.",
            )
            return

        title = "Editar lançamento" if editing else "Novo lançamento"
        window = self.create_modal(title, 640, 700)
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(2, weight=1)

        ctk.CTkLabel(window, text=title, font=("Segoe UI Semibold", 21)).grid(
            row=0, column=0, padx=26, pady=(20, 4), sticky="w"
        )
        ctk.CTkLabel(
            window,
            text=(
                "Preencha os dados abaixo. Algumas categorias permitem escolher se o lançamento aumenta ou diminui o saldo."
                if not editing else
                "Altere apenas o necessário. O sistema mantém backup e auditoria das modificações."
            ),
            font=("Segoe UI", 11), text_color=("#64748b", "#94a3b8"),
            wraplength=580, justify="left",
        ).grid(row=1, column=0, padx=26, pady=(0, 10), sticky="w")

        form_scroll = ctk.CTkScrollableFrame(window, corner_radius=13, height=460)
        form_scroll.grid(row=2, column=0, padx=26, sticky="nsew")
        form_scroll.grid_columnconfigure(0, weight=1)
        form = ctk.CTkFrame(form_scroll, fg_color="transparent")
        form.grid(row=0, column=0, sticky="ew")
        form.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(form, text="Funcionário", font=("Segoe UI Semibold", 11)).grid(row=0, column=0, padx=12, pady=(8, 3), sticky="w")
        employee_combo = ctk.CTkComboBox(form, values=employee_names, height=35, corner_radius=10, state="readonly")
        employee_combo.grid(row=1, column=0, padx=12, sticky="ew")

        ctk.CTkLabel(form, text="Data", font=("Segoe UI Semibold", 11)).grid(row=2, column=0, padx=12, pady=(11, 3), sticky="w")
        date_var = tk.StringVar(value=br_date(row["data"]) if editing else datetime.now().strftime("%d/%m/%Y"))
        date_widget = PortugueseDateEntry(form, textvariable=date_var, height=35, corner_radius=10)
        date_widget.grid(row=3, column=0, padx=12, sticky="ew")

        ctk.CTkLabel(form, text="Categoria", font=("Segoe UI Semibold", 11)).grid(row=4, column=0, padx=12, pady=(11, 3), sticky="w")
        category_combo = ctk.CTkComboBox(form, values=TIPOS_PADRAO, height=35, corner_radius=10, state="readonly")
        category_combo.grid(row=5, column=0, padx=12, sticky="ew")

        effect_box = ctk.CTkFrame(form, corner_radius=11)
        effect_box.grid(row=6, column=0, padx=12, pady=(10, 4), sticky="ew")
        effect_box.grid_columnconfigure(0, weight=1)
        effect_title = ctk.CTkLabel(effect_box, text="EFEITO NO SALDO", font=("Segoe UI Semibold", 10))
        effect_title.grid(row=0, column=0, padx=12, pady=(7, 0), sticky="w")
        effect_label = ctk.CTkLabel(effect_box, text="", font=("Segoe UI Semibold", 15), anchor="w", justify="left")
        effect_label.grid(row=1, column=0, padx=12, pady=(0, 8), sticky="ew")

        manual_effect_frame = ctk.CTkFrame(form, fg_color="transparent")
        manual_effect_frame.grid(row=7, column=0, padx=12, pady=(2, 2), sticky="ew")
        manual_effect_frame.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(manual_effect_frame, text="Como este lançamento afeta o saldo?", font=("Segoe UI Semibold", 11)).grid(row=0, column=0, pady=(0, 3), sticky="w")
        manual_effect = ctk.CTkSegmentedButton(
            manual_effect_frame, values=["Aumenta o saldo", "Diminui o saldo"], height=34,
        )
        manual_effect.grid(row=1, column=0, sticky="ew")

        original_category = normalize_spaces(row["categoria"]) if editing else ""
        category_changed = {"value": False}

        def categoria_valida(categoria: str) -> bool:
            categoria = normalize_spaces(categoria)
            if categoria in TIPOS_PADRAO:
                return True
            # Preserva lançamentos antigos cuja categoria possa não existir mais
            # na lista atual, desde que ela não tenha sido alterada durante a edição.
            return bool(editing and categoria == original_category)

        def selected_movement():
            current_category = normalize_spaces(category_combo.get())
            if not categoria_valida(current_category):
                return None
            if current_category.lower() in CATEGORIAS_MOVIMENTO_MANUAL:
                return MOVIMENTO_DEBITO if manual_effect.get() == "Diminui o saldo" else MOVIMENTO_CREDITO
            if editing and not category_changed["value"] and current_category == original_category:
                return normalizar_movimento(row["movimento"])
            return guess_movimento(current_category)

        def refit_window():
            def apply_fit():
                try:
                    if window.winfo_exists():
                        self.fit_scrollable_form_modal(
                            window, form_scroll, form, preferred_width=640,
                            min_body_height=250, max_body_height=500,
                        )
                except Exception:
                    pass
            window.after_idle(apply_fit)
            window.after(180, apply_fit)

        def update_effect(_value=None):
            current_category = normalize_spaces(category_combo.get())

            # Enquanto nenhuma categoria válida tiver sido escolhida, não presume
            # crédito nem débito e mantém o painel em estado neutro.
            if not categoria_valida(current_category):
                manual_effect_frame.grid_remove()
                effect_box.configure(fg_color=("#e2e8f0", "#334155"))
                effect_title.configure(text_color=("#475569", "#cbd5e1"))
                effect_label.configure(text="Selecione uma categoria", text_color=("#475569", "#e2e8f0"))
                refit_window()
                return

            allows_manual_movement = current_category.lower() in CATEGORIAS_MOVIMENTO_MANUAL
            if allows_manual_movement:
                manual_effect_frame.grid()
            else:
                manual_effect_frame.grid_remove()

            movimento = selected_movement()
            if movimento == MOVIMENTO_DEBITO:
                effect_box.configure(fg_color=("#fee2e2", "#3f1d24"))
                effect_title.configure(text_color=("#7f1d1d", "#fecaca"))
                effect_label.configure(text="DIMINUI o saldo do funcionário", text_color=("#991b1b", "#fca5a5"))
            else:
                effect_box.configure(fg_color=("#dcfce7", "#123524"))
                effect_title.configure(text_color=("#14532d", "#bbf7d0"))
                effect_label.configure(text="AUMENTA o saldo do funcionário", text_color=("#166534", "#86efac"))
            refit_window()

        def on_category_selected(_value=None):
            category_changed["value"] = True
            update_effect()

        category_combo.configure(command=on_category_selected)
        manual_effect.configure(command=lambda _value: update_effect())

        ctk.CTkLabel(form, text="Quantidade de tempo", font=("Segoe UI Semibold", 11)).grid(row=8, column=0, padx=12, pady=(11, 3), sticky="w")
        qty_frame = ctk.CTkFrame(form, fg_color="transparent")
        qty_frame.grid(row=9, column=0, padx=12, sticky="ew")
        qty_frame.grid_columnconfigure((0, 1, 2), weight=1)
        numeric_vcmd = (self.register(self.validate_numeric_entry), "%P")
        qty_entries = []
        for col, label in enumerate(("Dias", "Horas", "Minutos")):
            ctk.CTkLabel(qty_frame, text=label, font=("Segoe UI", 10)).grid(row=0, column=col, sticky="w", padx=(0 if col == 0 else 4, 0))
            entry = ctk.CTkEntry(qty_frame, height=34, corner_radius=10, validate="key", validatecommand=numeric_vcmd)
            entry.grid(row=1, column=col, padx=(0, 4) if col == 0 else ((4, 4) if col == 1 else (4, 0)), sticky="ew")
            qty_entries.append(entry)
        days_entry, hours_entry, minutes_entry = qty_entries

        ctk.CTkLabel(form, text="Observação", font=("Segoe UI Semibold", 11)).grid(row=10, column=0, padx=12, pady=(11, 3), sticky="w")
        observation_box = ctk.CTkTextbox(form, height=82, corner_radius=10)
        observation_box.grid(row=11, column=0, padx=12, pady=(0, 10), sticky="ew")

        if editing:
            employee_combo.set(row["funcionario"])
            category_combo.set(row["categoria"])
            manual_effect.set("Diminui o saldo" if normalize_spaces(row["movimento"]).startswith(MOVIMENTO_DEBITO) else "Aumenta o saldo")
            hpd = int(row["jornada_horas"] or DEFAULT_HOURS_PER_DAY)
            total = abs(int(row["minutos"] or 0))
            day_minutes = hpd * 60
            days = total // day_minutes
            rest = total % day_minutes
            hours = rest // 60
            minutes = rest % 60
            if days:
                days_entry.insert(0, str(days))
            if hours:
                hours_entry.insert(0, str(hours))
            if minutes:
                minutes_entry.insert(0, str(minutes))
            observation_box.insert("1.0", row["observacao"] or "")
        else:
            employee_combo.set("Selecione o funcionário")
            category_combo.set(CATEGORIA_PLACEHOLDER)
            manual_effect.set("Aumenta o saldo")

        def calculate_minutes(employee):
            fields = (("Dias", days_entry), ("Horas", hours_entry), ("Minutos", minutes_entry))
            for field_name, entry in fields:
                value = entry.get().strip()
                if value and not value.isdigit():
                    raise ValueError(f"O campo {field_name} aceita somente números.")
            days = int(days_entry.get() or "0")
            hours = int(hours_entry.get() or "0")
            minutes = int(minutes_entry.get() or "0")
            hpd = int(employee["jornada_horas"] or DEFAULT_HOURS_PER_DAY)
            total = days * hpd * 60 + hours * 60 + minutes
            parts = []
            if days:
                parts.append(f"{days} dia" + ("s" if days != 1 else ""))
            if hours:
                parts.append(f"{hours} hora" + ("s" if hours != 1 else ""))
            if minutes:
                parts.append(f"{minutes} minuto" + ("s" if minutes != 1 else ""))
            return total, " e ".join(parts) if parts else "0"

        def submit():
            try:
                employee_name = employee_combo.get().strip()
                employee = self.db.employee_by_name(employee_name)
                if not employee:
                    raise ValueError("Selecione um funcionário válido.")
                data_iso = parse_date(date_var.get().strip())
                if not data_iso:
                    raise ValueError("Informe uma data válida no formato DD/MM/AAAA.")
                categoria = normalize_spaces(category_combo.get())
                if not categoria_valida(categoria):
                    raise ValueError("Selecione uma categoria válida.")
                movimento = selected_movement()
                total_minutes, quantidade_texto = calculate_minutes(employee)
                if total_minutes <= 0:
                    raise ValueError("Informe dias, horas ou minutos para o lançamento.")
                observacao = observation_box.get("1.0", "end").strip()
                if not observacao:
                    raise ValueError("Informe uma observação para o lançamento.")

                if editing:
                    create_database_backup("alteracao_lancamento")
                    changed = self.db.update_entry(
                        entry_id, employee["id"], data_iso, movimento, categoria,
                        quantidade_texto, total_minutes, observacao,
                    )
                    self.refresh_entries()
                    self.refresh_dashboard()
                    window.destroy()
                    if changed:
                        self.show_info("Alterações salvas", "Alterações do lançamento salvas com sucesso.", parent=self)
                    else:
                        self.show_info("Nenhuma alteração", "Nenhuma alteração foi realizada.", parent=self)
                else:
                    self.db.add_entry(
                        employee["id"], data_iso, movimento, categoria,
                        quantidade_texto, total_minutes, observacao,
                    )
                    self.refresh_entries()
                    self.refresh_dashboard()
                    window.destroy()
                    self.show_info("Cadastro concluído", "Lançamento cadastrado com sucesso.", parent=self)
            except Exception as exc:
                self.show_error("Não foi possível salvar", str(exc), parent=window)

        buttons = ctk.CTkFrame(window, fg_color="transparent")
        buttons.grid(row=3, column=0, padx=26, pady=(14, 20), sticky="ew")
        buttons.grid_columnconfigure((0, 1), weight=1)
        ctk.CTkButton(
            buttons, text="Cancelar", command=window.destroy, height=38,
            fg_color="#475569", hover_color="#334155",
        ).grid(row=0, column=0, padx=(0, 6), sticky="ew")
        ctk.CTkButton(
            buttons, text="Salvar alterações" if editing else "Cadastrar lançamento",
            command=submit, height=38,
        ).grid(row=0, column=1, padx=(6, 0), sticky="ew")

        update_effect()
        refit_window()
        employee_combo.focus_set()

    def on_entry_select(self, _event=None):
        selected = self.tree_entries.selection()
        self.selected_entry_id = None
        if not selected:
            return
        values = self.tree_entries.item(selected[0], "values")
        if values:
            self.selected_entry_id = int(values[0])

    # Métodos mantidos para compatibilidade com chamadas antigas internas.
    def on_categoria_change(self, value):
        return guess_movimento(value)

    def on_movimento_change(self, value):
        return normalizar_movimento(value)

    def update_movimento_help(self, movimento=None):
        return movimento

    def save_employee(self):
        self.open_new_employee()

    def clear_employee_form(self):
        self.selected_employee_id = None
        if hasattr(self, "employee_row_widgets"):
            for data in self.employee_row_widgets.values():
                data["row"].configure(
                    fg_color=data["base_bg"],
                    border_width=0,
                    border_color=("#ffffff", "#111827"),
                )
                data["accent"].configure(fg_color="transparent")

    def calculate_form_minutes(self) -> tuple[int, str]:
        return 0, "0"

    def save_entry(self):
        self.open_new_entry()

    def clear_entry_form(self):
        self.selected_entry_id = None
        if hasattr(self, "tree_entries"):
            self.tree_entries.selection_remove(self.tree_entries.selection())

    def show_readonly_details(self, title: str, content: str, parent=None):
        """Exibe textos longos em uma janela própria, legível e copiável."""
        parent = parent if parent is not None else self
        window = self.create_modal(
            title, 760, 560, resizable=True, parent=parent, grab=True
        )
        self.responsive_minsize(window, 620, 440, floor_w=360, floor_h=280)
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(1, weight=1)

        ctk.CTkLabel(
            window,
            text=title,
            font=("Segoe UI Semibold", 20),
            anchor="w",
        ).grid(row=0, column=0, padx=18, pady=(16, 8), sticky="ew")

        text_box = ctk.CTkTextbox(window, wrap="word", font=("Segoe UI", 11))
        text_box.grid(row=1, column=0, padx=18, pady=(0, 10), sticky="nsew")
        text_box.insert("1.0", str(content or ""))
        text_box.configure(state="disabled")

        actions = ctk.CTkFrame(window, fg_color="transparent")
        actions.grid(row=2, column=0, padx=18, pady=(0, 16), sticky="ew")
        actions.grid_columnconfigure((0, 1), weight=1)

        def copy_details():
            self.clipboard_clear()
            self.clipboard_append(str(content or ""))
            self.update()
            self.show_info("Detalhes copiados", "Os detalhes foram copiados para a área de transferência.", parent=window)

        ctk.CTkButton(
            actions,
            text="Copiar detalhes",
            command=copy_details,
            fg_color="#475569",
            hover_color="#334155",
        ).grid(row=0, column=0, padx=(0, 6), sticky="ew")
        ctk.CTkButton(actions, text="Fechar", command=window.destroy).grid(row=0, column=1, padx=(6, 0), sticky="ew")
        window.protocol("WM_DELETE_WINDOW", window.destroy)
        window.focus_force()

    def prompt_support_authorization(self, scope: str, action_label: str) -> bool:
        """Desafio/resposta assinado offline, específico desta instalação e operação."""
        try:
            challenge = issue_request(self.db, scope)
        except SupportAuthorizationError as exc:
            self.show_error("Suporte não configurado", str(exc))
            return False

        result = {"authorized": False}
        window = self.create_modal("Autorização de suporte", 790, 510)
        window.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            window, text="Autorização exclusiva de suporte",
            font=("Segoe UI Semibold", 20),
        ).grid(row=0, column=0, padx=24, pady=(18, 4), sticky="w")
        ctk.CTkLabel(
            window,
            text=(f"Para {action_label}, envie a solicitação abaixo ao suporte autorizado. "
                  "O suporte devolverá uma resposta assinada, válida apenas para esta ação, "
                  "nesta instalação, por até 30 minutos e por um único uso."),
            justify="left", wraplength=730,
            text_color=("#475569", "#cbd5e1"),
        ).grid(row=1, column=0, padx=24, pady=(4, 10), sticky="ew")

        ctk.CTkLabel(window, text="Solicitação de suporte (copie o conteúdo completo)").grid(
            row=2, column=0, padx=24, sticky="w"
        )
        request_box = ctk.CTkTextbox(window, height=85, wrap="char")
        request_box.grid(row=3, column=0, padx=24, pady=(4, 8), sticky="ew")
        request_box.insert("1.0", challenge)
        request_box.configure(state="disabled")

        def copy_request():
            self.clipboard_clear()
            self.clipboard_append(challenge)
            self.update()

        ctk.CTkButton(
            window, text="Copiar solicitação", command=copy_request,
            fg_color="#475569", hover_color="#334155",
        ).grid(row=4, column=0, padx=24, pady=(0, 12), sticky="w")

        ctk.CTkLabel(window, text="Resposta assinada pelo suporte").grid(
            row=5, column=0, padx=24, sticky="w"
        )
        answer = ctk.CTkEntry(window, height=36, placeholder_text="CF41-...")
        answer.grid(row=6, column=0, padx=24, pady=(4, 4), sticky="ew")

        def confirm(_event=None):
            response = answer.get().strip()
            if not response:
                self.show_warning("Resposta obrigatória", "Informe a resposta enviada pelo suporte.", parent=window)
                return
            try:
                verify_and_consume(self.db, scope, response)
            except SupportAuthorizationError as exc:
                self.db.record_audit(
                    "FALHA AUTORIZAÇÃO DE SUPORTE",
                    detalhes=f"Solicitação recusada para {action_label}: {str(exc)}",
                )
                self.show_error("Autorização recusada", str(exc), parent=window)
                return
            self.db.record_audit(
                "AUTORIZAÇÃO DE SUPORTE ASSINADA",
                detalhes=f"Resposta assinada validada para {action_label}.",
            )
            result["authorized"] = True
            window.destroy()

        buttons = ctk.CTkFrame(window, fg_color="transparent")
        buttons.grid(row=7, column=0, padx=24, pady=(14, 18), sticky="ew")
        buttons.grid_columnconfigure((0, 1), weight=1)
        ctk.CTkButton(buttons, text="Cancelar", command=window.destroy,
                      fg_color="#475569", hover_color="#334155").grid(
            row=0, column=0, padx=(0, 6), sticky="ew"
        )
        ctk.CTkButton(buttons, text="Validar resposta", command=confirm,
                      fg_color="#991b1b", hover_color="#7f1d1d").grid(
            row=0, column=1, padx=(6, 0), sticky="ew"
        )
        answer.bind("<Return>", confirm)
        window.protocol("WM_DELETE_WINDOW", window.destroy)
        answer.focus_set()
        self.wait_window(window)
        return result["authorized"]

    def authorize_master_support(self, action: str, scope: str) -> bool:
        """Mantém a chamada administrativa, sem aceitar senha mestra legada."""
        return self.prompt_support_authorization(scope, action)

    def prompt_service_name(self, title="Cadastrar nome do serviço", current_name="", required=False):
        """Exibe uma janela simples para cadastrar ou editar o nome do serviço."""
        result = {"value": None}
        window = self.create_modal(title, 560, 300)
        window.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            window,
            text=title,
            font=("Segoe UI Semibold", 20),
        ).grid(row=0, column=0, padx=24, pady=(22, 6), sticky="w")

        intro = (
            "Informe o nome da unidade ou serviço que utiliza este sistema. "
            "Esse nome será exibido no cabeçalho principal do programa e nos relatórios preparados para assinatura."
        )
        ctk.CTkLabel(
            window,
            text=intro,
            justify="left",
            wraplength=500,
            text_color=("#475569", "#cbd5e1"),
        ).grid(row=1, column=0, padx=24, pady=(0, 14), sticky="w")

        ctk.CTkLabel(window, text="Nome do serviço").grid(
            row=2, column=0, padx=24, pady=(4, 2), sticky="w"
        )
        entry_name = ctk.CTkEntry(
            window,
            height=36,
            placeholder_text="Ex.: SAICA Resplandecer",
        )
        entry_name.grid(row=3, column=0, padx=24, sticky="ew")
        if current_name:
            entry_name.insert(0, current_name)
            entry_name.select_range(0, "end")

        def submit(_event=None):
            service_name = normalize_spaces(entry_name.get())
            if not service_name:
                self.show_error(
                    "Nome obrigatório",
                    "Informe o nome do serviço antes de continuar.",
                    parent=window,
                )
                entry_name.focus_set()
                return
            result["value"] = service_name
            window.destroy()

        def cancel():
            if required:
                close_app = self.ask_yes_no(
                    "Cadastro obrigatório",
                    "O nome do serviço é necessário para identificar o programa e os relatórios. "
                    "Deseja fechar o programa sem concluir o cadastro?",
                    parent=window,
                )
                if not close_app:
                    entry_name.focus_set()
                    return
            window.destroy()

        buttons = ctk.CTkFrame(window, fg_color="transparent")
        buttons.grid(row=4, column=0, padx=24, pady=22, sticky="ew")
        buttons.grid_columnconfigure((0, 1), weight=1)
        ctk.CTkButton(
            buttons,
            text="Sair" if required else "Cancelar",
            command=cancel,
            fg_color="#475569",
            hover_color="#334155",
        ).grid(row=0, column=0, padx=(0, 6), sticky="ew")
        ctk.CTkButton(buttons, text="Salvar", command=submit).grid(
            row=0, column=1, padx=(6, 0), sticky="ew"
        )

        entry_name.bind("<Return>", submit)
        window.protocol("WM_DELETE_WINDOW", cancel)
        entry_name.focus_set()
        self.wait_window(window)
        return result["value"]

    def ensure_service_name(self) -> bool:
        """Solicita o nome somente quando ainda não existe cadastro no banco atual."""
        if self.db.get_service_name():
            return True

        service_name = self.prompt_service_name(required=True)
        if not service_name:
            try:
                self.destroy()
            except Exception:
                pass
            return False

        create_database_backup("cadastro_nome_servico")
        self.db.set_service_name(service_name)
        self.db.record_audit(
            "NOME DO SERVIÇO CADASTRADO",
            detalhes=f"Serviço cadastrado: {service_name}.",
        )
        self.refresh_header_title()
        return True

    def edit_service_name(self):
        """Permite alterar o nome do serviço somente após autorização administrativa."""
        if not self.authorize_admin("editar o nome do serviço"):
            return

        current_name = self.db.get_service_name()
        service_name = self.prompt_service_name(
            title="Editar nome do serviço",
            current_name=current_name,
            required=False,
        )
        if not service_name or service_name == current_name:
            return

        create_database_backup("alteracao_nome_servico")
        self.db.set_service_name(service_name)
        self.db.record_audit(
            "NOME DO SERVIÇO ALTERADO",
            detalhes=f"Nome anterior: {current_name or '(não informado)'} | Novo nome: {service_name}.",
        )
        self.refresh_header_title()
        self.show_info(
            "Nome atualizado",
            "O nome do serviço foi atualizado com sucesso no cabeçalho principal e nos próximos relatórios.",
        )

    def prompt_new_admin_password(self, title="Criar senha administrativa", intro=None):
        result = {"value": None}
        window = self.create_modal(title, 500, 350)
        window.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            window,
            text=title,
            font=("Segoe UI Semibold", 19),
        ).grid(row=0, column=0, padx=24, pady=(22, 6), sticky="w")
        ctk.CTkLabel(
            window,
            text=intro or "A senha deve ter pelo menos 8 caracteres.",
            justify="left",
            wraplength=440,
            text_color=("#475569", "#cbd5e1"),
        ).grid(row=1, column=0, padx=24, pady=(0, 14), sticky="w")

        ctk.CTkLabel(window, text="Nova senha").grid(row=2, column=0, padx=24, pady=(4, 2), sticky="w")
        entry_password = ctk.CTkEntry(window, show="*", height=34)
        entry_password.grid(row=3, column=0, padx=24, sticky="ew")

        ctk.CTkLabel(window, text="Confirmar nova senha").grid(row=4, column=0, padx=24, pady=(12, 2), sticky="w")
        entry_confirm = ctk.CTkEntry(window, show="*", height=34)
        entry_confirm.grid(row=5, column=0, padx=24, sticky="ew")

        def submit(_event=None):
            password = entry_password.get()
            confirmation = entry_confirm.get()
            if len(password) < 8:
                self.show_error("Senha inválida", "A senha deve ter pelo menos 8 caracteres.", parent=window)
                return
            if password != confirmation:
                self.show_error("Senha inválida", "A confirmação não corresponde à nova senha.", parent=window)
                return
            result["value"] = password
            window.destroy()

        buttons = ctk.CTkFrame(window, fg_color="transparent")
        buttons.grid(row=6, column=0, padx=24, pady=22, sticky="ew")
        buttons.grid_columnconfigure((0, 1), weight=1)
        ctk.CTkButton(buttons, text="Cancelar", command=window.destroy, fg_color="#475569", hover_color="#334155").grid(row=0, column=0, padx=(0, 6), sticky="ew")
        ctk.CTkButton(buttons, text="Confirmar", command=submit).grid(row=0, column=1, padx=(6, 0), sticky="ew")

        entry_password.bind("<Return>", lambda _e: entry_confirm.focus_set())
        entry_confirm.bind("<Return>", submit)
        window.protocol("WM_DELETE_WINDOW", window.destroy)
        entry_password.focus_set()
        self.wait_window(window)
        return result["value"]

    def prompt_admin_authentication(self, action: str, locked_seconds: int = 0):
        result = {"kind": None, "value": None}
        height = 355 if locked_seconds else 340
        window = self.create_modal("Autorização administrativa", 500, height)
        window.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(window, text="Autorização administrativa", font=("Segoe UI Semibold", 19)).grid(row=0, column=0, padx=24, pady=(22, 6), sticky="w")
        message = f"Digite a senha administrativa para {action}."
        if locked_seconds:
            minutes = max(1, (locked_seconds + 59) // 60)
            message = (
                f"O acesso por senha está temporariamente bloqueado por aproximadamente {minutes} minuto(s). "
                "Use o código de recuperação ou tente novamente depois."
            )
        ctk.CTkLabel(
            window,
            text=message,
            justify="left",
            wraplength=440,
            text_color=("#475569", "#cbd5e1"),
        ).grid(row=1, column=0, padx=24, pady=(0, 14), sticky="w")

        entry_password = ctk.CTkEntry(window, show="*", height=34, state="disabled" if locked_seconds else "normal")
        entry_password.grid(row=2, column=0, padx=24, sticky="ew")

        def confirm(_event=None):
            if locked_seconds:
                return
            result["kind"] = "password"
            result["value"] = entry_password.get()
            window.destroy()

        def recover():
            result["kind"] = "recovery"
            window.destroy()

        def recover_by_support():
            result["kind"] = "support"
            window.destroy()

        ctk.CTkButton(window, text="Confirmar senha", command=confirm, state="disabled" if locked_seconds else "normal").grid(row=3, column=0, padx=24, pady=(14, 6), sticky="ew")
        ctk.CTkButton(window, text="Usar código de recuperação", command=recover, fg_color="#475569", hover_color="#334155").grid(row=4, column=0, padx=24, pady=6, sticky="ew")
        ctk.CTkButton(window, text="Solicitar recuperação ao suporte", command=recover_by_support, fg_color="#475569", hover_color="#334155").grid(row=5, column=0, padx=24, pady=6, sticky="ew")
        ctk.CTkButton(window, text="Cancelar", command=window.destroy, fg_color="transparent", border_width=1, text_color=("#334155", "#e2e8f0")).grid(row=6, column=0, padx=24, pady=(6, 18), sticky="ew")

        entry_password.bind("<Return>", confirm)
        window.protocol("WM_DELETE_WINDOW", window.destroy)
        if not locked_seconds:
            entry_password.focus_set()
        self.wait_window(window)
        return result["kind"], result["value"]

    def prompt_recovery_code(self):
        result = {"value": None}
        window = self.create_modal("Recuperar acesso administrativo", 520, 300)
        window.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(window, text="Código de recuperação", font=("Segoe UI Semibold", 19)).grid(row=0, column=0, padx=24, pady=(22, 6), sticky="w")
        ctk.CTkLabel(
            window,
            text="Informe o código gerado quando a senha administrativa foi criada ou alterada.",
            justify="left",
            wraplength=460,
            text_color=("#475569", "#cbd5e1"),
        ).grid(row=1, column=0, padx=24, pady=(0, 14), sticky="w")
        entry = ctk.CTkEntry(window, height=36, placeholder_text="XXXX-XXXX-XXXX-XXXX-XXXX")
        entry.grid(row=2, column=0, padx=24, sticky="ew")

        def submit(_event=None):
            value = entry.get().strip()
            if not value:
                self.show_warning("Código obrigatório", "Informe o código de recuperação.", parent=window)
                return
            result["value"] = value
            window.destroy()

        buttons = ctk.CTkFrame(window, fg_color="transparent")
        buttons.grid(row=3, column=0, padx=24, pady=22, sticky="ew")
        buttons.grid_columnconfigure((0, 1), weight=1)
        ctk.CTkButton(buttons, text="Cancelar", command=window.destroy, fg_color="#475569", hover_color="#334155").grid(row=0, column=0, padx=(0, 6), sticky="ew")
        ctk.CTkButton(buttons, text="Validar código", command=submit).grid(row=0, column=1, padx=(6, 0), sticky="ew")
        entry.bind("<Return>", submit)
        window.protocol("WM_DELETE_WINDOW", window.destroy)
        entry.focus_set()
        self.wait_window(window)
        return result["value"]

    def show_recovery_code_dialog(self, recovery_code: str):
        window = self.create_modal("Código de recuperação", 590, 360)
        window.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(window, text="Guarde este código de recuperação", font=("Segoe UI Semibold", 20)).grid(row=0, column=0, padx=24, pady=(22, 8), sticky="w")
        ctk.CTkLabel(
            window,
            text=(
                "Este código será mostrado somente agora. Ele permite redefinir a senha administrativa. "
                "Guarde-o em local seguro e não o deixe disponível aos usuários comuns."
            ),
            justify="left",
            wraplength=530,
            text_color=("#475569", "#cbd5e1"),
        ).grid(row=1, column=0, padx=24, pady=(0, 14), sticky="w")

        code_entry = ctk.CTkEntry(window, height=42, font=("Consolas", 17), justify="center")
        code_entry.grid(row=2, column=0, padx=24, sticky="ew")
        code_entry.insert(0, recovery_code)
        code_entry.configure(state="disabled")

        def copy_code():
            self.clipboard_clear()
            self.clipboard_append(recovery_code)
            self.update()
            self.show_info("Código copiado", "O código foi copiado para a área de transferência.", parent=window)

        def save_code():
            file = filedialog.asksaveasfilename(
                parent=window,
                title="Salvar código de recuperação",
                defaultextension=".txt",
                filetypes=[("Arquivo de texto", "*.txt")],
                initialfile="codigo_recuperacao_controle_folgas.txt",
            )
            if not file:
                return
            Path(file).write_text(
                "Código de recuperação administrativa - Controle de Folgas\n\n"
                f"{recovery_code}\n\n"
                "Guarde este arquivo em local seguro.\n",
                encoding="utf-8",
            )
            self.show_info("Código salvo", f"Código salvo com sucesso:\n{file}", parent=window)

        actions = ctk.CTkFrame(window, fg_color="transparent")
        actions.grid(row=3, column=0, padx=24, pady=(18, 8), sticky="ew")
        actions.grid_columnconfigure((0, 1), weight=1)
        ctk.CTkButton(actions, text="Copiar código", command=copy_code, fg_color="#475569", hover_color="#334155").grid(row=0, column=0, padx=(0, 6), sticky="ew")
        ctk.CTkButton(actions, text="Salvar em arquivo", command=save_code).grid(row=0, column=1, padx=(6, 0), sticky="ew")
        ctk.CTkButton(window, text="Fechar", command=window.destroy).grid(row=4, column=0, padx=24, pady=(8, 22), sticky="ew")
        window.protocol("WM_DELETE_WINDOW", window.destroy)
        self.wait_window(window)

    def setup_admin_password(self) -> bool:
        password = self.prompt_new_admin_password(
            title="Criar senha administrativa",
            intro=(
                "Esta senha protegerá o envio para a Lixeira, restaurações, acesso à Lixeira e à auditoria. "
                "Use pelo menos 8 caracteres."
            ),
        )
        if not password:
            return False
        create_database_backup("criacao_senha_administrativa")
        recovery_code = self.db.set_admin_credentials(password)
        self.db.record_audit("SENHA ADMINISTRATIVA CRIADA", detalhes="Credencial administrativa configurada no sistema.")
        self.show_recovery_code_dialog(recovery_code)
        return True

    def recover_admin_access(self) -> bool:
        code = self.prompt_recovery_code()
        if not code:
            return False
        if not self.db.verify_recovery_code(code):
            self.db.record_audit("FALHA NO CÓDIGO DE RECUPERAÇÃO", detalhes="Tentativa com código de recuperação inválido.")
            self.show_error("Código inválido", "O código de recuperação informado não é válido.")
            return False

        new_password = self.prompt_new_admin_password(
            title="Definir nova senha administrativa",
            intro="Código validado. Cadastre uma nova senha com pelo menos 8 caracteres.",
        )
        if not new_password:
            return False
        create_database_backup("recuperacao_senha_administrativa")
        recovery_code = self.db.set_admin_credentials(new_password)
        self.db.reset_admin_attempts()
        self.db.record_audit("SENHA REDEFINIDA POR RECUPERAÇÃO", detalhes="Acesso recuperado com o código de recuperação.")
        self.show_recovery_code_dialog(recovery_code)
        self.show_info("Acesso recuperado", "A nova senha administrativa foi cadastrada com sucesso.")
        return True

    def reset_admin_with_support(self, action: str) -> bool:
        """Redefine credenciais após autorização de suporte assinada."""
        new_password = self.prompt_new_admin_password(
            title="Redefinir senha administrativa",
            intro=(
                "Autorização assinada validada. Cadastre agora uma nova senha administrativa local "
                "com pelo menos 8 caracteres. Um novo código de recuperação será gerado."
            ),
        )
        if not new_password:
            self.db.record_audit(
                "RECUPERAÇÃO POR SUPORTE CANCELADA",
                detalhes=f"A redefinição foi cancelada durante a autorização para {action}.",
            )
            return False

        create_database_backup("recuperacao_por_suporte")
        recovery_code = self.db.set_admin_credentials(new_password)
        self.db.reset_admin_attempts()
        self.db.record_audit(
            "ACESSO RECUPERADO POR SUPORTE",
            detalhes=(
                f"Senha local redefinida por autorização assinada durante a solicitação para {action}. "
                "Um novo código de recuperação foi gerado e o anterior foi invalidado."
            ),
        )
        self.show_recovery_code_dialog(recovery_code)
        self.show_info(
            "Acesso recuperado",
            "A senha administrativa local foi redefinida com sucesso. "
            "A ação solicitada está autorizada.",
        )
        return True

    def authorize_admin(self, action: str) -> bool:
        self._admin_credentials_reset_during_auth = False

        if not self.db.admin_password_exists():
            return self.setup_admin_password()

        remaining = self.db.admin_lock_remaining_seconds()
        kind, password = self.prompt_admin_authentication(action, locked_seconds=remaining)
        if kind == "recovery":
            recovered = self.recover_admin_access()
            self._admin_credentials_reset_during_auth = recovered
            return recovered
        if kind == "support":
            if not self.prompt_support_authorization("ADMIN_RESET", "recuperar a senha administrativa"):
                return False
            recovered = self.reset_admin_with_support(action)
            self._admin_credentials_reset_during_auth = recovered
            return recovered
        if kind != "password":
            return False

        informed_password = password or ""

        if self.db.verify_admin_password(informed_password):
            self.db.reset_admin_attempts()
            self.db.record_audit("AUTORIZAÇÃO ADMINISTRATIVA", detalhes=f"Acesso autorizado para {action}.")
            return True

        attempts, locked = self.db.register_failed_admin_attempt()
        self.db.record_audit("FALHA DE SENHA ADMINISTRATIVA", detalhes=f"Tentativa inválida para {action}.")
        if locked:
            self.show_error(
                "Acesso bloqueado",
                f"Foram realizadas {ADMIN_MAX_ATTEMPTS} tentativas incorretas. "
                f"O acesso por senha ficará bloqueado por {ADMIN_LOCK_MINUTES} minutos.",
            )
        else:
            remaining_attempts = ADMIN_MAX_ATTEMPTS - attempts
            self.show_error(
                "Senha incorreta",
                f"A senha administrativa está incorreta. Restam {remaining_attempts} tentativa(s) antes do bloqueio temporário.",
            )
        return False

    def open_admin_options(self):
        if not self.db.admin_password_exists() and not self.setup_admin_password():
            return

        window = self.create_modal("Opções administrativas", 520, 570)
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(2, weight=1)

        ctk.CTkLabel(
            window, text="Opções administrativas", font=("Segoe UI Semibold", 21)
        ).grid(row=0, column=0, padx=24, pady=(18, 4), sticky="w")
        ctk.CTkLabel(
            window,
            text="As funções abaixo são protegidas e solicitam autorização administrativa.",
            justify="left",
            wraplength=460,
            text_color=("#475569", "#cbd5e1"),
        ).grid(row=1, column=0, padx=24, pady=(0, 10), sticky="w")

        # As opções ficam em uma região rolável. Assim, mesmo em 150% ou em
        # uma tela baixa, o botão Fechar permanece sempre acessível.
        options = ctk.CTkScrollableFrame(window, corner_radius=12, height=330)
        options.grid(row=2, column=0, padx=24, sticky="nsew")
        options.grid_columnconfigure(0, weight=1)

        option_buttons = [
            ("Enviar lançamento selecionado para a Lixeira", self.admin_delete_selected_entry, {}),
            ("Abrir Lixeira", self.open_trash, {}),
            ("Consultar auditoria", self.open_audit_log, {}),
            ("Editar nome do serviço", self.edit_service_name, {"fg_color": "#475569", "hover_color": "#334155"}),
            ("Alterar senha administrativa", self.change_admin_password, {"fg_color": "#475569", "hover_color": "#334155"}),
            ("Recuperar acesso com código", self.recover_admin_access, {"fg_color": "#475569", "hover_color": "#334155"}),
        ]
        for row_index, (text, command, colors) in enumerate(option_buttons):
            ctk.CTkButton(
                options, text=text, command=command, height=38, **colors
            ).grid(row=row_index, column=0, padx=8, pady=6, sticky="ew")

        ctk.CTkButton(
            window,
            text="Fechar",
            command=window.destroy,
            height=36,
            fg_color="transparent",
            border_width=1,
            text_color=("#334155", "#e2e8f0"),
        ).grid(row=3, column=0, padx=24, pady=(12, 18), sticky="ew")
        window.protocol("WM_DELETE_WINDOW", window.destroy)

    def change_admin_password(self):
        if not self.authorize_admin("alterar a senha administrativa"):
            return
        # Se a autorização ocorreu por código de recuperação ou por suporte,
        # a nova senha já foi cadastrada durante o próprio fluxo de recuperação.
        if self._admin_credentials_reset_during_auth:
            return
        new_password = self.prompt_new_admin_password(
            title="Alterar senha administrativa",
            intro="Informe e confirme a nova senha. O código de recuperação atual será substituído.",
        )
        if not new_password:
            return
        create_database_backup("alteracao_senha_administrativa")
        recovery_code = self.db.set_admin_credentials(new_password)
        self.db.record_audit("SENHA ADMINISTRATIVA ALTERADA", detalhes="Senha alterada e novo código de recuperação gerado.")
        self.show_recovery_code_dialog(recovery_code)
        self.show_info("Senha alterada", "A senha administrativa foi alterada com sucesso.")

    def prompt_delete_reason(self, row):
        result = {"value": None}
        window = self.create_modal("Motivo da exclusão", 620, 470)
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(1, weight=1)

        hpd = int(row["jornada_horas"] or DEFAULT_HOURS_PER_DAY)
        details = (
            f"Funcionário: {row['funcionario']}\n"
            f"Data: {br_date(row['data'])}\n"
            f"Movimento: {row['movimento']}\n"
            f"Categoria: {row['categoria']}\n"
            f"Quantidade: {row['quantidade_texto'] or format_minutes(abs(row['minutos']), hpd)}\n"
            f"Observação: {row['observacao'] or '(sem observação)'}"
        )

        ctk.CTkLabel(
            window, text="Enviar lançamento para a Lixeira", font=("Segoe UI Semibold", 20)
        ).grid(row=0, column=0, padx=24, pady=(18, 8), sticky="w")

        body = ctk.CTkScrollableFrame(window, corner_radius=12, height=300)
        body.grid(row=1, column=0, padx=24, sticky="nsew")
        body.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            body, text=details, justify="left", anchor="w", wraplength=540
        ).grid(row=0, column=0, padx=10, pady=(8, 12), sticky="ew")
        ctk.CTkLabel(body, text="Motivo obrigatório da exclusão").grid(
            row=1, column=0, padx=10, pady=(2, 4), sticky="w"
        )
        reason_box = ctk.CTkTextbox(body, height=100)
        reason_box.grid(row=2, column=0, padx=10, pady=(0, 10), sticky="ew")

        def submit():
            reason = normalize_spaces(reason_box.get("1.0", "end"))
            if not reason:
                self.show_warning("Motivo obrigatório", "Informe o motivo da exclusão.", parent=window)
                return
            result["value"] = reason
            window.destroy()

        buttons = ctk.CTkFrame(window, fg_color="transparent")
        buttons.grid(row=2, column=0, padx=24, pady=(12, 18), sticky="ew")
        buttons.grid_columnconfigure((0, 1), weight=1)
        ctk.CTkButton(
            buttons, text="Cancelar", command=window.destroy,
            fg_color="#475569", hover_color="#334155"
        ).grid(row=0, column=0, padx=(0, 6), sticky="ew")
        ctk.CTkButton(
            buttons, text="Continuar", command=submit,
            fg_color="#b45309", hover_color="#92400e"
        ).grid(row=0, column=1, padx=(6, 0), sticky="ew")
        window.protocol("WM_DELETE_WINDOW", window.destroy)
        reason_box.focus_set()
        self.wait_window(window)
        return result["value"]

    def admin_delete_selected_entry(self):
        if not self.selected_entry_id:
            self.show_warning("Selecione um lançamento", "Selecione na tabela o lançamento que deseja enviar para a Lixeira.")
            return
        row = self.db.entry_by_id(self.selected_entry_id, include_deleted=False)
        if not row:
            self.show_warning("Registro indisponível", "O lançamento selecionado não está mais disponível.")
            self.refresh_entries()
            return
        if not self.authorize_admin("enviar o lançamento selecionado para a Lixeira"):
            return

        reason = self.prompt_delete_reason(row)
        if not reason:
            return
        confirm_text = (
            "Confirma o envio deste lançamento para a Lixeira?\n\n"
            f"Funcionário: {row['funcionario']}\n"
            f"Data: {br_date(row['data'])}\n"
            f"Categoria: {row['categoria']}\n"
            f"Motivo: {reason}\n\n"
            "O registro deixará de participar dos saldos e relatórios, mas poderá ser restaurado."
        )
        if not self.ask_yes_no("Confirmar envio para a Lixeira", confirm_text):
            return

        try:
            create_database_backup("lancamento_para_lixeira")
            self.db.soft_delete_entry(self.selected_entry_id, reason)
            self.clear_entry_form()
            self.refresh_entries()
            self.refresh_dashboard()
            if hasattr(self, "tree_trash"):
                self.refresh_trash()
            self.show_info("Lançamento protegido", "O lançamento foi enviado para a Lixeira.")
        except Exception as exc:
            self.show_error("Erro", str(exc))

    def delete_selected_entry(self):
        """Mantido por compatibilidade com versões anteriores da interface."""
        self.admin_delete_selected_entry()

    def open_trash(self):
        if not self.authorize_admin("abrir a Lixeira"):
            return
        if hasattr(self, "trash_window") and self.trash_window is not None and self.trash_window.winfo_exists():
            self.trash_window.lift()
            self.refresh_trash()
            return

        self.trash_window = self.create_modal(
            "Lixeira de lançamentos", 1100, 650, resizable=True, grab=False
        )
        self.responsive_minsize(self.trash_window, 900, 540, floor_w=520, floor_h=360)
        self.trash_window.grid_columnconfigure(0, weight=1)
        self.trash_window.grid_rowconfigure(1, weight=1)

        header = ctk.CTkFrame(self.trash_window, corner_radius=0)
        header.grid(row=0, column=0, sticky="ew")
        header.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(header, text="Lixeira de lançamentos", font=("Segoe UI Semibold", 21)).grid(row=0, column=0, padx=18, pady=(14, 2), sticky="w")
        self.trash_count_label = ctk.CTkLabel(header, text="", text_color=("#475569", "#cbd5e1"))
        self.trash_count_label.grid(row=1, column=0, padx=18, pady=(0, 2), sticky="w")
        ctk.CTkLabel(
            header,
            text="Dê dois cliques em uma linha para visualizar todos os detalhes. Exclusões definitivas exigem uma resposta assinada do suporte.",
            text_color=("#475569", "#cbd5e1"),
            justify="left",
            anchor="w",
            wraplength=1040,
        ).grid(row=2, column=0, padx=18, pady=(0, 12), sticky="ew")

        table_frame = ctk.CTkFrame(self.trash_window, corner_radius=12)
        table_frame.grid(row=1, column=0, padx=12, pady=10, sticky="nsew")
        table_frame.grid_columnconfigure(0, weight=1)
        table_frame.grid_rowconfigure(0, weight=1)

        columns = ("id", "data", "funcionario", "movimento", "categoria", "quantidade", "excluido_em", "motivo")
        self.tree_trash = ttk.Treeview(table_frame, columns=columns, show="headings", style="Treeview")
        headers = [
            ("id", "ID", 55),
            ("data", "Data", 90),
            ("funcionario", "Funcionário", 170),
            ("movimento", "Mov.", 80),
            ("categoria", "Categoria", 145),
            ("quantidade", "Quantidade", 120),
            ("excluido_em", "Enviado em", 145),
            ("motivo", "Motivo", 280),
        ]
        for col, title, width in headers:
            self.tree_trash.heading(col, text=title)
            self.tree_trash.column(col, width=width, minwidth=50, anchor="w", stretch=(col == "motivo"))
        self.tree_trash.grid(row=0, column=0, padx=8, pady=8, sticky="nsew")
        self.tree_trash.bind("<Double-1>", self.open_selected_trash_details)
        self.tree_trash.bind("<Return>", self.open_selected_trash_details)
        yscroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree_trash.yview)
        yscroll.grid(row=0, column=1, sticky="ns", pady=8)
        xscroll = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree_trash.xview)
        xscroll.grid(row=1, column=0, sticky="ew", padx=8)
        self.tree_trash.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)

        actions = ctk.CTkFrame(self.trash_window, fg_color="transparent")
        actions.grid(row=2, column=0, padx=12, pady=(0, 12), sticky="ew")
        for i in range(5):
            actions.grid_columnconfigure(i, weight=1)
        ctk.CTkButton(actions, text="Restaurar selecionado", command=self.restore_selected_trash).grid(row=0, column=0, padx=4, sticky="ew")
        ctk.CTkButton(actions, text="Excluir definitivamente (Suporte)", command=self.permanent_delete_selected_trash, fg_color="#b91c1c", hover_color="#991b1b").grid(row=0, column=1, padx=4, sticky="ew")
        ctk.CTkButton(actions, text="Esvaziar Lixeira (Suporte)", command=self.empty_trash, fg_color="#7f1d1d", hover_color="#450a0a").grid(row=0, column=2, padx=4, sticky="ew")
        ctk.CTkButton(actions, text="Atualizar", command=self.refresh_trash, fg_color="#475569", hover_color="#334155").grid(row=0, column=3, padx=4, sticky="ew")
        ctk.CTkButton(actions, text="Fechar", command=self.trash_window.destroy, fg_color="#475569", hover_color="#334155").grid(row=0, column=4, padx=4, sticky="ew")

        self.trash_window.protocol("WM_DELETE_WINDOW", self.trash_window.destroy)
        self.refresh_trash()

    def refresh_trash(self):
        if not hasattr(self, "tree_trash") or not self.tree_trash.winfo_exists():
            return
        for item in self.tree_trash.get_children():
            self.tree_trash.delete(item)
        rows = self.db.trash_entries()
        for row in rows:
            try:
                deleted_at = datetime.fromisoformat(row["excluido_em"]).strftime("%d/%m/%Y %H:%M") if row["excluido_em"] else ""
            except Exception:
                deleted_at = row["excluido_em"] or ""
            self.tree_trash.insert(
                "",
                "end",
                values=(
                    row["id"],
                    br_date(row["data"]),
                    row["funcionario"],
                    row["movimento"],
                    row["categoria"],
                    row["quantidade_texto"] or "",
                    deleted_at,
                    row["motivo_exclusao"] or "",
                ),
            )
        self.trash_count_label.configure(text=f"{len(rows)} lançamento(s) armazenado(s) na Lixeira.")

    def selected_trash_entry_id(self):
        if not hasattr(self, "tree_trash"):
            return None
        selected = self.tree_trash.selection()
        if not selected:
            return None
        values = self.tree_trash.item(selected[0], "values")
        return int(values[0]) if values else None

    def open_selected_trash_details(self, _event=None):
        entry_id = self.selected_trash_entry_id()
        if not entry_id:
            return
        row = self.db.entry_by_id(entry_id, include_deleted=True)
        if not row:
            self.refresh_trash()
            return

        try:
            created_at = datetime.fromisoformat(row["criado_em"]).strftime("%d/%m/%Y %H:%M:%S") if row["criado_em"] else ""
        except Exception:
            created_at = row["criado_em"] or ""
        try:
            deleted_at = datetime.fromisoformat(row["excluido_em"]).strftime("%d/%m/%Y %H:%M:%S") if row["excluido_em"] else ""
        except Exception:
            deleted_at = row["excluido_em"] or ""

        hpd = int(row["jornada_horas"] or DEFAULT_HOURS_PER_DAY)
        quantidade = row["quantidade_texto"] or format_minutes(abs(int(row["minutos"] or 0)), hpd)
        details = (
            f"ID do lançamento: {row['id']}\n"
            f"Funcionário: {row['funcionario']}\n"
            f"Regime: {row['regime']}\n"
            f"Jornada diária: {hpd} horas\n"
            f"Data do lançamento: {br_date(row['data'])}\n"
            f"Movimento: {row['movimento']}\n"
            f"Categoria: {row['categoria']}\n"
            f"Quantidade: {quantidade}\n"
            f"Observação: {row['observacao'] or '(sem observação)'}\n"
            f"Origem: {row['origem'] or ''}\n"
            f"Criado em: {created_at}\n\n"
            f"ENVIADO PARA A LIXEIRA\n"
            f"Data e hora: {deleted_at}\n"
            f"Motivo: {row['motivo_exclusao'] or '(não informado)'}"
        )
        self.show_readonly_details("Detalhes do lançamento na Lixeira", details, parent=self.trash_window)

    def restore_selected_trash(self):
        entry_id = self.selected_trash_entry_id()
        if not entry_id:
            self.show_warning("Selecione um lançamento", "Selecione o lançamento que deseja restaurar.", parent=self.trash_window)
            return
        row = self.db.entry_by_id(entry_id, include_deleted=True)
        if not row:
            self.refresh_trash()
            return
        if not self.authorize_admin("restaurar um lançamento da Lixeira"):
            return
        if not self.ask_yes_no(
            "Confirmar restauração",
            f"Restaurar o lançamento de {row['funcionario']} em {br_date(row['data'])}?\n\nEle voltará a participar dos saldos e relatórios.",
            parent=self.trash_window,
        ):
            return
        try:
            create_database_backup("restauracao_lixeira")
            self.db.restore_entry(entry_id)
            self.refresh_trash()
            self.refresh_entries()
            self.refresh_dashboard()
            self.show_info("Restaurado", "O lançamento foi restaurado com sucesso.", parent=self.trash_window)
        except Exception as exc:
            self.show_error("Erro", str(exc), parent=self.trash_window)

    def permanent_delete_selected_trash(self):
        entry_id = self.selected_trash_entry_id()
        if not entry_id:
            self.show_warning("Selecione um lançamento", "Selecione o lançamento que deseja excluir definitivamente.", parent=self.trash_window)
            return
        row = self.db.entry_by_id(entry_id, include_deleted=True)
        if not row:
            self.refresh_trash()
            return
        if not self.authorize_master_support("excluir definitivamente um lançamento da Lixeira", f"TRASH_DELETE:{entry_id}"):
            return
        if not self.ask_yes_no(
            "Exclusão definitiva",
            (
                f"Excluir definitivamente o lançamento de {row['funcionario']} em {br_date(row['data'])}?\n\n"
                "Esta operação não poderá ser desfeita pela Lixeira. Um backup será criado antes da exclusão."
            ),
            parent=self.trash_window,
        ):
            return
        try:
            create_database_backup("exclusao_definitiva_lixeira")
            self.db.permanent_delete_entry(entry_id)
            self.refresh_trash()
            self.show_info("Excluído", "O lançamento foi excluído definitivamente.", parent=self.trash_window)
        except Exception as exc:
            self.show_error("Erro", str(exc), parent=self.trash_window)

    def empty_trash(self):
        count = len(self.db.trash_entries())
        if count == 0:
            self.show_info("Lixeira vazia", "Não há lançamentos na Lixeira.", parent=self.trash_window)
            return
        if not self.authorize_master_support("esvaziar definitivamente a Lixeira", "TRASH_EMPTY"):
            return
        if not self.ask_yes_no(
            "Esvaziar Lixeira",
            (
                f"Excluir definitivamente os {count} lançamento(s) da Lixeira?\n\n"
                "Esta ação não poderá ser desfeita pela Lixeira. Um backup será criado antes da operação."
            ),
            parent=self.trash_window,
        ):
            return
        try:
            create_database_backup("esvaziar_lixeira")
            deleted = self.db.empty_trash()
            self.refresh_trash()
            self.show_info("Lixeira esvaziada", f"{deleted} lançamento(s) foram excluídos definitivamente.", parent=self.trash_window)
        except Exception as exc:
            self.show_error("Erro", str(exc), parent=self.trash_window)

    def open_audit_log(self):
        if not self.authorize_admin("consultar a auditoria"):
            return
        window = self.create_modal(
            "Auditoria administrativa", 1100, 650, resizable=True, grab=False
        )
        self.responsive_minsize(window, 900, 520, floor_w=520, floor_h=340)
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(2, weight=1)

        ctk.CTkLabel(window, text="Auditoria administrativa", font=("Segoe UI Semibold", 21)).grid(row=0, column=0, padx=16, pady=(14, 2), sticky="w")
        ctk.CTkLabel(
            window,
            text="Dê dois cliques em uma linha para abrir os detalhes completos do registro de auditoria.",
            text_color=("#475569", "#cbd5e1"),
            justify="left",
            anchor="w",
        ).grid(row=1, column=0, padx=16, pady=(0, 8), sticky="ew")

        frame = ctk.CTkFrame(window, corner_radius=12)
        frame.grid(row=2, column=0, padx=12, pady=8, sticky="nsew")
        frame.grid_columnconfigure(0, weight=1)
        frame.grid_rowconfigure(0, weight=1)

        columns = ("data", "acao", "lancamento", "detalhes")
        tree = ttk.Treeview(frame, columns=columns, show="headings", style="Treeview")
        for col, title, width, stretch in [
            ("data", "Data e hora", 155, False),
            ("acao", "Ação", 260, False),
            ("lancamento", "Lançamento ID", 105, False),
            ("detalhes", "Detalhes", 540, True),
        ]:
            tree.heading(col, text=title)
            tree.column(col, width=width, minwidth=70, anchor="w", stretch=stretch)
        tree.grid(row=0, column=0, padx=8, pady=8, sticky="nsew")
        yscroll = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        yscroll.grid(row=0, column=1, sticky="ns", pady=8)
        xscroll = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        xscroll.grid(row=1, column=0, sticky="ew", padx=8)
        tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)

        audit_map = {}
        for row in self.db.audit_entries(limit=1000):
            try:
                created = datetime.fromisoformat(row["criado_em"]).strftime("%d/%m/%Y %H:%M:%S")
            except Exception:
                created = row["criado_em"]
            iid = str(row["id"])
            audit_map[iid] = row
            tree.insert(iid=iid, parent="", index="end", values=(created, row["acao"], row["lancamento_id"] or "", row["detalhes"] or ""))

        def open_audit_details(_event=None):
            selected = tree.selection()
            if not selected:
                return
            row = audit_map.get(selected[0])
            if not row:
                return
            try:
                created = datetime.fromisoformat(row["criado_em"]).strftime("%d/%m/%Y %H:%M:%S")
            except Exception:
                created = row["criado_em"] or ""
            lancamento_id = row["lancamento_id"] if row["lancamento_id"] is not None else "—"
            content = (
                f"Registro de auditoria: {row['id']}\n"
                f"Data e hora: {created}\n"
                f"Ação: {row['acao']}\n"
                f"Lançamento ID: {lancamento_id}\n\n"
                f"DETALHES COMPLETOS\n{row['detalhes'] or '(sem detalhes adicionais)'}"
            )
            self.show_readonly_details("Detalhes da Auditoria", content, parent=window)

        tree.bind("<Double-1>", open_audit_details)
        tree.bind("<Return>", open_audit_details)
        ctk.CTkButton(window, text="Fechar", command=window.destroy, width=140).grid(row=3, column=0, padx=12, pady=(4, 12), sticky="e")
        window.protocol("WM_DELETE_WINDOW", window.destroy)

    def toggle_theme(self):
        current = ctk.get_appearance_mode()
        new_theme = "light" if current == "Dark" else "dark"
        ctk.set_appearance_mode(new_theme)
        self.app_config["tema"] = new_theme
        save_app_config(self.app_config)
        self.after_idle(self.apply_treeview_theme)
        if hasattr(self, "employee_scroll"):
            self.after_idle(self.refresh_employees)


    def open_print_signatures_dialog(self):
        """Abre a seleção para impressão com layout responsivo e ações sempre visíveis."""
        if Workbook is None:
            self.show_error(
                "Biblioteca ausente",
                "Instale openpyxl: pip install openpyxl",
            )
            return

        if win32 is None:
            self.show_error(
                "Biblioteca ausente",
                (
                    "Para imprimir diretamente pelo Excel, instale o pywin32:\n\n"
                    "pip install pywin32"
                ),
            )
            return

        if not self.ensure_service_name():
            return

        employees = self.db.employees(only_active=True)
        if not employees:
            self.show_warning(
                "Sem funcionários",
                "Não há funcionários ativos cadastrados para impressão.",
            )
            return

        window = self.create_modal("Imprimir folhas de assinatura", 620, 535)
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(1, weight=1)

        ctk.CTkLabel(
            window,
            text="Imprimir folhas de assinatura",
            font=("Segoe UI Semibold", 20),
        ).grid(row=0, column=0, padx=24, pady=(18, 8), sticky="w")

        body = ctk.CTkScrollableFrame(window, corner_radius=12, height=380)
        body.grid(row=1, column=0, padx=24, sticky="nsew")
        body.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            body,
            text=(
                "Escolha um funcionário ou todos os funcionários ativos, informe o mês, o ano e a quantidade de cópias. "
                "Cada funcionário será impresso separadamente, preservando a paginação correta do próprio relatório."
            ),
            justify="left",
            wraplength=530,
            text_color=("#475569", "#cbd5e1"),
        ).grid(row=0, column=0, padx=10, pady=(8, 16), sticky="w")

        ctk.CTkLabel(body, text="Funcionário").grid(
            row=1, column=0, padx=10, pady=(4, 2), sticky="w"
        )
        employee_values = ["Todos os funcionários ativos"] + [e["nome"] for e in employees]
        combo_employee = ctk.CTkComboBox(body, values=employee_values, height=36, state="readonly")
        combo_employee.grid(row=2, column=0, padx=10, sticky="ew")
        combo_employee.set(employee_values[0])

        period_frame = ctk.CTkFrame(body, fg_color="transparent")
        period_frame.grid(row=3, column=0, padx=10, pady=(16, 0), sticky="ew")
        period_frame.grid_columnconfigure((0, 1), weight=1)

        ctk.CTkLabel(period_frame, text="Mês").grid(row=0, column=0, pady=(0, 2), sticky="w")
        ctk.CTkLabel(period_frame, text="Ano").grid(
            row=0, column=1, padx=(12, 0), pady=(0, 2), sticky="w"
        )

        month_values = MESES_FILTRO[1:]
        combo_month = ctk.CTkComboBox(period_frame, values=month_values, height=36, state="readonly")
        combo_month.grid(row=1, column=0, sticky="ew")
        combo_month.set(month_values[datetime.now().month - 1])

        current_year = str(datetime.now().year)
        year_values = sorted(set(self.db.years() + [current_year]), reverse=True)
        combo_year = ctk.CTkComboBox(period_frame, values=year_values, height=36, state="readonly")
        combo_year.grid(row=1, column=1, padx=(12, 0), sticky="ew")
        combo_year.set(current_year)

        copies_frame = ctk.CTkFrame(body, fg_color="transparent")
        copies_frame.grid(row=4, column=0, padx=10, pady=(16, 0), sticky="ew")
        copies_frame.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(copies_frame, text="Cópias por funcionário").grid(
            row=0, column=0, pady=(0, 2), sticky="w"
        )
        combo_copies = ctk.CTkComboBox(
            copies_frame,
            values=["1 cópia", "2 cópias"],
            height=36,
            state="readonly",
        )
        combo_copies.grid(row=1, column=0, sticky="ew")
        combo_copies.set("1 cópia")

        ctk.CTkLabel(
            body,
            text="A impressão usa a impressora definida como padrão no Windows.",
            text_color=("#64748b", "#94a3b8"),
        ).grid(row=5, column=0, padx=10, pady=(16, 10), sticky="w")

        buttons = ctk.CTkFrame(window, fg_color="transparent")
        buttons.grid(row=2, column=0, padx=24, pady=(12, 18), sticky="ew")
        buttons.grid_columnconfigure((0, 1), weight=1)

        def submit():
            employee_name = combo_employee.get().strip()
            year = combo_year.get().strip()
            month_label = combo_month.get().strip()
            month = month_value(month_label)
            copies = 2 if combo_copies.get().strip().startswith("2") else 1

            if not year or year == "Todos" or month == "Todos":
                self.show_error(
                    "Período obrigatório",
                    "Selecione um mês e um ano para imprimir.",
                    parent=window,
                )
                return

            if employee_name == "Todos os funcionários ativos":
                selected_employees = employees
            else:
                selected_employees = [e for e in employees if e["nome"] == employee_name]

            if not selected_employees:
                self.show_error(
                    "Funcionário não encontrado",
                    "O funcionário selecionado não foi encontrado entre os funcionários ativos.",
                    parent=window,
                )
                return

            period = period_label(year, month)
            total = len(selected_employees)
            plural_func = "s" if total != 1 else ""
            copia_label = "1 cópia" if copies == 1 else "2 cópias"
            if not self.ask_yes_no(
                "Confirmar impressão",
                (
                    f"Enviar o relatório de {total} funcionário{plural_func} para a impressora padrão?\n\n"
                    f"Período: {period}\n"
                    f"Cópias por funcionário: {copia_label}\n\n"
                    "Cada funcionário será impresso separadamente, com paginação própria."
                ),
                parent=window,
            ):
                return

            try:
                window.configure(cursor="watch")
                window.update_idletasks()
                self.print_signature_reports(selected_employees, year, month, copies=copies)
                self.show_info(
                    "Impressão enviada",
                    (
                        f"Relatório de {total} funcionário{plural_func} enviado para a impressora padrão.\n"
                        f"Cópias por funcionário: {copia_label}."
                    ),
                    parent=window,
                )
                window.destroy()
            except Exception as exc:
                self.show_error(
                    "Erro ao imprimir",
                    (
                        "Não foi possível enviar as folhas para impressão.\n\n"
                        f"Detalhes: {exc}"
                    ),
                    parent=window,
                )
            finally:
                try:
                    if window.winfo_exists():
                        window.configure(cursor="")
                except Exception:
                    pass

        ctk.CTkButton(
            buttons, text="Cancelar", command=window.destroy,
            fg_color="#475569", hover_color="#334155"
        ).grid(row=0, column=0, padx=(0, 6), sticky="ew")
        ctk.CTkButton(
            buttons, text="Imprimir", command=submit
        ).grid(row=0, column=1, padx=(6, 0), sticky="ew")
        window.protocol("WM_DELETE_WINDOW", window.destroy)

    def print_signature_reports(self, employees, year: str, month: str, copies: int = 1):
        """Gera um XLSX temporário e imprime cada funcionário separadamente, com paginação própria."""
        if win32 is None:
            raise RuntimeError("pywin32 não está instalado. Execute: pip install pywin32")

        temp_dir = Path(tempfile.gettempdir()) / "ControleFolgas"
        temp_dir.mkdir(parents=True, exist_ok=True)
        temp_path = temp_dir / f"assinaturas_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"

        wb = Workbook()
        default_ws = wb.active
        wb.remove(default_ws)
        self.create_signature_sheets(wb, employees, year=year, month=month, search="")
        if not wb.worksheets:
            raise RuntimeError("Nenhuma folha de assinatura foi gerada.")
        wb.active = 0
        wb.save(temp_path)

        excel = None
        book = None
        try:
            excel = win32.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.ScreenUpdating = False

            book = excel.Workbooks.Open(
                str(temp_path.resolve()),
                UpdateLinks=0,
                ReadOnly=True,
            )
            copies = 2 if int(copies) == 2 else 1

            # Cada funcionário ocupa uma worksheet própria. Imprimir worksheet por worksheet
            # faz o Excel recalcular &P e &N somente para aquela pessoa, evitando que a
            # paginação continue entre funcionários diferentes.
            for sheet_index in range(1, book.Worksheets.Count + 1):
                sheet = book.Worksheets(sheet_index)

                # Reforça diretamente no Excel as opções de impressão antes do PrintOut.
                # Isso evita que configurações locais/driver herdem modo rascunho ou
                # preto-e-branco e prejudiquem o fundo do cabeçalho.
                sheet.PageSetup.BlackAndWhite = False
                sheet.PageSetup.Draft = False

                # Alguns drivers de impressora ignoram Copies=2 quando o Excel envia
                # tudo em um único trabalho. Para garantir fisicamente duas vias, cada
                # cópia é enviada como um trabalho de impressão independente.
                for _copy_number in range(copies):
                    sheet.PrintOut(Copies=1, Collate=True)
        finally:
            if book is not None:
                try:
                    book.Close(SaveChanges=False)
                except Exception:
                    pass
            if excel is not None:
                try:
                    excel.Quit()
                except Exception:
                    pass
            try:
                temp_path.unlink()
            except OSError:
                pass

    def export_excel(self):
        if Workbook is None:
            self.show_error("Biblioteca ausente", "Instale openpyxl: pip install openpyxl")
            return

        if not self.ensure_service_name():
            return

        employee_id, selected_name, year, month, search = self.get_current_filters()

        nome_arquivo = "Relatorio_Controle_Folgas"
        if selected_name and selected_name != "Todos":
            nome_arquivo += f"_{safe_filename(selected_name)}"
        if month and month != "Todos":
            nome_arquivo += f"_{month}_{safe_filename(month_name_from_value(month))}"
        if year and year != "Todos":
            nome_arquivo += f"_{year}"
        nome_arquivo += f"_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

        file = filedialog.asksaveasfilename(
            title="Salvar relatório Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=nome_arquivo,
        )
        if not file:
            return
        try:
            self.create_excel_report(Path(file), employee_id=employee_id, employee_name=selected_name, year=year, month=month, search=search)
            self.show_info("Pronto", f"Relatório exportado com sucesso:\n{file}")
        except Exception as e:
            self.show_error("Erro ao exportar", str(e))

    def create_excel_report(self, path: Path, employee_id=None, employee_name="Todos", year="Todos", month="Todos", search=""):
        hpd = self.db.employee_hours(employee_id) if employee_id else DEFAULT_HOURS_PER_DAY
        wb = Workbook()
        ws = wb.active
        ws.title = "Lançamentos"
        rows = self.db.entries(funcionario_id=employee_id, year=year, month=month, search=search)
        balances = self.employee_running_balances(rows)

        ws.append(["Relatório gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
        ws.append(["Funcionário filtrado", employee_name or "Todos"])
        ws.append(["Ano filtrado", year or "Todos"])
        ws.append(["Mês filtrado", month_name_from_value(month)])
        ws.append(["Busca", search or ""])
        ws.append([])

        headers = [
            "ID", "Data", "Funcionário", "Movimento", "Categoria", "Quantidade",
            "Valor", "Saldo após lançamento", "Observação", "Origem"
        ]
        ws.append(headers)
        for r in rows:
            ws.append([
                r["id"], br_date(r["data"]), r["funcionario"], r["movimento"], r["categoria"], r["quantidade_texto"],
                format_minutes(r["minutos"], int(r["jornada_horas"])), balances.get(r["id"], ""), r["observacao"], r["origem"],
            ])
        self.format_sheet(ws, header_row=7)

        ws2 = wb.create_sheet("Resumo")
        ws2.append(["Funcionário", "Período", "Saldo do período", "Créditos", "Débitos abatidos", "Lançamentos"])
        employees = self.db.employees()
        if employee_id:
            employees = [e for e in employees if e["id"] == employee_id]
        for emp in employees:
            entries = self.db.entries(funcionario_id=emp["id"], year=year, month=month, search=search)
            creditos = sum(r["minutos"] for r in entries if r["minutos"] > 0)
            debitos = sum(r["minutos"] for r in entries if r["minutos"] < 0)
            saldo = creditos + debitos
            emp_hpd = int(emp["jornada_horas"])
            ws2.append([emp["nome"], period_label(year, month), format_minutes(saldo, emp_hpd), format_minutes(creditos, emp_hpd), format_minutes(debitos, emp_hpd), len(entries)])
        self.format_sheet(ws2)

        self.create_signature_sheets(wb, employees, year, month, search)

        wb.save(path)

    def create_signature_sheets(self, wb, employees, year="Todos", month="Todos", search=""):
        """Cria planilhas de assinatura prontas para impressão em A4 paisagem."""
        if not employees:
            return

        service_name = self.db.get_service_name() or "Serviço não informado"

        for idx, emp in enumerate(employees, start=1):
            hpd = int(emp["jornada_horas"])
            sheet_title = "Para Assinatura" if len(employees) == 1 else safe_sheet_name(f"Assinatura {idx}")
            ws = wb.create_sheet(sheet_title)
            ws.sheet_view.showGridLines = False

            rows = self.db.entries(
                funcionario_id=emp["id"],
                year=year,
                month=month,
                search=search,
            )
            rows = sorted(rows, key=lambda r: (r["data"], r["id"]))

            creditos = sum(int(r["minutos"] or 0) for r in rows if r["minutos"] > 0)
            debitos = sum(int(r["minutos"] or 0) for r in rows if r["minutos"] < 0)
            saldo_periodo = creditos + debitos
            data_limite = period_end_iso(year, month)
            saldo_final = (
                self.db.saldo_ate(emp["id"], data_limite)
                if data_limite
                else self.db.saldo(emp["id"])
            )

            # Título.
            ws.merge_cells("A1:G1")
            ws["A1"] = "CONTROLE DE FOLGAS E BANCO DE HORAS"
            ws["A1"].font = Font(bold=True, size=15)
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 24

            # Nome da unidade/serviço cadastrado na primeira abertura do sistema.
            ws.merge_cells("A2:G2")
            ws["A2"] = f"SERVIÇO: {service_name}"
            ws["A2"].font = Font(bold=True, size=11)
            ws["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[2].height = 20

            # Identificação.
            ws["A3"] = "Funcionário:"
            ws.merge_cells("B3:D3")
            ws["B3"] = emp["nome"]
            ws["E3"] = "Gerado em:"
            ws.merge_cells("F3:G3")
            ws["F3"] = datetime.now().strftime("%d/%m/%Y %H:%M")

            ws["A4"] = "Período:"
            ws.merge_cells("B4:D4")
            ws["B4"] = period_label(year, month)
            ws["E4"] = "Regime / jornada:"
            ws.merge_cells("F4:G4")
            ws["F4"] = f"{emp['regime']} - {hpd}h"

            # Resumo com células mescladas para evitar corte.
            ws["A6"] = "Créditos no período:"
            ws.merge_cells("B6:C6")
            ws["B6"] = format_minutes(creditos, hpd)

            ws["D6"] = "Débitos no período:"
            ws["E6"] = format_minutes(abs(debitos), hpd)

            ws["F6"] = "Saldo do período:"
            ws["G6"] = format_minutes(saldo_periodo, hpd)

            ws["A7"] = "Saldo final do banco:"
            ws.merge_cells("B7:C7")
            ws["B7"] = format_minutes(saldo_final, hpd)

            for row_number in (3, 4, 6, 7):
                ws.row_dimensions[row_number].height = 23

            header_row = 9
            headers = [
                "Data",
                "Movimento",
                "Categoria",
                "Observação",
                "Crédito",
                "Débito",
                "Saldo após lançamento",
            ]

            header_fill = PatternFill(fill_type="solid", fgColor="FF000000", bgColor="FF000000")
            header_font = Font(color="FFFFFFFF", bold=True, size=9)
            thin = Side(style="thin", color="FF000000")
            table_border = Border(left=thin, right=thin, top=thin, bottom=thin)

            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=header_row, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = table_border

            ws.row_dimensions[header_row].height = 28

            balances = self.employee_running_balances(rows)
            current_row = header_row + 1

            for r in rows:
                credito = format_minutes(r["minutos"], hpd) if r["minutos"] > 0 else ""
                debito = format_minutes(abs(r["minutos"]), hpd) if r["minutos"] < 0 else ""
                observacao = normalize_spaces(r["observacao"] or "")

                values = [
                    br_date(r["data"]),
                    r["movimento"],
                    r["categoria"],
                    observacao,
                    credito,
                    debito,
                    balances.get(r["id"], ""),
                ]

                for col, value in enumerate(values, start=1):
                    cell = ws.cell(row=current_row, column=col, value=value)
                    cell.border = table_border
                    cell.font = Font(size=9)
                    cell.alignment = Alignment(
                        horizontal="left" if col in (3, 4) else "center",
                        vertical="top",
                        wrap_text=True,
                    )

                # A altura cresce de acordo com a observação, evitando texto cortado.
                estimated_lines = max(1, (len(observacao) // 48) + 1)
                ws.row_dimensions[current_row].height = max(22, min(90, 15 * estimated_lines))
                current_row += 1

            if not rows:
                ws.merge_cells(
                    start_row=current_row,
                    start_column=1,
                    end_row=current_row,
                    end_column=7,
                )
                cell = ws.cell(
                    row=current_row,
                    column=1,
                    value="Nenhum lançamento encontrado para o filtro selecionado.",
                )
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = table_border
                ws.row_dimensions[current_row].height = 25
                current_row += 1

            # Declaração.
            current_row += 2
            declaration_row = current_row
            ws.merge_cells(
                start_row=declaration_row,
                start_column=1,
                end_row=declaration_row,
                end_column=7,
            )
            declaration = ws.cell(
                row=declaration_row,
                column=1,
                value=(
                    "Declaro estar ciente dos lançamentos, das justificativas apresentadas "
                    "e do saldo de banco de horas informado neste relatório."
                ),
            )
            declaration.font = Font(size=10)
            declaration.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[declaration_row].height = 32

            # Assinaturas.
            signature_line_row = declaration_row + 4
            ws.merge_cells(
                start_row=signature_line_row,
                start_column=1,
                end_row=signature_line_row,
                end_column=3,
            )
            ws.merge_cells(
                start_row=signature_line_row,
                start_column=5,
                end_row=signature_line_row,
                end_column=7,
            )
            ws.cell(row=signature_line_row, column=1).value = "_" * 42
            ws.cell(row=signature_line_row, column=5).value = "_" * 42
            ws.cell(row=signature_line_row, column=1).alignment = Alignment(horizontal="center")
            ws.cell(row=signature_line_row, column=5).alignment = Alignment(horizontal="center")

            label_row = signature_line_row + 1
            ws.merge_cells(start_row=label_row, start_column=1, end_row=label_row, end_column=3)
            ws.merge_cells(start_row=label_row, start_column=5, end_row=label_row, end_column=7)
            ws.cell(row=label_row, column=1).value = "Assinatura do funcionário"
            ws.cell(row=label_row, column=5).value = "Assinatura do responsável"
            ws.cell(row=label_row, column=1).alignment = Alignment(horizontal="center")
            ws.cell(row=label_row, column=5).alignment = Alignment(horizontal="center")

            date_row = label_row + 2
            ws.merge_cells(start_row=date_row, start_column=1, end_row=date_row, end_column=3)
            ws.merge_cells(start_row=date_row, start_column=5, end_row=date_row, end_column=7)
            ws.cell(row=date_row, column=1).value = "Data: ____/____/________"
            ws.cell(row=date_row, column=5).value = "Data: ____/____/________"
            ws.cell(row=date_row, column=1).alignment = Alignment(horizontal="center")
            ws.cell(row=date_row, column=5).alignment = Alignment(horizontal="center")

            # Formatação geral.
            bold_labels = ("A3", "E3", "A4", "E4", "A6", "D6", "F6", "A7")
            for address in bold_labels:
                ws[address].font = Font(bold=True, size=9)
                ws[address].alignment = Alignment(vertical="center", wrap_text=True)

            for row_cells in ws.iter_rows(min_row=3, max_row=7, min_col=1, max_col=7):
                for cell in row_cells:
                    if cell.value is not None:
                        cell.alignment = Alignment(vertical="center", wrap_text=True)
                        if cell.coordinate not in bold_labels:
                            cell.font = Font(size=9)

            widths = {
                "A": 12,   # Data
                "B": 14,   # Movimento
                "C": 20,   # Categoria
                "D": 43,   # Observação
                "E": 16,   # Crédito
                "F": 16,   # Débito
                "G": 23,   # Saldo
            }
            for col, width in widths.items():
                ws.column_dimensions[col].width = width

            # Configuração de impressão.
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            # Força impressão completa, sem o modo rascunho/preto-e-branco que alguns Excel/PCs podem herdar.
            ws.page_setup.blackAndWhite = False
            ws.page_setup.draft = False
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.page_setup.horizontalCentered = True
            ws.page_setup.verticalCentered = False
            ws.page_margins = PageMargins(
                left=0.20,
                right=0.20,
                top=0.35,
                bottom=0.35,
                header=0.15,
                footer=0.15,
            )

            ws.print_title_rows = f"1:{header_row}"
            ws.print_area = f"A1:G{date_row}"
            ws.freeze_panes = f"A{header_row + 1}"

            # Rodapé discreto com paginação.
            ws.oddFooter.center.text = "Página &P de &N"
            ws.oddFooter.center.size = 8

    def format_sheet(self, ws, header_row=1):
        header_fill = PatternFill(fill_type="solid", fgColor="FF000000", bgColor="FF000000")
        header_font = Font(color="FFFFFFFF", bold=True)
        thin = Side(style="thin", color="FF000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for cell in ws[header_row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = border
        for row in ws.iter_rows(min_row=header_row + 1):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                max_len = max(max_len, len(str(cell.value or "")))
            ws.column_dimensions[col_letter].width = min(max_len + 2, 55)
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = ws.dimensions


def main():
    startup_splash_update("Inicializando janela principal...")
    app = None
    try:
        app = App()
        app.mainloop()
    finally:
        # Garante que um erro inesperado nunca deixe o splash preso na tela.
        startup_splash_close()


if __name__ == "__main__":
    main()
